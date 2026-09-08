"""Exportación de los reportes del acopio (Plan 087 — Etapa 6).

CONVENCIÓN DEL PROYECTO, NO UNA NUEVA
CSV con separador `;` y BOM UTF-8, igual que `contable/views_reportes.py::exportar_mayor_csv`. El
BOM es lo que hace que Excel abra el archivo sin pedir un asistente de importación.

EN EL ARCHIVO VAN NÚMEROS CRUDOS, NO FORMATO es-AR
En pantalla se usa `|formato_ar`, que es regla inflexible. En un CSV con separador `;` un
`1.234,56` es ambiguo y Excel lo puede leer como texto; en XLSX el formato lo pone la celda con
`number_format`. Formatear acá rompería la planilla en la máquina del organismo.
"""
import csv
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

CERO = Decimal('0.00')

MIME_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
FORMATO_IMPORTE = '#,##0.00'
FORMATO_FECHA = 'dd/mm/yyyy'


def _nombre(base, extension):
    return f"{base}_{timezone.localdate().strftime('%Y%m%d')}.{extension}"


def csv_response(base, encabezados, filas):
    """CSV con `;` y BOM, la convención ya establecida en el ERP."""
    respuesta = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    respuesta['Content-Disposition'] = f'attachment; filename="{_nombre(base, "csv")}"'

    respuesta.write('﻿')
    escritor = csv.writer(respuesta, delimiter=';')
    escritor.writerow(encabezados)
    for fila in filas:
        escritor.writerow(fila)
    return respuesta


def xlsx_response(base, titulo, encabezados, filas, *, columnas_importe=(), columnas_fecha=(),
                  totales=None):
    """XLSX con encabezado estilado y formato numérico en la celda.

    `columnas_importe` y `columnas_fecha` son índices 0-based: el formato lo pone la celda, no el
    texto, que es lo que permite que el organismo sume la columna en su propia planilla.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:31]                       # Excel no admite hojas con nombre más largo

    relleno = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fuente_encabezado = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fuente_datos = Font(name="Arial", size=9)
    linea = Side(border_style="thin", color="CBD5E1")
    borde = Border(top=linea, bottom=linea, left=linea, right=linea)

    ws.append(list(encabezados))
    for columna in range(1, len(encabezados) + 1):
        celda = ws.cell(row=1, column=columna)
        celda.fill = relleno
        celda.font = fuente_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for fila in filas:
        ws.append(_valores_para_excel(fila))

    ultima = ws.max_row
    if totales is not None:
        ws.append(_valores_para_excel(totales))
        ultima = ws.max_row
        for columna in range(1, len(encabezados) + 1):
            ws.cell(row=ultima, column=columna).font = Font(name="Arial", size=9, bold=True)

    for numero_fila in range(2, ultima + 1):
        for columna in range(1, len(encabezados) + 1):
            celda = ws.cell(row=numero_fila, column=columna)
            celda.border = borde
            if celda.font.bold is not True:
                celda.font = fuente_datos
            if (columna - 1) in columnas_importe:
                celda.number_format = FORMATO_IMPORTE
            elif (columna - 1) in columnas_fecha:
                celda.number_format = FORMATO_FECHA

    for columna, encabezado in enumerate(encabezados, start=1):
        ancho = max(11, min(38, len(str(encabezado)) + 4))
        ws.column_dimensions[get_column_letter(columna)].width = ancho

    ws.freeze_panes = 'A2'

    respuesta = HttpResponse(content_type=MIME_XLSX)
    respuesta['Content-Disposition'] = f'attachment; filename="{_nombre(base, "xlsx")}"'
    wb.save(respuesta)
    return respuesta


def _valores_para_excel(fila):
    """`Decimal` va como `float` para que Excel lo trate como número y no como texto."""
    return [float(v) if isinstance(v, Decimal) else v for v in fila]
