"""Plan 087 — humo de las pantallas de reportes y de sus exportaciones.

Fija además las reglas transversales del proyecto: filtro de `condic` en todo reporte con
importes, formato es-AR en pantalla, y números crudos en el archivo exportado.
"""
from datetime import date
from decimal import Decimal

from django.urls import reverse

from .test_plan087_reportes import ReportesBaseTestCase


class PantallasBaseTestCase(ReportesBaseTestCase):

    def setUp(self):
        super().setUp()
        self.user.is_staff = True
        self.user.save(update_fields=['is_staff'])
        self.client.force_login(self.user)
        sesion = self.client.session
        sesion['empresa_id'] = self.empresa.id
        sesion['sucursal_id'] = self.sucursal.id
        sesion.save()


class PortadaTests(PantallasBaseTestCase):

    def test_la_portada_lista_los_reportes(self):
        r = self.client.get(reverse('agro_reportes'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Planilla FET')
        self.assertContains(r, 'DDJJ de Existencias')
        self.assertContains(r, 'Libro de Retenciones')
        self.assertContains(r, 'Tablero de Margen')


class PlanillaFETPantallaTests(PantallasBaseTestCase):

    def test_renderiza_con_totales(self):
        self._liquidar()

        r = self.client.get(reverse('agro_reporte_fet'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'PRODUCTOR RI')
        self.assertContains(r, 'Totales')

    def test_ofrece_el_filtro_de_condicion_y_arranca_en_real(self):
        """Los reportes oficiales son declaraciones: lo que se declara es la lente fiscal."""
        r = self.client.get(reverse('agro_reporte_fet'))

        self.assertContains(r, 'name="condic"')
        self.assertContains(r, 'Presupuestado')
        self.assertContains(r, '<option value="1" selected>Real</option>', html=False)

    def test_los_importes_salen_en_formato_argentino(self):
        self._liquidar()

        r = self.client.get(reverse('agro_reporte_fet'))

        self.assertContains(r, '1.000.000,00')

    def test_la_grilla_respeta_el_filtro(self):
        self._liquidar()

        con = self.client.get(reverse('agro_reporte_fet_grilla'), {'condic': '1'})
        sin = self.client.get(reverse('agro_reporte_fet_grilla'), {'condic': '2'})

        self.assertContains(con, 'PRODUCTOR RI')
        self.assertContains(sin, 'No hay liquidaciones confirmadas')

    def test_csv_tiene_bom_y_punto_y_coma(self):
        """Convención del proyecto: es lo que hace que Excel lo abra sin asistente."""
        self._liquidar()

        r = self.client.get(reverse('agro_reporte_fet_csv'))

        self.assertEqual(r['Content-Type'], 'text/csv; charset=utf-8-sig')
        self.assertIn('attachment; filename="planilla_fet_', r['Content-Disposition'])
        contenido = r.content.decode('utf-8-sig')
        self.assertTrue(r.content.startswith('﻿'.encode()))
        self.assertIn('Romaneo;Asiento;Letra', contenido)

    def test_el_csv_lleva_numeros_crudos_y_no_formato_argentino(self):
        """Un `1.234,56` dentro de un CSV con separador `;` es ambiguo y Excel lo lee como texto."""
        self._liquidar()

        contenido = self.client.get(reverse('agro_reporte_fet_csv')).content.decode('utf-8-sig')

        self.assertIn('1000000.00', contenido)
        self.assertNotIn('1.000.000,00', contenido)

    def test_xlsx_responde_con_el_mime_correcto(self):
        self._liquidar()

        r = self.client.get(reverse('agro_reporte_fet_xlsx'))

        self.assertEqual(
            r['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('attachment; filename="planilla_fet_', r['Content-Disposition'])
        self.assertTrue(r.content.startswith(b'PK'))          # un .xlsx es un ZIP

    def test_el_xlsx_trae_la_fila_de_totales(self):
        import io

        import openpyxl

        self._liquidar()
        r = self.client.get(reverse('agro_reporte_fet_xlsx'))
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb.active

        self.assertEqual(ws.cell(row=1, column=1).value, 'Romaneo')
        self.assertEqual(ws.cell(row=ws.max_row, column=1).value, 'TOTALES')
        # El importe va como NÚMERO, no como texto: el organismo tiene que poder sumarlo.
        self.assertIsInstance(ws.cell(row=2, column=14).value, (int, float))


class ResumenAcopioPantallaTests(PantallasBaseTestCase):

    def test_renderiza(self):
        self._romaneo_confirmado(kilos='1000')

        r = self.client.get(reverse('agro_reporte_acopio'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'BURLEY')
        self.assertContains(r, 'B1F')

    def test_ofrece_el_filtro_de_condicion(self):
        r = self.client.get(reverse('agro_reporte_acopio'))

        self.assertContains(r, 'name="condic"')
        self.assertContains(r, 'Presupuestado')

    def test_csv(self):
        self._romaneo_confirmado(kilos='1000')

        contenido = self.client.get(
            reverse('agro_reporte_acopio_csv')).content.decode('utf-8-sig')

        self.assertIn('Variedad;Grupo;', contenido)
        self.assertIn('SUBTOTAL', contenido)


class ExistenciasPantallaTests(PantallasBaseTestCase):

    def test_renderiza_con_la_fecha_de_hoy_por_defecto(self):
        self._romaneo_confirmado(kilos='1000')

        r = self.client.get(reverse('agro_reporte_existencias'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'BURLEY')

    def test_una_fecha_anterior_no_declara_nada(self):
        self._romaneo_confirmado(kilos='1000')          # 2026-09-01

        r = self.client.get(reverse('agro_reporte_existencias_grilla'),
                            {'fecha': '2026-08-31', 'condic': '1'})

        self.assertContains(r, 'No hay movimientos de tabaco')

    def test_ofrece_el_filtro_de_condicion(self):
        r = self.client.get(reverse('agro_reporte_existencias'))

        self.assertContains(r, 'name="condic"')

    def test_csv(self):
        self._romaneo_confirmado(kilos='1000')

        contenido = self.client.get(
            reverse('agro_reporte_existencias_csv'),
            {'fecha': '2026-09-30'}).content.decode('utf-8-sig')

        self.assertIn('Variedad;Galpón;', contenido)
        self.assertIn('1000.00', contenido)


class LibroRetencionesPantallaTests(PantallasBaseTestCase):

    def test_renderiza_con_totales_por_organismo(self):
        self._liquidar()

        r = self.client.get(reverse('agro_reporte_retenciones'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'A depositar por organismo')
        self.assertContains(r, 'Al liquidar')

    def test_muestra_las_dos_procedencias_despues_del_pago(self):
        liq = self._liquidar()
        self._pagar([liq], liq.total)

        r = self.client.get(reverse('agro_reporte_retenciones_grilla'))

        self.assertContains(r, 'Al liquidar')
        self.assertContains(r, 'Al pagar')

    def test_el_combo_de_conceptos_sale_del_maestro(self):
        """Es extensible por diseño: una lista fija se desactualizaría al primer alta."""
        r = self.client.get(reverse('agro_reporte_retenciones'))

        self.assertContains(r, 'EEAOC')
        self.assertContains(r, 'RET-GCIAS')

    def test_ofrece_el_filtro_de_condicion(self):
        r = self.client.get(reverse('agro_reporte_retenciones'))

        self.assertContains(r, 'name="condic"')

    def test_csv(self):
        self._liquidar()

        contenido = self.client.get(
            reverse('agro_reporte_retenciones_csv')).content.decode('utf-8-sig')

        self.assertIn('Fecha;Período;Origen;', contenido)
        self.assertIn('EEAOC', contenido)


class TableroPantallaTests(PantallasBaseTestCase):

    def _lote_vendido(self, kilos='1000', precio='1500'):
        from facturacion.models import TipoComprobante, Venta, VentaItem

        from verticalidades.agricola.tabaco.services import lotes as lote_svc

        lote = self._lote_armado(kilos=kilos)
        tipo, _ = TipoComprobante.objects.get_or_create(
            codigo='087', defaults={'detalle': 'FACTURA A', 'signo': 1})
        venta = Venta.objects.create(
            empresa=self.empresa, sucursal=self.sucursal, tipo=tipo, punto=2, numero=1,
            fecha=date(2026, 9, 20), cliente=self.productor, usuario=self.user, estado=0, condic=1)
        VentaItem.objects.create(venta=venta, producto=self.producto, cantidad=Decimal(kilos),
                                 precio_unitario=Decimal(precio), iva_alicuota=Decimal('21'),
                                 total=Decimal(kilos) * Decimal(precio))
        return lote_svc.asignar_venta(lote, venta, self.user)

    def test_renderiza_agrupado_por_campania(self):
        self._lote_vendido()

        r = self.client.get(reverse('agro_reporte_tablero'))

        self.assertEqual(r.status_code, 200)
        self.assertContains(r, self.campania.codigo)
        self.assertContains(r, 'Totales')

    def test_cambiar_la_agrupacion_cambia_las_columnas(self):
        self._lote_vendido()

        por_clase = self.client.get(reverse('agro_reporte_tablero_grilla'),
                                    {'agrupar': 'clase'})
        por_productor = self.client.get(reverse('agro_reporte_tablero_grilla'),
                                        {'agrupar': 'productor'})

        self.assertContains(por_clase, 'B1F')
        self.assertContains(por_clase, 'Coef.')
        self.assertContains(por_productor, 'PRODUCTOR RI')
        self.assertContains(por_productor, 'CUIT')

    def test_ofrece_el_filtro_de_condicion(self):
        r = self.client.get(reverse('agro_reporte_tablero'))

        self.assertContains(r, 'name="condic"')

    def test_los_importes_salen_en_formato_argentino(self):
        self._lote_vendido()

        r = self.client.get(reverse('agro_reporte_tablero'))

        self.assertContains(r, '1.500.000,00')

    def test_csv(self):
        self._lote_vendido()

        contenido = self.client.get(
            reverse('agro_reporte_tablero_csv')).content.decode('utf-8-sig')

        self.assertIn('Campaña;Lotes;Fardos;', contenido)


class AislamientoPorEmpresaTests(PantallasBaseTestCase):
    """Regla inflexible: ningún reporte puede mostrar datos de otra empresa."""

    def test_los_reportes_no_ven_otra_empresa(self):
        self._liquidar()

        from empresas.models import Empresa
        ajena = Empresa.objects.create(nombre="OTRA", cuit="30999999993",
                                       tipo_actividad='AGRICOLA')
        sesion = self.client.session
        sesion['empresa_id'] = ajena.id
        sesion.save()

        for nombre in ('agro_reporte_fet_grilla', 'agro_reporte_acopio_grilla',
                       'agro_reporte_retenciones_grilla', 'agro_reporte_tablero_grilla'):
            with self.subTest(reporte=nombre):
                self.assertNotContains(self.client.get(reverse(nombre)), 'PRODUCTOR RI')
