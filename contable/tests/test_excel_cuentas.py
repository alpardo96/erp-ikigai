import io
import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.contrib.auth.models import User
from empresas.models import Empresa
from contable.models import Cuenta
from contable.services.excel_service import (
    generar_excel_cuentas,
    procesar_captura_excel_cuentas
)

class ExcelCuentasTestCase(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="EMPRESA DE TEST PLAN DE CUENTAS",
            cuit="30111111118"
        )
        self.user = User.objects.create_user(username="testuser", password="password123")

        # Cuentas iniciales
        self.cta_activo = Cuenta.objects.create(
            empresa=self.empresa,
            jerarquia="1",
            cuenta="ACTIVO",
            imputable=0,
            tipo="A"
        )
        self.cta_caja = Cuenta.objects.create(
            empresa=self.empresa,
            jerarquia="1.1",
            cuenta="CAJA Y BANCOS",
            imputable=0,
            tipo="A",
            sumariza=self.cta_activo
        )
        self.cta_caja_efectivo = Cuenta.objects.create(
            empresa=self.empresa,
            jerarquia="1.1.01",
            cuenta="CAJA GENERAL",
            imputable=1,
            tipo="A",
            tipo_disponibilidad="EFE",
            sumariza=self.cta_caja
        )

    def test_generar_excel_cuentas(self):
        queryset = Cuenta.objects.filter(empresa=self.empresa).order_by('jerarquia')
        wb = generar_excel_cuentas(queryset)
        
        ws = wb.active
        self.assertEqual(ws.title, "Plan de Cuentas")
        # Verificar fila de encabezados
        self.assertEqual(ws.cell(row=1, column=1).value, "ID")
        self.assertEqual(ws.cell(row=1, column=2).value, "Sumariza ID")
        self.assertEqual(ws.cell(row=1, column=3).value, "Jerarquía")
        self.assertEqual(ws.cell(row=1, column=4).value, "Nombre Cuenta")

        # Verificar filas de datos (3 cuentas creadas)
        self.assertEqual(ws.max_row, 4)
        # Fila 4: Cuenta 1.1.01 (CAJA GENERAL) que sumariza en 1.1 (CAJA Y BANCOS)
        self.assertEqual(ws.cell(row=4, column=2).value, self.cta_caja.id)
        self.assertEqual(ws.cell(row=4, column=3).value, "1.1.01")
        self.assertEqual(ws.cell(row=4, column=4).value, "CAJA GENERAL")

    def test_capturar_excel_actualizar_y_crear_cuentas(self):
        # Crear un archivo Excel en memoria con 3 filas:
        # Fila 1: Actualizar CAJA GENERAL (ID existente)
        # Fila 2: Crear nueva cuenta sin ID (celda vacía)
        # Fila 3: Crear nueva cuenta con ID inventado inexistente (el sistema debe asignarle ID automático)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Sumariza ID", "Jerarquía", "Nombre Cuenta", "Imputable (1=Sí, 0=No)", "Tipo (A/P/N/R)", "Código (Legacy)", "RG 830", "Tipo Disponibilidad"])

        # Fila 1 (Actualizar ID existente)
        ws.append([self.cta_caja_efectivo.id, self.cta_caja.id, "1.1.01", "caja general en pesos modificada", 1, "A", 101, "", "EFE"])
        # Fila 2 (Nueva cuenta sin ID pero con Sumariza ID explícito)
        ws.append(["", self.cta_caja.id, "1.1.02", "banco nacion cuenta corriente", 1, "A", 102, "", "BCO"])
        # Fila 3 (Nueva cuenta con ID inventado inexistente 99999)
        ws.append([99999, self.cta_caja.id, "1.1.03", "banco galicia cta cte", 1, "A", 103, "", "BCO"])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        resultado = procesar_captura_excel_cuentas(self.empresa, self.user, buffer)

        self.assertEqual(resultado['actualizados'], 1)
        self.assertEqual(resultado['creados'], 2)
        self.assertEqual(len(resultado['errores']), 0)

        # 1. Verificar actualización de CAJA GENERAL
        self.cta_caja_efectivo.refresh_from_db()
        self.assertEqual(self.cta_caja_efectivo.cuenta, "CAJA GENERAL EN PESOS MODIFICADA")
        self.assertEqual(self.cta_caja_efectivo.codigo, 101)
        self.assertEqual(self.cta_caja_efectivo.sumariza, self.cta_caja)

        # 2. Verificar creación de cuenta sin ID
        nueva_banco = Cuenta.objects.get(empresa=self.empresa, jerarquia="1.1.02")
        self.assertEqual(nueva_banco.cuenta, "BANCO NACION CUENTA CORRIENTE")
        self.assertEqual(nueva_banco.tipo_disponibilidad, "BCO")
        self.assertEqual(nueva_banco.sumariza, self.cta_caja)
        self.assertIsNotNone(nueva_banco.id)

        # 3. Verificar creación de cuenta con ID inexistente 99999 (ID NO debe ser 99999 sino el asignado por BD)
        nueva_galicia = Cuenta.objects.get(empresa=self.empresa, jerarquia="1.1.03")
        self.assertEqual(nueva_galicia.cuenta, "BANCO GALICIA CTA CTE")
        self.assertEqual(nueva_galicia.sumariza, self.cta_caja)
        self.assertNotEqual(nueva_galicia.id, 99999)
        self.assertIsNotNone(nueva_galicia.id)

    def test_views_excel_exportar_y_modal(self):
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

        self.client.force_login(self.user)

        # Test exportar Excel completo
        response_export = self.client.get(reverse('config_cuentacontable_exportar_excel'))
        self.assertEqual(response_export.status_code, 200)
        self.assertEqual(
            response_export['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        # Test abrir modal capturar
        response_modal = self.client.get(reverse('config_cuentacontable_modal_capturar_excel'))
        self.assertEqual(response_modal.status_code, 200)
        self.assertContains(response_modal, "Recapturar / Importar Plan de Cuentas")
