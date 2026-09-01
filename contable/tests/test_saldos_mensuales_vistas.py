"""Balance de Saldos Mensuales — pruebas de vistas y drill-down (Plan 047, fases 4-6).

Cubre los puntos 21 y 22 del plan de pruebas, más la integración de la página y el partial HTMX.
"""
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from contable.models import Cuenta
from contable.services.asientos import crear_asiento
from empresas.models import Empresa, Ejercicio
from usuarios.models import Perfil

User = get_user_model()
CERO = Decimal("0.00")


def _neto(respuesta):
    """Σ(debe − haber) de todo lo que `libro_mayor_rows` devolvió.

    La vista agrupa por cuenta en `cuentas_data`, así que hay que aplanar antes de sumar.
    """
    total = Decimal("0.00")
    for bloque in respuesta.context['cuentas_data']:
        for mov in bloque['movimientos']:
            total += mov.debe - mov.haber
    return total


class _BaseVistas(TestCase):
    """Fixtures compartidas. Sin tests propios: heredar de una TestCase con métodos de test
    los volvería a correr en cada subclase."""

    def setUp(self):
        self.empresa = Empresa.objects.create(
            nombre="EMPRESA VISTAS SA", cuit="30111111117",
            direccion="Calle 1", correo="v@e.com",
        )
        self.ejercicio = Ejercicio.objects.create(
            empresa=self.empresa, ejercicio="Ejercicio 2026",
            inicio=timezone.datetime(2026, 1, 1).date(),
            cierre=timezone.datetime(2026, 12, 31).date(),
        )
        self.user = User.objects.create_user(username="vistas_user", password="pw123456")
        self.perfil = Perfil.objects.create(usuario=self.user)
        self.perfil.empresas.add(self.empresa)

        self.caja = Cuenta.objects.create(
            empresa=self.empresa, jerarquia="111001", cuenta="CAJA", imputable=1, tipo="A")
        self.ventas = Cuenta.objects.create(
            empresa=self.empresa, jerarquia="410101", cuenta="VENTAS", imputable=1, tipo="R")

        self.client.login(username="vistas_user", password="pw123456")
        session = self.client.session
        session['empresa_id'] = self.empresa.pk
        session['ejercicio_id'] = self.ejercicio.pk
        session.save()

    def _asiento(self, mes, condic=1, monto="1000.00", dia=15):
        return crear_asiento(
            empresa=self.empresa,
            fecha=timezone.datetime(2026, mes, dia).date(),
            concepto=f"MOV {mes}", condic=condic,
            lineas=[
                {'cuenta': self.caja, 'debe': Decimal(monto), 'haber': CERO},
                {'cuenta': self.ventas, 'debe': CERO, 'haber': Decimal(monto)},
            ],
            usuario=self.user,
        )

    def _apertura(self, monto="900.00"):
        return crear_asiento(
            empresa=self.empresa, fecha=self.ejercicio.inicio,
            concepto="APERTURA", condic=5,
            lineas=[
                {'cuenta': self.caja, 'debe': Decimal(monto), 'haber': CERO},
                {'cuenta': self.ventas, 'debe': CERO, 'haber': Decimal(monto)},
            ],
            usuario=self.user,
        )


class SaldosMensualesVistasTests(_BaseVistas):
    # -- Página y partial -------------------------------------------------------------------

    def test_pagina_responde_y_arma_las_columnas_del_ejercicio(self):
        self._asiento(mes=3)
        response = self.client.get(reverse('contable_saldos_mensuales'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['periodos']), 12)
        self.assertEqual(response.context['periodos'][0]['clave'], "202601")
        self.assertContains(response, "Ene-26")
        self.assertContains(response, "Dic-26")

    def test_importes_en_formato_es_ar(self):
        self._asiento(mes=3, monto="1234567.89")
        response = self.client.get(reverse('contable_saldos_mensuales'))
        self.assertContains(response, "1.234.567,89")

    def test_htmx_devuelve_solo_el_partial(self):
        self._asiento(mes=3)
        response = self.client.get(
            reverse('contable_saldos_mensuales'), HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        # El partial no trae el layout base.
        self.assertNotContains(response, "<html")
        self.assertContains(response, "filtros-saldos-mensuales")

    def test_endpoint_de_datos_aplica_los_filtros(self):
        self._asiento(mes=2, condic=1, monto="100.00")
        self._asiento(mes=2, condic=2, monto="20.00")

        url = reverse('saldos_mensuales_datos')
        response = self.client.get(url, {'filtros_aplicados': '1', 'condic': ['1']})
        fila = next(f for f in response.context['filas'] if f['cuenta'].id == self.caja.id)
        self.assertEqual(fila['meses'][1], Decimal("100.00"))

        response = self.client.get(url, {'filtros_aplicados': '1', 'condic': ['1', '2']})
        fila = next(f for f in response.context['filas'] if f['cuenta'].id == self.caja.id)
        self.assertEqual(fila['meses'][1], Decimal("120.00"))

    def test_destildar_todas_las_condiciones_avisa_y_no_consulta(self):
        self._asiento(mes=2)
        response = self.client.get(reverse('saldos_mensuales_datos'), {'filtros_aplicados': '1'})
        self.assertEqual(response.context['filas'], [])
        self.assertTrue(response.context['aviso'])

    def test_empresa_sin_ejercicios_no_rompe(self):
        Ejercicio.objects.filter(empresa=self.empresa).delete()
        response = self.client.get(reverse('contable_saldos_mensuales'))
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context['ejercicio'])
        self.assertTrue(response.context['aviso'])

    def test_requiere_login(self):
        self.client.logout()
        response = self.client.get(reverse('contable_saldos_mensuales'))
        self.assertEqual(response.status_code, 302)

    # -- Drill-down (21 y 22) ---------------------------------------------------------------

    def test_21_los_querystrings_llevan_el_condic_explicito(self):
        self._asiento(mes=3)
        self._apertura()
        response = self.client.get(reverse('contable_saldos_mensuales'))

        qs_meses = response.context['qs_meses']
        qs_apertura = response.context['qs_apertura']
        qs_total = response.context['qs_total']

        self.assertIn("condic=1", qs_meses)
        self.assertNotIn("condic=5", qs_meses)      # la apertura no entra en una columna mensual
        self.assertIn("condic=5", qs_apertura)
        self.assertNotIn("condic=1", qs_apertura)   # ni un movimiento en la columna Apertura
        self.assertIn("condic=1", qs_total)
        self.assertIn("condic=5", qs_total)
        # El ejercicio viaja siempre, para que el mayor no traiga asientos de otro.
        for qs in (qs_meses, qs_apertura, qs_total):
            self.assertIn(f"ejercicio_id={self.ejercicio.pk}", qs)

    def test_21b_la_celda_enlaza_al_mayor_con_el_rango_del_mes(self):
        self._asiento(mes=3)
        response = self.client.get(reverse('contable_saldos_mensuales'))
        html = response.content.decode()
        self.assertIn("fecha_desde=2026-03-01&amp;fecha_hasta=2026-03-31", html)
        self.assertIn(reverse('mayor_cuenta_modal', args=[self.caja.id]), html)

    def test_22_el_mayor_concilia_exactamente_con_la_celda(self):
        """El caso que se rompería: apertura fechada dentro del primer mes.

        Sin el `condic` explícito en el link, el mayor de enero mostraría también el asiento de
        apertura —su fecha cae en el rango— mientras la celda lo excluye.
        """
        self._apertura(monto="900.00")
        self._asiento(mes=1, dia=20, monto="100.00")

        pagina = self.client.get(reverse('contable_saldos_mensuales'))
        fila = next(f for f in pagina.context['filas'] if f['cuenta'].id == self.caja.id)
        valor_celda = fila['meses'][0]
        self.assertEqual(valor_celda, Decimal("100.00"))

        from urllib.parse import parse_qsl
        params = dict(parse_qsl(pagina.context['qs_meses']))
        params['condic'] = pagina.context['condics']
        params['fecha_desde'] = '2026-01-01'
        params['fecha_hasta'] = '2026-01-31'

        mayor = self.client.get(
            reverse('libro_mayor_rows'), {**params, 'cuenta_id': self.caja.id})
        self.assertEqual(mayor.status_code, 200)

        self.assertEqual(_neto(mayor), valor_celda)

        # Y la celda Apertura concilia con su propio link.
        params_ap = dict(parse_qsl(pagina.context['qs_apertura']))
        params_ap['fecha_desde'] = '2026-01-01'
        params_ap['fecha_hasta'] = '2026-12-31'
        mayor_ap = self.client.get(
            reverse('libro_mayor_rows'), {**params_ap, 'cuenta_id': self.caja.id})
        self.assertEqual(_neto(mayor_ap), fila['apertura'])

    def test_22b_el_mayor_filtra_por_ejercicio(self):
        """Un asiento de otro ejercicio con fecha solapada no debe aparecer en el mayor."""
        otro = Ejercicio.objects.create(
            empresa=self.empresa, ejercicio="Ejercicio 2027",
            inicio=timezone.datetime(2027, 1, 1).date(),
            cierre=timezone.datetime(2027, 12, 31).date(),
        )
        self._asiento(mes=4, monto="100.00")
        crear_asiento(
            empresa=self.empresa, fecha=timezone.datetime(2026, 4, 10).date(),
            concepto="AJENO", condic=1, ejercicio=otro,
            lineas=[
                {'cuenta': self.caja, 'debe': Decimal("777.00"), 'haber': CERO},
                {'cuenta': self.ventas, 'debe': CERO, 'haber': Decimal("777.00")},
            ],
            usuario=self.user,
        )

        respuesta = self.client.get(reverse('libro_mayor_rows'), {
            'cuenta_id': self.caja.id, 'condic': ['1'],
            'ejercicio_id': self.ejercicio.pk,
            'fecha_desde': '2026-04-01', 'fecha_hasta': '2026-04-30',
        })
        self.assertEqual(_neto(respuesta), Decimal("100.00"))

    def test_22c_el_mayor_filtra_por_modulo(self):
        self._asiento(mes=5, monto="100.00")
        crear_asiento(
            empresa=self.empresa, fecha=timezone.datetime(2026, 5, 10).date(),
            concepto="COMPRAS", condic=1, modulo=5,
            lineas=[
                {'cuenta': self.caja, 'debe': Decimal("50.00"), 'haber': CERO},
                {'cuenta': self.ventas, 'debe': CERO, 'haber': Decimal("50.00")},
            ],
            usuario=self.user,
        )
        respuesta = self.client.get(reverse('libro_mayor_rows'), {
            'cuenta_id': self.caja.id, 'condic': ['1'], 'modulo': 5,
            'fecha_desde': '2026-05-01', 'fecha_hasta': '2026-05-31',
        })
        self.assertEqual(_neto(respuesta), Decimal("50.00"))


class SaldosMensualesExcelTests(_BaseVistas):
    """12, 13 y 28: el × −1 vive SÓLO acá, y sólo en modo 'Sólo Resultados'."""

    def _hoja(self, **filtros):
        from io import BytesIO
        import openpyxl

        respuesta = self.client.get(reverse('exportar_saldos_mensuales_excel'), filtros)
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        return respuesta, openpyxl.load_workbook(BytesIO(respuesta.content)).active

    def _buscar(self, ws, detalle):
        """Fila del Excel cuya columna D (Detalle) coincide."""
        for fila in ws.iter_rows(min_row=5):
            if fila[3].value == detalle:
                return fila
        return None

    def test_28_el_excel_tiene_una_columna_por_periodo(self):
        self._asiento(mes=3)
        respuesta, ws = self._hoja()
        self.assertIn('attachment;', respuesta['Content-Disposition'])
        self.assertIn('saldos_mensuales_', respuesta['Content-Disposition'])

        # Encabezado (fila 4): A..E fijas, F=Apertura, 12 meses, Total, Tipo y las 4 de clasif.
        encabezado = [c.value for c in ws[4]]
        self.assertEqual(encabezado[:6], ["Codigo", "Sumariza", "Jerarquia", "Detalle", "Imp", "Apertura"])
        self.assertEqual(encabezado[6:18], [int(f"2026{m:02d}") for m in range(1, 13)])
        self.assertEqual(encabezado[18:24], ["Total", "Tipo", "Bce", "Pres", "Econ", "Fciero"])

    def test_12_modo_todas_no_invierte(self):
        self._asiento(mes=3, monto="1000.00")   # Caja al debe, Ventas al haber
        _, ws = self._hoja()
        ventas = self._buscar(ws, "VENTAS")
        # Signo natural: un ingreso está en el haber, o sea negativo.
        self.assertEqual(ventas[18].value, -1000.0)         # Total
        self.assertEqual(ventas[6 + 2].value, -1000.0)      # marzo

    def test_12b_modo_resultados_invierte(self):
        self._asiento(mes=3, monto="1000.00")
        _, ws = self._hoja(filtros_aplicados='1', condic=['1'], alcance='resultados')
        ventas = self._buscar(ws, "VENTAS")
        # Invertido: el ingreso sale POSITIVO, que es lo que el usuario necesita en el informe.
        self.assertEqual(ventas[18].value, 1000.0)
        self.assertEqual(ventas[6 + 2].value, 1000.0)
        # Y la cuenta patrimonial ni aparece: el alcance filtra a tipo='R'.
        self.assertIsNone(self._buscar(ws, "CAJA"))

    def test_13_totales_en_modo_resultados_son_la_ganancia(self):
        """Ingresos 1000 y egresos 400 ⇒ el Excel muestra +600 de ganancia en el mes."""
        egresos = Cuenta.objects.create(
            empresa=self.empresa, jerarquia="510001", cuenta="COSTO",
            imputable=1, tipo="R")
        self._asiento(mes=3, monto="1000.00")
        crear_asiento(
            empresa=self.empresa, fecha=timezone.datetime(2026, 3, 20).date(),
            concepto="COSTO", condic=1,
            lineas=[
                {'cuenta': egresos, 'debe': Decimal("400.00"), 'haber': CERO},
                {'cuenta': self.caja, 'debe': CERO, 'haber': Decimal("400.00")},
            ],
            usuario=self.user,
        )
        _, ws = self._hoja(filtros_aplicados='1', condic=['1'], alcance='resultados')
        totales = self._buscar(ws, "TOTALES:")
        self.assertEqual(totales[6 + 2].value, 600.0)   # marzo
        self.assertEqual(totales[18].value, 600.0)      # total del ejercicio

    def test_13b_totales_en_modo_todas_cierran_en_cero(self):
        self._asiento(mes=3, monto="1000.00")
        _, ws = self._hoja()
        totales = self._buscar(ws, "TOTALES:")
        for i in range(12):
            self.assertEqual(totales[6 + i].value, 0.0)
        self.assertEqual(totales[18].value, 0.0)
