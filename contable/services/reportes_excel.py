from datetime import date

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone


def _configurar_hoja(ws, titulo, empresa, ejercicio=None, fecha_desde=None, fecha_hasta=None):
    """Aplica estilos comunes y titulo a la hoja de Excel."""
    # Titulos
    ws.merge_cells('A1:F1')
    empresa_nombre = empresa.nombre.upper() if empresa else ''
    if ejercicio:
        ej_limpio = ejercicio.ejercicio.split("-")[-1].strip() if "-" in ejercicio.ejercicio else ejercicio.ejercicio
        ejercicio_nombre = f". {ej_limpio}"
    else:
        ejercicio_nombre = ""
    ws['A1'] = f"{empresa_nombre}{ejercicio_nombre}".strip(" .")
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')
    
    ws.merge_cells('A2:F2')
    ws['A2'] = titulo
    ws['A2'].font = Font(size=12, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center')
    
    ws.merge_cells('A3:F3')
    rango = ""
    f_desde_str = fecha_desde.strftime('%d/%m/%Y') if isinstance(fecha_desde, date) else fecha_desde
    f_hasta_str = fecha_hasta.strftime('%d/%m/%Y') if isinstance(fecha_hasta, date) else fecha_hasta
    
    if f_desde_str and f_hasta_str:
        rango = f"Desde {f_desde_str} hasta {f_hasta_str} | "
    elif f_desde_str:
        rango = f"Desde {f_desde_str} | "
    elif f_hasta_str:
        rango = f"Hasta {f_hasta_str} | "
        
    ws['A3'] = f"{rango}Emisión: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    ws['A3'].font = Font(size=10, italic=True)
    ws['A3'].alignment = Alignment(horizontal='right')


def exportar_diario_excel(asientos, empresa):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libro Diario"
    
    _configurar_hoja(ws, "LIBRO DIARIO", empresa)
    
    # Encabezados
    headers = ['Fecha', 'Asiento Nro', 'Concepto', 'Debe', 'Haber']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5") # Indigo 600
        cell.alignment = Alignment(horizontal='center')
    
    row_num = 6
    border_bottom = Border(bottom=Side(style='thin'))
    
    for asiento in asientos:
        # Cabecera de asiento
        ws.cell(row=row_num, column=1).value = asiento.fecha.strftime('%d/%m/%Y')
        ws.cell(row=row_num, column=2).value = asiento.asiento_id
        ws.cell(row=row_num, column=3).value = asiento.concepto
        ws.cell(row=row_num, column=3).font = Font(bold=True)
        row_num += 1
        
        total_debe = 0
        total_haber = 0
        
        # Lineas del asiento
        for linea in asiento.lineas.all().order_by('orden'):
            ws.cell(row=row_num, column=3).value = f"    {linea.cuenta.cuenta}"
            ws.cell(row=row_num, column=4).value = float(linea.debe) if linea.debe else None
            ws.cell(row=row_num, column=5).value = float(linea.haber) if linea.haber else None
            
            # Formato moneda
            ws.cell(row=row_num, column=4).number_format = '#,##0.00'
            ws.cell(row=row_num, column=5).number_format = '#,##0.00'
            
            total_debe += linea.debe
            total_haber += linea.haber
            row_num += 1
            
        # Totales del asiento
        ws.cell(row=row_num, column=3).value = "TOTAL ASIENTO"
        ws.cell(row=row_num, column=3).font = Font(italic=True, bold=True)
        ws.cell(row=row_num, column=3).alignment = Alignment(horizontal='right')
        
        c_debe = ws.cell(row=row_num, column=4)
        c_debe.value = float(total_debe)
        c_debe.number_format = '#,##0.00'
        c_debe.font = Font(bold=True)
        c_debe.border = border_bottom
        
        c_haber = ws.cell(row=row_num, column=5)
        c_haber.value = float(total_haber)
        c_haber.number_format = '#,##0.00'
        c_haber.font = Font(bold=True)
        c_haber.border = border_bottom
        
        row_num += 2 # Espacio entre asientos

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="libro_diario_{timezone.localdate().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


def exportar_mayor_excel(cuentas_data, empresa, ejercicio=None, fecha_desde=None, fecha_hasta=None, columnas_sel=None):
    from contable.services.reportes_mayor import COLUMNAS_MAYOR_CATALOGO, obtener_valor_columna_movimiento

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libro Mayor"
    
    _configurar_hoja(ws, "LIBRO MAYOR", empresa, ejercicio, fecha_desde, fecha_hasta)
    
    if not columnas_sel:
        columnas_sel = [c['clave'] for c in COLUMNAS_MAYOR_CATALOGO if c['default']]
        
    headers = [c['nombre'] for c in COLUMNAS_MAYOR_CATALOGO if c['clave'] in columnas_sel]
    keys = [c['clave'] for c in COLUMNAS_MAYOR_CATALOGO if c['clave'] in columnas_sel]
    
    # Encabezados en fila 5
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal='center')
        
    start_data_row = 6
    row_num = start_data_row
    
    for cta_data in cuentas_data:
        for mov in cta_data['movimientos']:
            for col_num, key in enumerate(keys, 1):
                val = obtener_valor_columna_movimiento(mov, cta_data, key)
                cell = ws.cell(row=row_num, column=col_num)
                if key in ['debe', 'haber', 'saldo', 'monto_asiento', 'cotizacion', 'debe_divisa', 'haber_divisa']:
                    cell.value = float(val) if val else 0.0
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal='right')
                elif key in ['asiento_id', 'numero_diario', 'orden_linea', 'cuenta_id']:
                    cell.value = int(val) if (val != '' and val is not None) else ''
                    cell.alignment = Alignment(horizontal='center')
                else:
                    cell.value = str(val) if val is not None else ''
            row_num += 1
            
    end_data_row = max(row_num - 1, start_data_row)
    
    # Formatear como Tabla de Excel si hay filas
    if end_data_row >= start_data_row:
        max_col_letter = get_column_letter(len(headers))
        tab = Table(displayName="LibroMayorTabla", ref=f"A5:{max_col_letter}{end_data_row}")
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws.add_table(tab)
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="libro_mayor_{timezone.localdate().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


def exportar_balance_excel(balance_data, empresa, fecha_hasta, context_data=None, ejercicio=None, fecha_desde=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sumas y Saldos"
    
    _configurar_hoja(ws, "BALANCE DE COMPROBACIÓN DE SUMAS Y SALDOS", empresa, ejercicio, fecha_desde, fecha_hasta)
    
    headers = ['ID', 'Jerarquía', 'Detalle', 'Apertura', 'Debe', 'Haber', 'Saldo']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal='center')
        
    row_num = 6
    
    blue_font = Font(color="2563EB", bold=True)
    
    for item in balance_data:
        c1 = ws.cell(row=row_num, column=1)
        c1.value = item['cuenta'].id
        
        c2 = ws.cell(row=row_num, column=2)
        c2.value = item['cuenta'].jerarquia
        
        c3 = ws.cell(row=row_num, column=3)
        c3.value = item['cuenta'].cuenta
        
        c4 = ws.cell(row=row_num, column=4)
        c4.value = float(item['apertura'])
        c4.number_format = '#,##0.00'
        
        c5 = ws.cell(row=row_num, column=5)
        c5.value = float(item['periodo_debe'])
        c5.number_format = '#,##0.00'
        
        c6 = ws.cell(row=row_num, column=6)
        c6.value = float(item['periodo_haber'])
        c6.number_format = '#,##0.00'
        
        c7 = ws.cell(row=row_num, column=7)
        c7.value = float(item['saldo'])
        c7.number_format = '#,##0.00'
        
        if item['cuenta'].imputable == 0:
            for cell in [c1, c2, c3, c4, c5, c6, c7]:
                cell.font = blue_font
                
        row_num += 1
        
    if context_data and 'total_apertura' in context_data:
        ws.cell(row=row_num, column=3).value = "TOTALES GENERALES"
        ws.cell(row=row_num, column=3).font = Font(bold=True)
        ws.cell(row=row_num, column=4).value = float(context_data['total_apertura'])
        ws.cell(row=row_num, column=5).value = float(context_data['total_debe'])
        ws.cell(row=row_num, column=6).value = float(context_data['total_haber'])
        ws.cell(row=row_num, column=7).value = float(context_data['total_saldo'])
        for col in range(4, 8):
            cell = ws.cell(row=row_num, column=col)
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)
        row_num += 1
        
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="balance_{timezone.localdate().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response


# =============================================================================
# BALANCE DE SALDOS MENSUALES (Plan 047)
# =============================================================================

_SM_COL_CODIGO = 1      # A
_SM_COL_SUMARIZA = 2    # B
_SM_COL_JERARQUIA = 3   # C
_SM_COL_DETALLE = 4     # D
_SM_COL_IMP = 5         # E
_SM_COL_APERTURA = 6    # F


def exportar_saldos_mensuales_excel(datos, empresa):
    ejercicio = datos.get('ejercicio')
    periodos = datos['periodos']
    invertir = datos.get('alcance') == 'resultados'
    signo = -1 if invertir else 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Saldos Mensuales"

    titulo = "BALANCE DE SALDOS MENSUALES"
    if invertir:
        titulo += " — CUENTAS DE RESULTADO"
    _configurar_hoja(
        ws, titulo, empresa, ejercicio,
        ejercicio.inicio if ejercicio else None,
        ejercicio.cierre if ejercicio else None,
    )

    col_total = _SM_COL_APERTURA + len(periodos) + 1
    col_tipo = col_total + 1

    encabezados = [
        (_SM_COL_CODIGO, "Codigo"), (_SM_COL_SUMARIZA, "Sumariza"),
        (_SM_COL_JERARQUIA, "Jerarquia"), (_SM_COL_DETALLE, "Detalle"),
        (_SM_COL_IMP, "Imp"), (_SM_COL_APERTURA, "Apertura"),
        (col_total, "Total"), (col_tipo, "Tipo"),
        (col_tipo + 1, "Bce"), (col_tipo + 2, "Pres"),
        (col_tipo + 3, "Econ"), (col_tipo + 4, "Fciero"),
    ]
    for i, p in enumerate(periodos):
        encabezados.append((_SM_COL_APERTURA + 1 + i, int(p['clave'])))

    fila_encabezado = 4
    for col, valor in encabezados:
        celda = ws.cell(row=fila_encabezado, column=col)
        celda.value = valor
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="4F81BD")
        celda.alignment = Alignment(horizontal='center')
        if isinstance(valor, int):
            celda.number_format = '000000'

    ultima_col = col_tipo + 4
    row = fila_encabezado + 1

    for f in datos['filas']:
        cuenta = f['cuenta']
        ws.cell(row=row, column=_SM_COL_CODIGO).value = cuenta.codigo or cuenta.id
        ws.cell(row=row, column=_SM_COL_SUMARIZA).value = cuenta.sumariza_id or 0
        ws.cell(row=row, column=_SM_COL_JERARQUIA).value = cuenta.jerarquia
        ws.cell(row=row, column=_SM_COL_DETALLE).value = cuenta.cuenta
        ws.cell(row=row, column=_SM_COL_IMP).value = cuenta.imputable

        ws.cell(row=row, column=_SM_COL_APERTURA).value = float(f['apertura']) * signo
        for i, valor in enumerate(f['meses']):
            ws.cell(row=row, column=_SM_COL_APERTURA + 1 + i).value = float(valor) * signo
        ws.cell(row=row, column=col_total).value = float(f['total']) * signo

        ws.cell(row=row, column=col_tipo).value = cuenta.tipo
        ws.cell(row=row, column=col_tipo + 1).value = cuenta.id_bce or 0
        ws.cell(row=row, column=col_tipo + 2).value = cuenta.id_pre or 0
        ws.cell(row=row, column=col_tipo + 3).value = cuenta.id_ec or 0
        ws.cell(row=row, column=col_tipo + 4).value = cuenta.id_fc or 0

        if cuenta.imputable == 0:
            for col in range(1, ultima_col + 1):
                ws.cell(row=row, column=col).font = Font(bold=True, color="0000FF")
        row += 1

    totales = datos['totales']
    row += 1
    ws.cell(row=row, column=_SM_COL_DETALLE).value = "TOTALES:"
    ws.cell(row=row, column=_SM_COL_APERTURA).value = float(totales['apertura']) * signo
    for i, valor in enumerate(totales['meses']):
        ws.cell(row=row, column=_SM_COL_APERTURA + 1 + i).value = float(valor) * signo
    ws.cell(row=row, column=col_total).value = float(totales['total']) * signo
    for col in range(1, ultima_col + 1):
        celda = ws.cell(row=row, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="4F81BD")

    for fila in ws.iter_rows(min_row=fila_encabezado + 1, max_row=row,
                             min_col=_SM_COL_APERTURA, max_col=col_total):
        for celda in fila:
            celda.number_format = '#,##0.00'

    ws.column_dimensions['A'].width = 9
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 38
    ws.column_dimensions['E'].width = 5
    for i in range(_SM_COL_APERTURA, col_total + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 16
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=_SM_COL_APERTURA)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = (
        f'attachment; filename="saldos_mensuales_'
        f'{timezone.localdate().strftime("%Y%m%d")}.xlsx"')
    wb.save(response)
    return response
