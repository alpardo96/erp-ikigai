"""
Servicio de Exportación e Importación/Captura Masiva en Excel para el Plan de Cuentas.
ERP Ikigai 2.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.db import transaction
from contable.models import Cuenta

# Mapeo de columnas disponibles para exportación y captura de Cuentas Contables
COLUMNAS_CUENTA_MAP = {
    'id': 'ID',
    'sumariza_id': 'Sumariza ID',
    'jerarquia': 'Jerarquía',
    'cuenta': 'Nombre Cuenta',
    'imputable': 'Imputable (1=Sí, 0=No)',
    'tipo': 'Tipo (A/P/N/R)',
    'codigo': 'Código (Legacy)',
    'rg_830': 'RG 830',
    'tipo_disponibilidad': 'Tipo Disponibilidad',
    'id_pre': 'ID Pre',
    'id_bce': 'ID Bce',
    'id_ec': 'ID EC',
    'id_fc': 'ID FC',
}


def generar_excel_cuentas(queryset):
    """
    Genera un libro de Excel (.xlsx) con el plan de cuentas provisto en el queryset.
    """
    columnas_keys = list(COLUMNAS_CUENTA_MAP.keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Plan de Cuentas"

    # Estilos del encabezado (Dark slate slate-900)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    border_light = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(top=border_light, bottom=border_light, left=border_light, right=border_light)

    # Escribir Encabezados
    headers = [COLUMNAS_CUENTA_MAP[k] for k in columnas_keys]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Escribir Filas de Datos
    for cta in queryset:
        row_data = [
            cta.id,
            cta.sumariza_id if cta.sumariza_id is not None else '',
            cta.jerarquia or '',
            (cta.cuenta or '').upper(),
            cta.imputable if cta.imputable is not None else 0,
            cta.tipo or '',
            cta.codigo if cta.codigo is not None else '',
            cta.rg_830 if cta.rg_830 is not None else '',
            cta.tipo_disponibilidad or '',
            cta.id_pre if cta.id_pre is not None else '',
            cta.id_bce if cta.id_bce is not None else '',
            cta.id_ec if cta.id_ec is not None else '',
            cta.id_fc if cta.id_fc is not None else '',
        ]
        ws.append(row_data)

    # Formatear filas y autoajustar anchos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(columnas_keys)):
        for cell in row:
            cell.font = data_font
            cell.border = cell_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb


def procesar_captura_excel_cuentas(empresa, usuario, archivo_excel):
    """
    Procesa la captura e importación masiva de cuentas contables desde un archivo Excel.

    Reglas de negocio:
    1. Convierte a MAYÚSCULAS el nombre de la cuenta (campo `cuenta`).
    2. Si 'ID' viene y EXISTE en la empresa activa -> Actualiza todos los campos excepto el ID.
    3. Si 'ID' no se encuentra en la empresa activa (o está vacío/inválido) -> Trata la fila como una
       cuenta NUEVA y deja que la base de datos (PostgreSQL/Django) asigne automáticamente el ID.
    4. Resuelve dinámicamente el vínculo `sumariza` buscando por `sumariza_id` o la jerarquía padre correspondiente en la empresa.
    """
    resultado = {
        'actualizados': 0,
        'creados': 0,
        'errores': []
    }

    try:
        wb = openpyxl.load_workbook(archivo_excel, data_only=True)
        ws = wb.active
    except Exception as e:
        resultado['errores'].append(f"Error al abrir el archivo Excel: {str(e)}")
        return resultado

    # Leer encabezados (primera fila)
    headers = []
    for cell in ws[1]:
        val = str(cell.value or '').strip()
        headers.append(val)

    if not headers:
        resultado['errores'].append("El archivo Excel está vacío o no contiene encabezados.")
        return resultado

    # Mapeo inverso de encabezados a nombres de campo
    normalized_headers = {}
    for idx, h in enumerate(headers):
        h_lower = h.lower()
        found_key = None
        for k, v in COLUMNAS_CUENTA_MAP.items():
            if k.lower() == h_lower or v.lower() == h_lower:
                found_key = k
                break
        if found_key:
            normalized_headers[idx] = found_key
        else:
            # Coincidencias flexibles por nombres parciales
            if 'id' in h_lower and 'sumariza' not in h_lower and 'pre' not in h_lower and 'bce' not in h_lower and 'ec' not in h_lower and 'fc' not in h_lower:
                normalized_headers[idx] = 'id'
            elif 'sumariza' in h_lower:
                normalized_headers[idx] = 'sumariza_id'
            elif 'jerarq' in h_lower or 'codigo jer' in h_lower:
                normalized_headers[idx] = 'jerarquia'
            elif 'nombre' in h_lower or 'cuenta' in h_lower or 'descripcion' in h_lower:
                normalized_headers[idx] = 'cuenta'
            elif 'imput' in h_lower:
                normalized_headers[idx] = 'imputable'
            elif 'tipo' in h_lower and 'disponib' not in h_lower:
                normalized_headers[idx] = 'tipo'
            elif 'legacy' in h_lower or 'cod' in h_lower:
                normalized_headers[idx] = 'codigo'
            elif 'rg' in h_lower or '830' in h_lower:
                normalized_headers[idx] = 'rg_830'
            elif 'disponib' in h_lower:
                normalized_headers[idx] = 'tipo_disponibilidad'
            elif 'id_pre' in h_lower or 'pre' in h_lower:
                normalized_headers[idx] = 'id_pre'
            elif 'id_bce' in h_lower or 'bce' in h_lower:
                normalized_headers[idx] = 'id_bce'
            elif 'id_ec' in h_lower or 'ec' in h_lower:
                normalized_headers[idx] = 'id_ec'
            elif 'id_fc' in h_lower or 'fc' in h_lower:
                normalized_headers[idx] = 'id_fc'

    # Cache preexistente de cuentas de la empresa por ID y por Jerarquía
    cuentas_por_id = {c.id: c for c in Cuenta.objects.filter(empresa=empresa)}
    cuentas_por_jerarquia = {c.jerarquia: c for c in Cuenta.objects.filter(empresa=empresa)}

    with transaction.atomic():
        row_idx = 1
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_idx += 1
            if not any(row):
                continue  # Fila totalmente vacía

            row_dict = {}
            for col_idx, val in enumerate(row):
                if col_idx in normalized_headers:
                    row_dict[normalized_headers[col_idx]] = val

            jerarquia_val = str(row_dict.get('jerarquia') or '').strip()
            cuenta_val = str(row_dict.get('cuenta') or '').strip().upper()

            if not jerarquia_val or not cuenta_val:
                resultado['errores'].append(f"Fila {row_idx}: Omitida por no contar con 'Jerarquía' o 'Nombre Cuenta' válido.")
                continue

            # Parsear Imputable (1 o 0)
            raw_imp = row_dict.get('imputable')
            if str(raw_imp).strip().upper() in ['1', 'SI', 'SÍ', 'TRUE', 'S']:
                imputable_val = 1
            else:
                try:
                    imputable_val = int(float(str(raw_imp)))
                    imputable_val = 1 if imputable_val != 0 else 0
                except (ValueError, TypeError):
                    imputable_val = 0

            # Parsear Tipo (A, P, N, R)
            tipo_raw = str(row_dict.get('tipo') or '').strip().upper()
            if tipo_raw.startswith('A'): tipo_val = 'A'
            elif tipo_raw.startswith('P'): tipo_val = 'P'
            elif tipo_raw.startswith('N'): tipo_val = 'N'
            elif tipo_raw.startswith('R'): tipo_val = 'R'
            else: tipo_val = 'A'  # Valor por defecto seguro

            # Helper para enteros opcionales
            def _clean_int(v):
                if v is None or str(v).strip() == '':
                    return None
                try:
                    return int(float(str(v).strip()))
                except (ValueError, TypeError):
                    return None

            codigo_val = _clean_int(row_dict.get('codigo'))
            sumariza_id_val = _clean_int(row_dict.get('sumariza_id'))
            rg_830_val = _clean_int(row_dict.get('rg_830'))
            id_pre_val = _clean_int(row_dict.get('id_pre'))
            id_bce_val = _clean_int(row_dict.get('id_bce'))
            id_ec_val = _clean_int(row_dict.get('id_ec'))
            id_fc_val = _clean_int(row_dict.get('id_fc'))

            # Parsear tipo_disponibilidad
            disp_raw = str(row_dict.get('tipo_disponibilidad') or '').strip().upper()
            valid_disps = ['EFE', 'DOL', 'VAL', 'BCO', 'TAR', 'OTR']
            tipo_disp_val = disp_raw if disp_raw in valid_disps else ''

            # Buscar ID si viene especificado
            raw_id = row_dict.get('id')
            cta_id = _clean_int(raw_id)

            # Buscar si la cuenta ya existe en la empresa activa por ID
            cuenta_obj = None
            if cta_id and cta_id in cuentas_por_id:
                cuenta_obj = cuentas_por_id[cta_id]

            # Buscar cuenta padre (sumariza) por ID explícito o por prefijo de jerarquía
            sumariza_obj = None
            if sumariza_id_val and sumariza_id_val in cuentas_por_id:
                sumariza_obj = cuentas_por_id[sumariza_id_val]
            elif '.' in jerarquia_val:
                padre_jerarquia = jerarquia_val.rsplit('.', 1)[0]
                if padre_jerarquia in cuentas_por_jerarquia:
                    sumariza_obj = cuentas_por_jerarquia[padre_jerarquia]

            if cuenta_obj:
                # --- ACTUALIZACIÓN DE CUENTA EXISTENTE ---
                cuenta_obj.jerarquia = jerarquia_val
                cuenta_obj.cuenta = cuenta_val
                cuenta_obj.imputable = imputable_val
                cuenta_obj.tipo = tipo_val
                cuenta_obj.codigo = codigo_val
                cuenta_obj.rg_830 = rg_830_val
                cuenta_obj.tipo_disponibilidad = tipo_disp_val
                cuenta_obj.id_pre = id_pre_val
                cuenta_obj.id_bce = id_bce_val
                cuenta_obj.id_ec = id_ec_val
                cuenta_obj.id_fc = id_fc_val
                cuenta_obj.sumariza = sumariza_obj
                if usuario:
                    cuenta_obj.modificado_por = usuario
                cuenta_obj.save()

                resultado['actualizados'] += 1
                cuentas_por_jerarquia[jerarquia_val] = cuenta_obj
            else:
                # --- CREACIÓN DE NUEVA CUENTA ---
                # Si el usuario puso un ID que no existía en la empresa activa, NO se fuerza ese ID,
                # sino que PostgreSQL le asignará el ID correspondiente de forma automática.
                nueva_cta = Cuenta.objects.create(
                    empresa=empresa,
                    jerarquia=jerarquia_val,
                    cuenta=cuenta_val,
                    imputable=imputable_val,
                    tipo=tipo_val,
                    codigo=codigo_val,
                    rg_830=rg_830_val,
                    tipo_disponibilidad=tipo_disp_val,
                    id_pre=id_pre_val,
                    id_bce=id_bce_val,
                    id_ec=id_ec_val,
                    id_fc=id_fc_val,
                    sumariza=sumariza_obj,
                    creado_por=usuario if usuario else None,
                    modificado_por=usuario if usuario else None,
                )
                resultado['creados'] += 1
                cuentas_por_id[nueva_cta.id] = nueva_cta
                cuentas_por_jerarquia[jerarquia_val] = nueva_cta

    return resultado
