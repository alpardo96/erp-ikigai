import io
import openpyxl
from decimal import Decimal
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth.models import User
from empresas.models import Empresa
from productos.models import Producto, Marca, Rubro, Familia
from productos.services.excel_service import (
    generar_excel_productos,
    procesar_captura_excel_productos
)


class ExcelProductosTestCase(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = User.objects.create_user(username='testuser', password='password123')
        self.empresa = Empresa.objects.create(
            nombre='EMPRESA TEST SA',
            cuit='30711111118'
        )

        # Crear marca, rubro y producto inicial
        self.marca = Marca.objects.create(empresa=self.empresa, detalle='SHELL')
        self.rubro = Rubro.objects.create(empresa=self.empresa, detalle='LUBRICANTES')
        self.familia = Familia.objects.create(empresa=self.empresa, rubro=self.rubro, detalle='ACEITES SINTETICOS')
        
        self.producto = Producto.objects.create(
            empresa=self.empresa,
            cod_prov='ACE-500',
            cod_fab='SH-500',
            detalle='ACEITE SHELL HELIX 5W30',
            marca=self.marca,
            rubro=self.rubro,
            familia=self.familia,
            cto_adq=Decimal('1000.00'),
            cto_rep=Decimal('1200.00'),
            margen=Decimal('30.00'),
            precio_neto=Decimal('1560.00'),
            precio_total=Decimal('1887.60'),
            alic_iva=Decimal('21.00')
        )

        # Iniciar sesión
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()
        self.client.force_login(self.user)

    def test_exportar_excel_completo(self):
        url = reverse('producto_exportar_excel_completo')
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        # Cargar workbook desde los datos de la respuesta
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, 'ID')
        self.assertEqual(ws.cell(row=2, column=1).value, self.producto.id)
        self.assertEqual(ws.cell(row=2, column=4).value, 'ACEITE SHELL HELIX 5W30')

    def test_exportar_excel_seleccion(self):
        url = reverse('producto_exportar_excel_seleccion')
        response = self.client.post(url, {'columnas': ['id', 'detalle', 'precio_total']})
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.max_column, 3)
        self.assertEqual(ws.cell(row=1, column=1).value, 'ID')
        self.assertEqual(ws.cell(row=1, column=2).value, 'Detalle')
        self.assertEqual(ws.cell(row=1, column=3).value, 'Precio Final')

    def test_capturar_excel_actualizar_y_crear(self):
        # Crear un archivo Excel en memoria con 2 filas:
        # Fila 1: Producto existente (con ID) -> actualizar costo y detalle
        # Fila 2: Producto nuevo (sin ID, con Marca/Rubro nuevos en minúsculas) -> debe crearlos en MAYÚSCULAS
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['ID', 'Cód. Prov', 'Cód. Fab', 'Detalle', 'Marca', 'Rubro', 'Familia', 'Costo Rep.', 'Margen', 'IVA (%)'])

        # Fila 1: Actualizar producto existente #self.producto.id
        ws.append([self.producto.id, 'ACE-500-MOD', 'SH-500-MOD', 'aceite shell helix 5w30 modificado', 'SHELL', 'LUBRICANTES', 'ACEITES SINTETICOS', 1500.0, 35.0, 21.0])

        # Fila 2: Crear producto nuevo (sin ID)
        ws.append(['', 'FIL-100', 'MN-100', 'filtro de aceite mann', 'Mann Filter', 'Filtros', 'Filtros de Aceite', 800.0, 40.0, 21.0])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        resultado = procesar_captura_excel_productos(self.empresa, self.user, buffer)

        self.assertEqual(resultado['actualizados'], 1)
        self.assertEqual(resultado['creados'], 1)
        self.assertGreaterEqual(resultado['entidades_creadas'], 2)

        # Verificar actualización de Fila 1
        prod_mod = Producto.objects.get(id=self.producto.id)
        self.assertEqual(prod_mod.detalle, 'ACEITE SHELL HELIX 5W30 MODIFICADO') # Convertido a MAYÚSCULAS
        self.assertEqual(prod_mod.cod_prov, 'ACE-500-MOD')
        self.assertEqual(prod_mod.cto_rep, Decimal('1500.00'))

        # Verificar creación de Fila 2 (Producto Nuevo)
        prod_nuevo = Producto.objects.get(cod_prov='FIL-100')
        self.assertIsNotNone(prod_nuevo.id)
        self.assertNotEqual(prod_nuevo.id, self.producto.id)
        self.assertEqual(prod_nuevo.detalle, 'FILTRO DE ACEITE MANN')
        self.assertEqual(prod_nuevo.marca.detalle, 'MANN FILTER') # Creada automáticamente en MAYÚSCULAS
        self.assertEqual(prod_nuevo.rubro.detalle, 'FILTROS') # Creado automáticamente en MAYÚSCULAS
        self.assertEqual(prod_nuevo.familia.detalle, 'FILTROS DE ACEITE')
