"""
Servicio de Exportación e Importación/Captura Masiva en Excel para Productos.
ERP Ikigai 2.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from decimal import Decimal, InvalidOperation
from django.db import transaction
from django.utils.timezone import now
from productos.models import Producto, Marca, Rubro, Familia
from facturacion.models import ClienteProveedor

# Mapeo de columnas disponibles para exportación y captura
COLUMNAS_PRODUCTO_MAP = {
    'id': 'ID',
    'cod_prov': 'Cód. Prov',
    'cod_fab': 'Cód. Fab',
    'detalle': 'Detalle',
    'proveedor': 'Proveedor',
    'minimo': 'Mínimo',
    'ptopedir': 'Pto. Pedir',
    'creden': 'Credencial',
    'moneda': 'Moneda',
    'alic_iva': 'IVA (%)',
    'margen': 'Margen (%)',
    'marca': 'Marca',
    'rubro': 'Rubro',
    'familia': 'Familia',
    'subprod': 'Subproductos',
    'stock_global': 'Stock Total',
    'cto_adq': 'Costo Adq.',
    'cto_rep': 'Costo Rep.',
    'precio_neto': 'Precio Neto',
    'precio_total': 'Precio Final',
    'cotiz_cpra': 'Cotiz. Compra',
    'fec_adq': 'Fec. Adq.',
    'fec_act': 'Fec. Act.',
}

def generar_excel_productos(queryset, columnas_seleccionadas=None):
    """
    Genera un libro de Excel (.xlsx) con los productos provistos en el queryset
    y las columnas solicitadas (o todas si no se especifica selección).
    """
    if not columnas_seleccionadas:
        columnas_seleccionadas = list(COLUMNAS_PRODUCTO_MAP.keys())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    # Estilos del encabezado
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark slate
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    border_light = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(top=border_light, bottom=border_light, left=border_light, right=border_light)

    # Escribir Encabezados
    headers = [COLUMNAS_PRODUCTO_MAP.get(col, col) for col in columnas_seleccionadas]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Escribir Filas de Datos
    for prod in queryset:
        row_data = []
        for col_key in columnas_seleccionadas:
            if col_key == 'id':
                row_data.append(prod.id)
            elif col_key == 'cod_prov':
                row_data.append(prod.cod_prov or '')
            elif col_key == 'cod_fab':
                row_data.append(prod.cod_fab or '')
            elif col_key == 'detalle':
                row_data.append((prod.detalle or '').upper())
            elif col_key == 'proveedor':
                row_data.append(prod.proveedor.razon_social if prod.proveedor else '')
            elif col_key == 'minimo':
                row_data.append(float(prod.minimo or 0))
            elif col_key == 'ptopedir':
                row_data.append(float(prod.ptopedir or 0))
            elif col_key == 'creden':
                row_data.append("SI" if prod.creden else "NO")
            elif col_key == 'moneda':
                row_data.append(prod.get_moneda_display())
            elif col_key == 'alic_iva':
                row_data.append(float(prod.alic_iva_porc or 21.0))
            elif col_key == 'margen':
                row_data.append(float(prod.margen or 0))
            elif col_key == 'marca':
                row_data.append((prod.marca.detalle if prod.marca else '').upper())
            elif col_key == 'rubro':
                row_data.append((prod.rubro.detalle if prod.rubro else '').upper())
            elif col_key == 'familia':
                row_data.append((prod.familia.detalle if prod.familia else '').upper())
            elif col_key == 'subprod':
                row_data.append("SI" if prod.subprod else "NO")
            elif col_key == 'stock_global':
                row_data.append(float(prod.stock_global or 0))
            elif col_key == 'cto_adq':
                row_data.append(float(prod.cto_adq or 0))
            elif col_key == 'cto_rep':
                row_data.append(float(prod.cto_rep or 0))
            elif col_key == 'precio_neto':
                row_data.append(float(prod.precio_neto or 0))
            elif col_key == 'precio_total':
                row_data.append(float(prod.precio_total or 0))
            elif col_key == 'cotiz_cpra':
                row_data.append(float(prod.cotiz_cpra or 0))
            elif col_key == 'fec_adq':
                row_data.append(prod.fec_adq.strftime("%d/%m/%Y") if prod.fec_adq else '')
            elif col_key == 'fec_act':
                row_data.append(prod.fec_act.strftime("%d/%m/%Y") if prod.fec_act else '')
            else:
                row_data.append('')
        
        ws.append(row_data)

    # Formatear filas y autoajustar anchos
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(columnas_seleccionadas)):
        for cell in row:
            cell.font = data_font
            cell.border = cell_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb


def procesar_captura_excel_productos(empresa, usuario, archivo_excel):
    """
    Procesa la captura e importación masiva de productos desde un archivo Excel.
    
    Reglas:
    1. Convierte a MAYÚSCULAS todos los campos de texto (detalle, cod_prov, cod_fab, marca, rubro, familia).
    2. Si 'ID' ya existe en la empresa, actualiza todos los campos excepto el ID.
    3. Si 'ID' no se encuentra o está vacío, crea un nuevo producto dejando que la BD asigne el ID automáticamente.
    4. Si 'Marca', 'Rubro' o 'Familia' no existen en la BD de la empresa, los crea automáticamente en MAYÚSCULAS.
    """
    resultado = {
        'actualizados': 0,
        'creados': 0,
        'entidades_creadas': 0,
        'errores': []
    }

    try:
        wb = openpyxl.load_workbook(archivo_excel, data_only=True)
        ws = wb.active
    except Exception as e:
        resultado['errores'].append(f"Error al abrir el archivo Excel: {str(e)}")
        return resultado

    # Leer encabezados (primera fila)
    headers = []
    for cell in ws[1]:
        val = str(cell.value or '').strip()
        headers.append(val)

    if not headers:
        resultado['errores'].append("El archivo Excel está vacío o no contiene encabezados.")
        return resultado

    # Mapeo inverso de encabezados a nombres de campo
    normalized_headers = {}
    for idx, h in enumerate(headers):
        h_lower = h.lower()
        # Buscar coincidencia en COLUMNAS_PRODUCTO_MAP
        found_key = None
        for k, v in COLUMNAS_PRODUCTO_MAP.items():
            if k.lower() == h_lower or v.lower() == h_lower:
                found_key = k
                break
        if found_key:
            normalized_headers[idx] = found_key
        else:
            # Coincidencias adicionales por nombres comunes
            if 'id' in h_lower: normalized_headers[idx] = 'id'
            elif 'prov' in h_lower and 'cod' in h_lower: normalized_headers[idx] = 'cod_prov'
            elif 'fab' in h_lower and 'cod' in h_lower: normalized_headers[idx] = 'cod_fab'
            elif 'det' in h_lower or 'desc' in h_lower or 'nombre' in h_lower: normalized_headers[idx] = 'detalle'
            elif 'proveed' in h_lower: normalized_headers[idx] = 'proveedor'
            elif 'min' in h_lower: normalized_headers[idx] = 'minimo'
            elif 'pedir' in h_lower or 'ptoped' in h_lower: normalized_headers[idx] = 'ptopedir'
            elif 'cred' in h_lower: normalized_headers[idx] = 'creden'
            elif 'mon' in h_lower: normalized_headers[idx] = 'moneda'
            elif 'iva' in h_lower: normalized_headers[idx] = 'alic_iva'
            elif 'marg' in h_lower: normalized_headers[idx] = 'margen'
            elif 'marca' in h_lower: normalized_headers[idx] = 'marca'
            elif 'rubro' in h_lower: normalized_headers[idx] = 'rubro'
            elif 'familia' in h_lower: normalized_headers[idx] = 'familia'
            elif 'subp' in h_lower or 'serie' in h_lower: normalized_headers[idx] = 'subprod'
            elif 'adq' in h_lower or 'costo ad' in h_lower: normalized_headers[idx] = 'cto_adq'
            elif 'rep' in h_lower or 'costo re' in h_lower: normalized_headers[idx] = 'cto_rep'
            elif 'neto' in h_lower: normalized_headers[idx] = 'precio_neto'
            elif 'precio' in h_lower or 'total' in h_lower or 'final' in h_lower: normalized_headers[idx] = 'precio_total'

    # Cache de relaciones para evitar queries redundantes
    marcas_cache = {m.detalle.upper(): m for m in Marca.objects.filter(empresa=empresa)}
    rubros_cache = {r.detalle.upper(): r for r in Rubro.objects.filter(empresa=empresa)}
    familias_cache = {f.detalle.upper(): f for f in Familia.objects.filter(empresa=empresa)}
    proveedores_cache = {p.razon_social.upper(): p for p in ClienteProveedor.objects.filter(empresa=empresa, tipo_entidad=2)}

    with transaction.atomic():
        row_idx = 1
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_idx += 1
            if not any(row):
                continue  # Fila vacía

            # Extraer diccionario de valores de la fila
            row_dict = {}
            for col_idx, val in enumerate(row):
                if col_idx in normalized_headers:
                    row_dict[normalized_headers[col_idx]] = val

            detalle_val = str(row_dict.get('detalle') or '').strip().upper()
            if not detalle_val:
                resultado['errores'].append(f"Fila {row_idx}: Omitida por no tener 'Detalle' del producto.")
                continue

            # 1. Determinar ID si viene
            prod_id = None
            raw_id = row_dict.get('id')
            if raw_id is not None and str(raw_id).strip() != '':
                try:
                    prod_id = int(float(str(raw_id).strip()))
                except (ValueError, TypeError):
                    prod_id = None

            # 2. Procesar relaciones (Marca, Rubro, Familia, Proveedor)
            # Marca
            marca_obj = None
            marca_raw = str(row_dict.get('marca') or '').strip().upper()
            if marca_raw:
                if marca_raw in marcas_cache:
                    marca_obj = marcas_cache[marca_raw]
                else:
                    marca_obj = Marca.objects.create(
                        empresa=empresa,
                        detalle=marca_raw,
                        creado_por=usuario,
                        modificado_por=usuario
                    )
                    marcas_cache[marca_raw] = marca_obj
                    resultado['entidades_creadas'] += 1

            # Rubro
            rubro_obj = None
            rubro_raw = str(row_dict.get('rubro') or '').strip().upper()
            if rubro_raw:
                if rubro_raw in rubros_cache:
                    rubro_obj = rubros_cache[rubro_raw]
                else:
                    rubro_obj = Rubro.objects.create(
                        empresa=empresa,
                        detalle=rubro_raw,
                        creado_por=usuario,
                        modificado_por=usuario
                    )
                    rubros_cache[rubro_raw] = rubro_obj
                    resultado['entidades_creadas'] += 1

            # Familia
            familia_obj = None
            familia_raw = str(row_dict.get('familia') or '').strip().upper()
            if familia_raw:
                if familia_raw in familias_cache:
                    familia_obj = familias_cache[familia_raw]
                else:
                    familia_obj = Familia.objects.create(
                        empresa=empresa,
                        rubro=rubro_obj,
                        detalle=familia_raw,
                        creado_por=usuario,
                        modificado_por=usuario
                    )
                    familias_cache[familia_raw] = familia_obj
                    resultado['entidades_creadas'] += 1

            # Proveedor
            proveedor_obj = None
            prov_raw = str(row_dict.get('proveedor') or '').strip().upper()
            if prov_raw:
                if prov_raw in proveedores_cache:
                    proveedor_obj = proveedores_cache[prov_raw]
                elif prov_raw.isdigit():
                    prov_by_id = ClienteProveedor.objects.filter(empresa=empresa, id=int(prov_raw)).first()
                    if prov_by_id:
                        proveedor_obj = prov_by_id

            # 3. Parsear valores numéricos y booleanos
            def to_decimal(v, default=Decimal('0.00')):
                if v is None or str(v).strip() == '':
                    return default
                try:
                    # reemplazar coma por punto
                    clean_str = str(v).strip().replace(',', '.')
                    return Decimal(clean_str)
                except (InvalidOperation, ValueError, TypeError):
                    return default

            def to_bool(v):
                if isinstance(v, bool):
                    return v
                s = str(v or '').strip().upper()
                return s in ['SI', 'SÍ', 'TRUE', '1', 'VERDADERO', 'X']

            def to_moneda(v):
                s = str(v or '').strip().upper()
                if 'DOL' in s or 'U$S' in s or '$US' in s:
                    return 'DOL'
                if 'EUR' in s or '60' in s:
                    return '60'
                return 'PES'

            def to_iva(v):
                d = to_decimal(v, Decimal('21.00'))
                if d in [Decimal('21.00'), Decimal('10.50'), Decimal('0.00'), Decimal('27.00'), Decimal('5.00'), Decimal('2.50')]:
                    return d
                return Decimal('21.00')

            cod_prov = str(row_dict.get('cod_prov') or '').strip().upper() or None
            cod_fab = str(row_dict.get('cod_fab') or '').strip().upper() or None
            minimo = to_decimal(row_dict.get('minimo'))
            ptopedir = to_decimal(row_dict.get('ptopedir'))
            creden = to_bool(row_dict.get('creden'))
            subprod = to_bool(row_dict.get('subprod'))
            moneda = to_moneda(row_dict.get('moneda'))
            alic_iva = to_iva(row_dict.get('alic_iva'))
            margen = to_decimal(row_dict.get('margen'))
            cto_adq = to_decimal(row_dict.get('cto_adq'))
            cto_rep = to_decimal(row_dict.get('cto_rep'))
            precio_neto = to_decimal(row_dict.get('precio_neto'))
            precio_total = to_decimal(row_dict.get('precio_total'))
            cotiz_cpra = to_decimal(row_dict.get('cotiz_cpra'))

            # Si precio_total no viene especificado pero viene cto_adq/rep y margen, calcularlo
            if precio_total == Decimal('0.00') and (cto_adq > 0 or cto_rep > 0):
                costo_base = cto_rep if cto_rep > 0 else cto_adq
                precio_neto_calc = costo_base * (Decimal('1.00') + (margen / Decimal('100.00')))
                precio_total_calc = precio_neto_calc * (Decimal('1.00') + (alic_iva / Decimal('100.00')))
                if precio_neto == Decimal('0.00'):
                    precio_neto = round(precio_neto_calc, 2)
                precio_total = round(precio_total_calc, 2)

            # 4. Actualizar o Crear Producto
            producto_existente = None
            if prod_id:
                producto_existente = Producto.objects.filter(id=prod_id, empresa=empresa).first()

            if producto_existente:
                # Corregir / Actualizar todos los campos EXCEPTO el ID
                producto_existente.detalle = detalle_val
                if cod_prov is not None: producto_existente.cod_prov = cod_prov
                if cod_fab is not None: producto_existente.cod_fab = cod_fab
                if proveedor_obj: producto_existente.proveedor = proveedor_obj
                producto_existente.minimo = minimo
                producto_existente.ptopedir = ptopedir
                producto_existente.creden = creden
                producto_existente.subprod = subprod
                producto_existente.moneda = moneda
                producto_existente.alic_iva = alic_iva
                producto_existente.margen = margen
                if marca_obj: producto_existente.marca = marca_obj
                if rubro_obj: producto_existente.rubro = rubro_obj
                if familia_obj: producto_existente.familia = familia_obj
                if cto_adq > 0: producto_existente.cto_adq = cto_adq
                if cto_rep > 0: producto_existente.cto_rep = cto_rep
                if precio_neto > 0: producto_existente.precio_neto = precio_neto
                if precio_total > 0: producto_existente.precio_total = precio_total
                if cotiz_cpra > 0: producto_existente.cotiz_cpra = cotiz_cpra
                producto_existente.fec_act = now().date()
                producto_existente.modificado_por = usuario
                producto_existente.save()
                resultado['actualizados'] += 1
            else:
                # Crear Producto Nuevo dejando que la máquina asigne el ID
                nuevo_prod = Producto.objects.create(
                    empresa=empresa,
                    detalle=detalle_val,
                    cod_prov=cod_prov,
                    cod_fab=cod_fab,
                    proveedor=proveedor_obj,
                    minimo=minimo,
                    ptopedir=ptopedir,
                    creden=creden,
                    subprod=subprod,
                    moneda=moneda,
                    alic_iva=alic_iva,
                    margen=margen,
                    marca=marca_obj,
                    rubro=rubro_obj,
                    familia=familia_obj,
                    cto_adq=cto_adq,
                    cto_rep=cto_rep,
                    precio_neto=precio_neto,
                    precio_total=precio_total,
                    cotiz_cpra=cotiz_cpra,
                    fec_act=now().date(),
                    creado_por=usuario,
                    modificado_por=usuario
                )
                resultado['creados'] += 1

    return resultado
