"""Pruebas del listado de Facturas Pendientes (Plan 056).

Cubren lo que hace distinto a este listado: el aislamiento multiempresa, la traducción de
los rangos de `saldo` del VFP (Pagadas / Pendientes, negativos incluidos), el filtro de
condición, la regla `pagado = total - saldo` en ambas operaciones, los totales calculados
sobre el conjunto completo aunque la grilla se trunque, y que el módulo NO escriba nada.
"""
import io
from datetime import date
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse

from empresas.models import Ejercicio, Empresa, Sucursal
from facturacion.models import ClienteProveedor, Compra, TipoComprobante, Venta
from facturacion.services.facturas_pendientes import (
    LIMITE_GRILLA, FiltroFacturas, consultar,
)
from tesoreria.models import OrdenPago, OrdenPagoAplicacion

User = get_user_model()

DESDE = date(2026, 1, 1)
HASTA = date(2026, 12, 31)


class BaseFacturasPendientes(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="EMPRESA FP", cuit="30777777771")
        self.sucursal = Sucursal.objects.create(empresa=self.empresa, nombre="CASA CENTRAL")
        self.ejercicio = Ejercicio.objects.create(
            empresa=self.empresa, ejercicio="FP 2026", inicio=DESDE, cierre=HASTA)
        self.usuario = User.objects.create_user(username="fpuser", password="password123")
        self.tipo = TipoComprobante.objects.create(codigo="FA", detalle="FACTURA A")

        self.proveedor = ClienteProveedor.objects.create(
            empresa=self.empresa, razon_social="PROVEEDOR UNO SA", tipo_entidad=2)
        self.cliente = ClienteProveedor.objects.create(
            empresa=self.empresa, razon_social="CLIENTE UNO SA", tipo_entidad=1)

        self.client = Client()
        self.client.login(username="fpuser", password="password123")
        sesion = self.client.session
        sesion['empresa_id'] = self.empresa.id
        sesion['ejercicio_id'] = self.ejercicio.id
        sesion.save()

    def crear_compra(self, numero, total, saldo, *, condic=1, empresa=None,
                     proveedor=None, fecha=date(2026, 6, 1)):
        compra = Compra.objects.create(
            empresa=empresa or self.empresa,
            sucursal=self.sucursal,
            ejercicio=self.ejercicio,
            usuario=self.usuario,
            proveedor=proveedor or self.proveedor,
            tipo=self.tipo,
            fecha=fecha,
            periodo="202606",
            punto=1,
            numero=numero,
            condic=condic,
            neto=Decimal(str(total)),
            total=Decimal(str(total)),
        )
        # Se fija el saldo por UPDATE para no volver a disparar los signals de contabilización.
        Compra.objects.filter(pk=compra.pk).update(
            saldo=Decimal(str(saldo)), pagado=Decimal(str(total)) - Decimal(str(saldo)))
        compra.refresh_from_db()
        return compra

    def crear_venta(self, numero, total, saldo, *, condic=1, estado=0,
                    fecha=date(2026, 6, 1)):
        # `Venta.save()` recalcula `total` desde los importes: se carga por `neto`.
        venta = Venta.objects.create(
            empresa=self.empresa,
            sucursal=self.sucursal,
            ejercicio=self.ejercicio,
            usuario=self.usuario,
            cliente=self.cliente,
            tipo=self.tipo,
            fecha=fecha,
            periodo="202606",
            punto=1,
            numero=numero,
            condic=condic,
            estado=estado,
            neto=Decimal(str(total)),
        )
        Venta.objects.filter(pk=venta.pk).update(saldo=Decimal(str(saldo)))
        venta.refresh_from_db()
        return venta

    def filtro(self, **kwargs):
        base = dict(empresa_id=self.empresa.id, desde=DESDE, hasta=HASTA, operacion='C')
        base.update(kwargs)
        return FiltroFacturas(**base)


class AislamientoMultiempresaTests(BaseFacturasPendientes):
    def test_no_devuelve_comprobantes_de_otra_empresa(self):
        otra = Empresa.objects.create(nombre="EMPRESA AJENA", cuit="30777777772")
        proveedor_ajeno = ClienteProveedor.objects.create(
            empresa=otra, razon_social="PROVEEDOR AJENO SA", tipo_entidad=2)

        self.crear_compra(1, 1000, 1000)
        self.crear_compra(2, 5000, 5000, empresa=otra, proveedor=proveedor_ajeno)

        filas, totales, _ = consultar(self.filtro())

        self.assertEqual(len(filas), 1)
        self.assertEqual(totales.cantidad, 1)
        self.assertEqual(totales.total, Decimal('1000.00'))

    def test_entidad_de_otra_empresa_se_ignora_en_el_filtro(self):
        """Pasar por querystring un id ajeno no debe filtrar por él ni filtrar por nada."""
        otra = Empresa.objects.create(nombre="EMPRESA AJENA", cuit="30777777773")
        ajeno = ClienteProveedor.objects.create(
            empresa=otra, razon_social="AJENO SA", tipo_entidad=2)
        self.crear_compra(1, 1000, 1000)

        respuesta = self.client.get(
            reverse('facturas_pendientes_grilla'),
            {'desde': DESDE, 'hasta': HASTA, 'operacion': 'C', 'entidad': ajeno.pk})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['totales'].cantidad, 1)


class EstadoDePagoTests(BaseFacturasPendientes):
    def setUp(self):
        super().setUp()
        self.cancelada = self.crear_compra(1, 1000, 0)
        self.pendiente = self.crear_compra(2, 2000, 500)
        # Nota de crédito sin aplicar: saldo NEGATIVO. El VFP la contaba como pendiente
        # (rango `-1e11 .. -1e-10`) y acá tiene que pasar lo mismo.
        self.nota_credito = self.crear_compra(3, -300, -300)

    def test_pagadas_solo_saldo_cero(self):
        filas, totales, _ = consultar(self.filtro(estado='pagadas'))
        self.assertEqual(totales.cantidad, 1)
        self.assertEqual([f.pk for f in filas], [self.cancelada.pk])

    def test_pendientes_incluye_saldos_negativos(self):
        filas, totales, _ = consultar(self.filtro(estado='pendientes'))
        self.assertEqual(totales.cantidad, 2)
        self.assertEqual(
            sorted(f.pk for f in filas),
            sorted([self.pendiente.pk, self.nota_credito.pk]))

    def test_todas_no_filtra(self):
        _, totales, _ = consultar(self.filtro(estado='todas'))
        self.assertEqual(totales.cantidad, 3)


class FiltroCondicionTests(BaseFacturasPendientes):
    def setUp(self):
        super().setUp()
        for i, condic in enumerate((1, 2, 3, 4), start=1):
            self.crear_compra(i, 100 * i, 100 * i, condic=condic)

    def test_por_defecto_trae_las_cuatro_condiciones(self):
        _, totales, _ = consultar(self.filtro())
        self.assertEqual(totales.cantidad, 4)

    def test_filtra_por_condicion(self):
        filas, totales, _ = consultar(self.filtro(condics=(2,)))
        self.assertEqual(totales.cantidad, 1)
        self.assertEqual(filas[0].condic, 2)
        self.assertEqual(filas[0].condic_nombre, 'Presupuestado')

    def test_filtra_por_varias_condiciones(self):
        _, totales, _ = consultar(self.filtro(condics=(1, 3)))
        self.assertEqual(totales.cantidad, 2)

    def test_sin_checkboxes_marcados_trae_todo(self):
        """Un filtro vacío significa "todas", no "ninguna": un listado en blanco confunde."""
        respuesta = self.client.get(
            reverse('facturas_pendientes_grilla'),
            {'desde': DESDE, 'hasta': HASTA, 'operacion': 'C'})
        self.assertEqual(respuesta.context['totales'].cantidad, 4)


class PagadoDerivadoTests(BaseFacturasPendientes):
    """`pagado = total - saldo` en las dos operaciones (Plan 056 §2.3)."""

    def test_compra_pagado_coincide_con_las_aplicaciones_de_op(self):
        from contable.services.saldos import recalcular_saldo_compra

        compra = self.crear_compra(1, 1000, 1000)
        orden = OrdenPago.objects.create(
            empresa=self.empresa, sucursal=self.sucursal, ejercicio=self.ejercicio,
            proveedor=self.proveedor, fecha=date(2026, 6, 15), total=Decimal('400.00'))
        OrdenPagoAplicacion.objects.create(
            orden_pago=orden, compra=compra,
            importe=Decimal('400.00'), importe_pesos=Decimal('400.00'))
        recalcular_saldo_compra(compra.pk)

        filas, totales, _ = consultar(self.filtro())

        self.assertEqual(filas[0].pagado, Decimal('400.00'))
        self.assertEqual(filas[0].saldo, Decimal('600.00'))
        self.assertEqual(totales.pagado, Decimal('400.00'))

    def test_venta_pagado_es_total_menos_saldo(self):
        """La Venta no tiene campo `pagado`: sale de `cobrado + Σ recibos aplicados`."""
        self.crear_venta(1, 1000, 250)

        filas, totales, _ = consultar(self.filtro(operacion='V'))

        self.assertEqual(filas[0].total, Decimal('1000.00'))
        self.assertEqual(filas[0].pagado, Decimal('750.00'))
        self.assertEqual(filas[0].saldo, Decimal('250.00'))
        self.assertEqual(totales.pagado, Decimal('750.00'))

    def test_los_totales_cierran(self):
        self.crear_compra(1, 1000, 400)
        self.crear_compra(2, 2500, 0)
        self.crear_compra(3, -300, -300)

        _, totales, _ = consultar(self.filtro())

        self.assertEqual(totales.total, totales.pagado + totales.saldo)


class VentasAnuladasTests(BaseFacturasPendientes):
    def setUp(self):
        super().setUp()
        self.activa = self.crear_venta(1, 1000, 1000)
        self.anulada = self.crear_venta(2, 5000, 5000, estado=1)
        self.pendiente_autorizacion = self.crear_venta(3, 700, 700, estado=2)

    def test_anuladas_excluidas_por_defecto(self):
        filas, totales, _ = consultar(self.filtro(operacion='V'))
        pks = [f.pk for f in filas]
        self.assertNotIn(self.anulada.pk, pks)
        # Las pendientes de autorización sí se listan.
        self.assertIn(self.pendiente_autorizacion.pk, pks)
        self.assertEqual(totales.cantidad, 2)

    def test_anuladas_no_contaminan_los_totales(self):
        _, totales, _ = consultar(self.filtro(operacion='V'))
        self.assertEqual(totales.total, Decimal('1700.00'))

    def test_incluir_anuladas_las_muestra_marcadas(self):
        filas, totales, _ = consultar(
            self.filtro(operacion='V', incluir_anuladas=True))
        self.assertEqual(totales.cantidad, 3)
        anulada = next(f for f in filas if f.pk == self.anulada.pk)
        self.assertTrue(anulada.anulada)


class TotalesYTruncadoTests(BaseFacturasPendientes):
    def test_totales_sobre_el_conjunto_completo_aunque_se_trunque(self):
        cantidad = LIMITE_GRILLA + 5
        for i in range(1, cantidad + 1):
            self.crear_compra(i, 100, 100)

        filas, totales, truncado = consultar(self.filtro(), limite=LIMITE_GRILLA)

        self.assertTrue(truncado)
        self.assertEqual(len(filas), LIMITE_GRILLA)
        # El pie no miente: cuenta y suma TODO, no sólo lo visible.
        self.assertEqual(totales.cantidad, cantidad)
        self.assertEqual(totales.total, Decimal('100.00') * cantidad)

    def test_sin_truncado_no_avisa(self):
        self.crear_compra(1, 100, 100)
        filas, _, truncado = consultar(self.filtro(), limite=LIMITE_GRILLA)
        self.assertFalse(truncado)
        self.assertEqual(len(filas), 1)


class AcumuladoTests(BaseFacturasPendientes):
    def test_acumulado_global_termina_en_el_total_de_saldo(self):
        self.crear_compra(1, 1000, 1000)
        self.crear_compra(2, 500, 500)
        self.crear_compra(3, 300, 0)

        filas, totales, _ = consultar(self.filtro())

        self.assertEqual(filas[-1].acum_global, totales.saldo)

    def test_acumulado_de_grupo_reinicia_por_entidad(self):
        otro = ClienteProveedor.objects.create(
            empresa=self.empresa, razon_social="ZZZ PROVEEDOR DOS SA", tipo_entidad=2)
        self.crear_compra(1, 1000, 1000)
        self.crear_compra(2, 500, 500)
        self.crear_compra(3, 700, 700, proveedor=otro)

        filas, _, _ = consultar(self.filtro())

        # Ordenado por razón social: PROVEEDOR UNO (2 filas) y después ZZZ PROVEEDOR DOS.
        self.assertEqual(filas[0].acum_grupo, Decimal('1000.00'))
        self.assertEqual(filas[1].acum_grupo, Decimal('1500.00'))
        self.assertEqual(filas[2].acum_grupo, Decimal('700.00'))
        self.assertEqual(filas[2].acum_global, Decimal('2200.00'))


class ExportacionesTests(BaseFacturasPendientes):
    def setUp(self):
        super().setUp()
        self.crear_compra(1, 1210, 1210)
        self.crear_compra(2, 605, 0)

    def test_excel_devuelve_xlsx_con_los_datos(self):
        respuesta = self.client.get(
            reverse('facturas_pendientes_excel'),
            {'desde': DESDE, 'hasta': HASTA, 'operacion': 'C'})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        ws = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
        self.assertEqual(ws.cell(row=5, column=1).value, 'ID Asiento')
        self.assertEqual(ws.cell(row=5, column=10).value, 'Acum.')
        self.assertEqual(ws.cell(row=6, column=6).value, 'PROVEEDOR UNO SA')
        self.assertEqual(ws.cell(row=6, column=7).value, 1210)

    def test_pdf_devuelve_application_pdf(self):
        respuesta = self.client.get(
            reverse('facturas_pendientes_pdf'),
            {'desde': DESDE, 'hasta': HASTA, 'operacion': 'C'})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')

    def test_pantalla_y_grilla_responden(self):
        self.assertEqual(self.client.get(reverse('facturas_pendientes')).status_code, 200)
        self.assertEqual(
            self.client.get(reverse('facturas_pendientes_grilla')).status_code, 200)


class SoloLecturaTests(BaseFacturasPendientes):
    def test_recorrer_el_modulo_no_altera_los_importes(self):
        """El listado no escribe: `pagado`/`saldo` los mantiene el circuito de tesorería."""
        compra = self.crear_compra(1, 1000, 400)
        antes = (compra.total, compra.pagado, compra.saldo)

        params = {'desde': DESDE, 'hasta': HASTA, 'operacion': 'C'}
        self.client.get(reverse('facturas_pendientes'), params)
        self.client.get(reverse('facturas_pendientes_grilla'), params)
        self.client.get(reverse('facturas_pendientes_excel'), params)
        self.client.get(reverse('facturas_pendientes_pdf'), params)

        compra.refresh_from_db()
        self.assertEqual((compra.total, compra.pagado, compra.saldo), antes)


class SinEmpresaTests(TestCase):
    def test_redirige_si_no_hay_empresa_en_sesion(self):
        User.objects.create_user(username="sinempresa", password="password123")
        cliente = Client()
        cliente.login(username="sinempresa", password="password123")

        respuesta = cliente.get(reverse('facturas_pendientes'))

        self.assertEqual(respuesta.status_code, 302)
        self.assertIn(reverse('seleccion_empresa'), respuesta.url)
