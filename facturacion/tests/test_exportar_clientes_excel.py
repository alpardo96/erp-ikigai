from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth import get_user_model
from empresas.models import Empresa, Sucursal
from facturacion.models import ClienteProveedor, Jurisdiccion
import openpyxl
import io

User = get_user_model()


class ExportarClientesExcelTestCase(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="EMPRESA EXPORT TEST", cuit="30777777778")
        self.user = User.objects.create_user(username="exportuser", password="password123")

        # Usamos get_or_create para evitar colisión si la jurisdicción 901 ya fue precargada en la migración
        self.jurisdiccion, _ = Jurisdiccion.objects.get_or_create(codigo=901, defaults={'nombre': "BUENOS AIRES"})

        self.cliente1 = ClienteProveedor.objects.create(
            empresa=self.empresa,
            razon_social="CLIENTE ALPHA SRL",
            tipo_entidad=1,
            cuit="30111111118",
            domicilio="AV CORRIENTES 1234",
            localidad="CABA",
            jurisdiccion=self.jurisdiccion,
            saldo_inicial=1000.50,
            saldo=1500.00
        )

        self.proveedor1 = ClienteProveedor.objects.create(
            empresa=self.empresa,
            razon_social="PROVEEDOR BETA SA",
            tipo_entidad=2,
            cuit="30222222224",
            domicilio="CALLE FALSA 123",
            localidad="LA PLATA",
            jurisdiccion=self.jurisdiccion,
            saldo_inicial=0.00,
            saldo=500.00
        )

        self.client = Client()
        self.client.login(username="exportuser", password="password123")

        # Set session variable for active empresa
        session = self.client.session
        session['empresa_id'] = self.empresa.id
        session.save()

    def test_exportar_excel_todos(self):
        response = self.client.get(reverse('clientes_exportar_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        # Cargar libro Excel desde el contenido descargado
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Verificar encabezados en la fila 5
        self.assertEqual(ws.cell(row=5, column=1).value, 'ID')
        self.assertEqual(ws.cell(row=5, column=2).value, 'Razón Social')
        self.assertEqual(ws.cell(row=5, column=3).value, 'Tipo')
        self.assertEqual(ws.cell(row=5, column=25).value, 'Observaciones')

        # Verificar datos exportados (fila 6 y 7)
        razones = [ws.cell(row=6, column=2).value, ws.cell(row=7, column=2).value]
        self.assertIn("CLIENTE ALPHA SRL", razones)
        self.assertIn("PROVEEDOR BETA SA", razones)

    def test_exportar_excel_filtrado_tipo(self):
        response = self.client.get(reverse('clientes_exportar_excel') + '?tipo=2')
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        # Fila 6 debe ser PROVEEDOR BETA SA
        self.assertEqual(ws.cell(row=6, column=2).value, 'PROVEEDOR BETA SA')
        # Fila 7 debe estar vacía (solo hay 1 proveedor)
        self.assertIsNone(ws.cell(row=7, column=2).value)

    def test_exportar_excel_filtrado_busqueda(self):
        response = self.client.get(reverse('clientes_exportar_excel') + '?q=ALPHA')
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active

        self.assertEqual(ws.cell(row=6, column=2).value, 'CLIENTE ALPHA SRL')
        self.assertIsNone(ws.cell(row=7, column=2).value)
