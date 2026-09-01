import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone

def exportar_clientes_excel_service(clientes, empresa, filtros):
    """
    Genera un archivo Excel (.xlsx) con todos los campos de la tabla ClienteProveedor
    para los registros filtrados (sin truncar a 50 filas).
    
    `filtros` es un diccionario con {'q': ..., 'tipo_nombre': ...}
    """
    clientes = list(clientes)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clientes y Proveedores"

    is_armeria = bool(empresa and empresa.tipo_actividad and empresa.tipo_actividad.upper() == 'ARMERIA')

    # Definición de encabezados de las columnas de ClienteProveedor
    headers = [
        'ID',
        'Razón Social',
        'Tipo',
        'Tipo Doc',
        'CUIT / DNI',
        'Fecha Nacimiento',
        'Domicilio',
        'C. Postal',
        'Localidad',
        'Provincia / Jurisdicción',
        'Contacto',
        'Teléfono',
        'Correo',
        'Condición IVA',
        'Ingresos Brutos',
        'Saldo Inicial',
        'Saldo Actual',
        'Límite Crédito',
        'Objetivo Mensual',
        'Clasificación',
        'Exige Orden Compra',
        'Cta Patrimonial',
        'Cta Resultado',
        'Código Anterior',
        'Observaciones'
    ]
    if is_armeria:
        headers.extend(['CLU', 'Vencimiento CLU', 'Es Policía'])

    ncols = len(headers)
    ultima_col = get_column_letter(ncols)

    # Título principal de la empresa
    ws.merge_cells(f'A1:{ultima_col}1')
    ws['A1'] = empresa.nombre if empresa else 'ERP Ikigai'
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')

    # Subtítulo del reporte
    ws.merge_cells(f'A2:{ultima_col}2')
    ws['A2'] = "LISTADO DE CLIENTES Y PROVEEDORES"
    ws['A2'].font = Font(size=12, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center')

    # Información de filtros y fecha de emisión
    q_str = filtros.get('q') or '—'
    tipo_str = filtros.get('tipo_nombre') or 'Todos'
    ws.merge_cells(f'A3:{ultima_col}3')
    ws['A3'] = (
        f"Filtro Búsqueda: '{q_str}'   |   Tipo: {tipo_str}   |   "
        f"Registros Exportados: {len(clientes)}   |   "
        f"Emitido: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    )
    ws['A3'].font = Font(size=10, italic=True)
    ws['A3'].alignment = Alignment(horizontal='right')

    # Estilizado de la fila de encabezados de la tabla (Fila 5)
    header_fill = PatternFill("solid", fgColor="0F172A")  # Slate 900
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, text in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.value = text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Índices de columnas numéricas (monetarias): Saldo Inicial(16), Saldo Actual(17), Límite Crédito(18), Objetivo Mensual(19)
    col_num_indices = {16, 17, 18, 19}

    row_idx = 6
    for c in clientes:
        tipo_desc = 'Cliente' if c.tipo_entidad == 1 else ('Proveedor' if c.tipo_entidad == 2 else 'Otro')
        juris_desc = c.jurisdiccion.nombre if c.jurisdiccion else ''
        f_nac = c.fecha_nacimiento.strftime('%d/%m/%Y') if c.fecha_nacimiento else ''
        usa_oc_desc = 'SÍ' if c.usa_orden_compra else 'NO'

        val_saldo_ini = float(c.saldo_inicial or 0)
        val_saldo_act = float(c.saldo or 0)
        val_limite = float(c.limite or 0)
        val_objetivo = float(c.objetivo_mensual or 0)

        row_values = [
            c.codigo_id,
            c.razon_social or '',
            tipo_desc,
            c.get_tipo_documento_display() if hasattr(c, 'get_tipo_documento_display') else (c.tipo_documento or ''),
            c.cuit or '',
            f_nac,
            c.domicilio or '',
            c.codigo_postal or '',
            c.localidad or '',
            juris_desc,
            c.contacto or '',
            c.telefono or '',
            c.correo or '',
            c.condicion_iva or '',
            c.get_tipo_iibb_display() if hasattr(c, 'get_tipo_iibb_display') else (c.tipo_iibb or ''),
            val_saldo_ini,
            val_saldo_act,
            val_limite,
            val_objetivo,
            c.clasificacion or '',
            usa_oc_desc,
            c.cta_pat if c.cta_pat else '',
            c.cta_res if c.cta_res else '',
            c.codigo_anterior or '',
            c.observaciones or ''
        ]

        if is_armeria:
            armeria_obj = getattr(c, 'armeria', None)
            clu_val = armeria_obj.clu if armeria_obj else ''
            clu_vto_val = armeria_obj.clu_vto.strftime('%d/%m/%Y') if (armeria_obj and armeria_obj.clu_vto) else ''
            policia_val = ('SÍ' if armeria_obj.es_policia else 'NO') if armeria_obj else 'NO'
            row_values.extend([clu_val, clu_vto_val, policia_val])

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            if col_idx in col_num_indices:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in (1, 3, 4, 5, 6, 8, 14, 15, 21, 22, 23, 24, 26, 27, 28):
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(horizontal='left')

        row_idx += 1

    # Ajuste automático del ancho de las columnas
    for col_idx in range(1, ncols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row < 5:  # Ignorar celdas combinadas del título
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Configuración de respuesta HTTP para descarga
    fecha_archivo = timezone.localdate().strftime("%Y%m%d")
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="clientes_proveedores_{fecha_archivo}.xlsx"'
    wb.save(response)

    return response
