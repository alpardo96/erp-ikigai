import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from decimal import Decimal


def exportar_compras_excel(compras, empresa, filtros):
    """Genera el .xlsx del listado de compras (con los filtros aplicados) y lo devuelve
    como respuesta HTTP de descarga. `filtros` = {'desde', 'hasta', 'proveedor_nombre',
    'condic_nombre'}. Incluye todos los importes, la condición, la imputación contable
    y el número de asiento."""
    from contable.models import Cuenta

    compras = list(compras)
    # Mapa de imputación contable (cta_imputacion es un pk de Cuenta, no un FK).
    cta_ids = {c.cta_imputacion for c in compras if c.cta_imputacion}
    cta_map = {}
    if cta_ids:
        for cu in Cuenta.objects.filter(pk__in=cta_ids):
            cta_map[cu.pk] = f"{cu.jerarquia} · {cu.cuenta}"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compras"

    # Columnas: las numéricas (para totales y formato) son de la 8 a la 15.
    headers = ['Fecha', 'Comprobante', 'Proveedor', 'CUIT', 'Condición', 'Imputación', 'Asiento',
               'Neto', 'No Gravado', 'Exento', 'IVA', 'Perc. IIBB', 'Perc. IVA', 'Otros', 'Total']
    ncols = len(headers)
    col_num_ini = 8  # primera columna numérica (Neto)
    ultima = get_column_letter(ncols)

    ws.merge_cells(f'A1:{ultima}1')
    ws['A1'] = empresa.nombre if empresa else 'ERP Ikigai'
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'A2:{ultima}2')
    ws['A2'] = "LISTADO DE COMPRAS"
    ws['A2'].font = Font(size=12, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center')

    desde = filtros.get('desde') or '—'
    hasta = filtros.get('hasta') or '—'
    prov = filtros.get('proveedor_nombre') or 'Todos'
    cond = filtros.get('condic_nombre') or 'Todas'
    ws.merge_cells(f'A3:{ultima}3')
    ws['A3'] = (f"Período: {desde} a {hasta}   |   Proveedor: {prov}   |   Condición: {cond}   |   "
                f"Emitido: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}")
    ws['A3'].font = Font(size=10, italic=True)
    ws['A3'].alignment = Alignment(horizontal='right')

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col)
        cell.value = h
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")  # slate-900
        cell.alignment = Alignment(horizontal='center')

    row = 6
    # Acumuladores por columna numérica (Neto, No Grav, Exento, IVA, P.IIBB, P.IVA, Otros, Total)
    campos_num = ['neto', 'no_gravado', 'exento', 'iva', 'p_iibb', 'p_iva', 'otros', 'total']
    totales = {f: Decimal('0') for f in campos_num}

    for c in compras:
        codigo = c.tipo.codigo if c.tipo else ''
        ws.cell(row=row, column=1).value = c.fecha.strftime('%d/%m/%Y') if c.fecha else ''
        ws.cell(row=row, column=2).value = f"{codigo} {c.punto:04d}-{c.numero:08d}"
        ws.cell(row=row, column=3).value = c.proveedor.razon_social if c.proveedor else ''
        ws.cell(row=row, column=4).value = c.proveedor.cuit if c.proveedor else ''
        ws.cell(row=row, column=5).value = {1: 'Real', 2: 'Presupuesto'}.get(c.condic, '')
        ws.cell(row=row, column=6).value = cta_map.get(c.cta_imputacion, '')
        ws.cell(row=row, column=7).value = c.asiento_id or ''
        for i, f in enumerate(campos_num):
            val = getattr(c, f, 0) or 0
            cell = ws.cell(row=row, column=col_num_ini + i)
            cell.value = float(val)
            cell.number_format = '#,##0.00'
            totales[f] += val
        row += 1

    # Fila de totales
    ws.cell(row=row, column=7).value = "TOTALES"
    ws.cell(row=row, column=7).font = Font(bold=True)
    ws.cell(row=row, column=7).alignment = Alignment(horizontal='right')
    border_top = Border(top=Side(style='thin'))
    for i, f in enumerate(campos_num):
        cell = ws.cell(row=row, column=col_num_ini + i)
        cell.value = float(totales[f])
        cell.number_format = '#,##0.00'
        cell.font = Font(bold=True)
        cell.border = border_top

    widths = [12, 20, 38, 14, 13, 34, 10, 14, 13, 13, 14, 13, 13, 13, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="compras_{timezone.localdate().strftime("%Y%m%d")}.xlsx"'
    wb.save(response)
    return response
