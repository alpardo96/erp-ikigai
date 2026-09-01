"""Exportación a Excel del listado de Facturas Pendientes — Plan 056 §3.4.

Equivalente moderno del `cmdExcel.Click` del VFP: mismas columnas útiles (incluida la suma
corrida `Acum.`), sin `Cantidad` ni `litros` (herencia de verticales viejas, siempre en cero)
y sin `F.Pago` (no existe en nuestro modelo).

Estilo calcado de `clientes_excel.py` para que todos los .xlsx del ERP se vean igual.
"""
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.http import HttpResponse
from django.utils import timezone

from facturacion.services.facturas_pendientes import OPERACIONES

HEADERS = [
    'ID Asiento',
    'Fecha',
    'Período',
    'Comprobante',
    'Cód.',
    'Cliente/Proveedor',
    'Total',
    'Pagado',
    'Saldo',
    'Acum.',
    'Condición',
    'C/V',
    'Neto',
    'IVA',
    'No Gravado',
    'Exento',
    'Otros',
    'Clasificación',
    'Descripción',
]

#: 1-based sobre HEADERS: Total, Pagado, Saldo, Acum., Neto, IVA, No Gravado, Exento, Otros.
COLUMNAS_MONETARIAS = {7, 8, 9, 10, 13, 14, 15, 16, 17}
COLUMNAS_CENTRADAS = {1, 2, 3, 5, 11, 12}


def exportar_facturas_pendientes_excel(filas, totales, empresa, filtro):
    """Devuelve el `HttpResponse` con el .xlsx. `filas` viene sin truncar."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas Pendientes"

    ncols = len(HEADERS)
    ultima_col = get_column_letter(ncols)

    ws.merge_cells(f'A1:{ultima_col}1')
    ws['A1'] = empresa.nombre if empresa else 'ERP Ikigai'
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'A2:{ultima_col}2')
    ws['A2'] = "LISTADO DE FACTURAS PENDIENTES"
    ws['A2'].font = Font(size=12, bold=True)
    ws['A2'].alignment = Alignment(horizontal='center')

    ws.merge_cells(f'A3:{ultima_col}3')
    ws['A3'] = (
        f"{filtro.descripcion()}   |   Comprobantes: {totales.cantidad}   |   "
        f"Emitido: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    )
    ws['A3'].font = Font(size=10, italic=True)
    ws['A3'].alignment = Alignment(horizontal='right')

    header_fill = PatternFill("solid", fgColor="0F172A")  # Slate 900
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, text in enumerate(HEADERS, 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.value = text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    row_idx = 6
    for fila in filas:
        valores = [
            fila.asiento_id or '',
            fila.fecha.strftime('%d/%m/%Y') if fila.fecha else '',
            fila.periodo,
            fila.comprobante,
            fila.entidad_id,
            fila.entidad_nombre,
            float(fila.total),
            float(fila.pagado),
            float(fila.saldo),
            float(fila.acum_global),
            fila.condic_nombre,
            fila.operacion,
            float(fila.neto),
            float(fila.iva),
            float(fila.no_gravado),
            float(fila.exento),
            float(fila.otros),
            fila.entidad_clasificacion,
            fila.descripcion,
        ]
        for col_idx, val in enumerate(valores, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            if col_idx in COLUMNAS_MONETARIAS:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif col_idx in COLUMNAS_CENTRADAS:
                cell.alignment = Alignment(horizontal='center')
            else:
                cell.alignment = Alignment(horizontal='left')
        row_idx += 1

    # Fila de totales, equivalente a las cajas del pie de la pantalla.
    if filas:
        fila_tot = row_idx + 1
        ws.cell(row=fila_tot, column=6).value = 'TOTALES'
        ws.cell(row=fila_tot, column=6).alignment = Alignment(horizontal='right')
        for col_idx, valor in ((7, totales.total), (8, totales.pagado), (9, totales.saldo)):
            cell = ws.cell(row=fila_tot, column=col_idx)
            cell.value = float(valor)
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right')
        for col_idx in range(1, ncols + 1):
            ws.cell(row=fila_tot, column=col_idx).font = Font(bold=True)

    for col_idx in range(1, ncols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row < 5:  # Las celdas combinadas del encabezado no cuentan.
                continue
            max_len = max(max_len, len(str(cell.value or '')))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    etiqueta = dict(OPERACIONES).get(filtro.operacion, filtro.operacion).lower()
    fecha_archivo = timezone.localdate().strftime("%Y%m%d")
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = (
        f'attachment; filename="facturas_pendientes_{etiqueta}_{fecha_archivo}.xlsx"')
    wb.save(response)
    return response
