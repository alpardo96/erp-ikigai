import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from decimal import Decimal
from datetime import date, datetime

def exportar_ventas_producto_excel_service(items, empresa, filtros):
    """
    Genera y devuelve una respuesta HttpResponse con la planilla .xlsx estilizada
    del Reporte de Ventas por Producto.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos Vendidos"

    # Encabezados de 38 columnas (sin subfamilia)
    headers = [
        'ID.Vta', 'ID.Asto', 'Fecha', 'Comprobante', 'id.cod', 'Cliente', 'Domicilio', 'Credencial',
        'Codigo', 'Cod Prov', 'Producto', 'Calibre', 'Serie', 'CUIM', 'Cantidad', 'cto.Rep.',
        'fec.Act.', 'moneda', 'cotiz', 'alic.iva', 'pcio.T.', 'Neto', 'Total',
        'Stock', 'id.pro', 'Proveedor', 'id.rubro', 'rubro', 'id.flia', 'familia',
        'id.marca', 'marca', 'id.vdor', 'Vendedor',
        'Cto.Adq.', 'Margen', 'Mg.Rubro', 'Sucursal'
    ]

    ncols = len(headers)

    # Fila 1: Título institucional (sin combinar celdas)
    empresa_nombre = empresa.nombre.upper() if empresa else 'ERP IKIGAI'
    ws['A1'] = f"{empresa_nombre} - LISTADO DE PRODUCTOS VENDIDOS"
    ws['A1'].font = Font(size=14, bold=True, color="1E293B")
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    # Fila 2: Subtítulo con parámetros de período y filtros (sin combinar celdas)
    desde = filtros.get('desde') or '—'
    hasta = filtros.get('hasta') or '—'
    suc = filtros.get('sucursal_nombre') or 'Todas'
    cli = filtros.get('cliente_nombre') or 'Todos'
    ws['A2'] = (f"Periodo: {desde} al {hasta} | Sucursal: {suc} | Cliente: {cli} | "
                f"Emisión: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}")
    ws['A2'].font = Font(size=10, italic=True, color="64748B")
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')

    # Fila 4: Encabezados de la tabla
    header_fill = PatternFill("solid", fgColor="0F172A") # slate-900 elegante
    header_font = Font(bold=True, color="FFFFFF", size=10)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Filas de datos (empezando en fila 5)
    row_idx = 5
    tot_cant = Decimal("0.00")
    tot_neto = Decimal("0.00")
    tot_total = Decimal("0.00")

    border_light = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Stock por (producto, sucursal) en UNA consulta
    from productos.models import StockSucursal
    _productos = {i.producto_id for i in items if i.producto_id}
    _sucursales = {i.venta.sucursal_id for i in items if i.venta_id and i.venta.sucursal_id}
    stock_por_clave = {}
    if _productos and _sucursales:
        stock_por_clave = {
            (s['producto_id'], s['sucursal_id']): s['cantidad']
            for s in StockSucursal.objects.filter(
                producto_id__in=_productos, sucursal_id__in=_sucursales
            ).values('producto_id', 'sucursal_id', 'cantidad')
        }

    for item in items:
        vta = item.venta
        prod = item.producto
        cli_obj = vta.cliente if vta else None
        prov_obj = prod.proveedor if prod else None
        rubro_obj = prod.rubro if prod else None
        flia_obj = prod.familia if prod else None
        marca_obj = prod.marca if prod else None
        vendedor_obj = vta.vendedor if vta else None

        cant = Decimal(str(item.cantidad or 0))
        pcio_vta = Decimal(str(item.precio_unitario or 0))
        total = Decimal(str(item.total or 0))
        alic_iva = Decimal(str(item.iva_alicuota or 0))

        # Neto de la línea
        factor_iva = Decimal("1.00") + (alic_iva / Decimal("100.00"))
        neto = (total / factor_iva).quantize(Decimal("0.01")) if factor_iva else total

        pcio_t = (total / cant).quantize(Decimal("0.01")) if cant else pcio_vta

        # Acumuladores
        tot_cant += cant
        tot_neto += neto
        tot_total += total

        comprobante_str = f"{vta.tipo.codigo if vta.tipo else ''} {vta.punto:05d}-{vta.numero:08d}" if vta else ''

        # Fechas tipo Date puro (sin timestamp hora)
        fecha_date = None
        if vta and vta.fecha:
            if isinstance(vta.fecha, datetime):
                fecha_date = vta.fecha.date()
            elif isinstance(vta.fecha, date):
                fecha_date = vta.fecha
            else:
                try:
                    fecha_date = datetime.strptime(str(vta.fecha)[:10], '%Y-%m-%d').date()
                except:
                    fecha_date = None

        fec_act_date = None
        if prod and prod.fec_act:
            if isinstance(prod.fec_act, datetime):
                fec_act_date = prod.fec_act.date()
            elif isinstance(prod.fec_act, date):
                fec_act_date = prod.fec_act
            else:
                try:
                    fec_act_date = datetime.strptime(str(prod.fec_act)[:10], '%Y-%m-%d').date()
                except:
                    fec_act_date = None

        domicilio_str = (cli_obj.domicilio_completo if cli_obj else '') or (vta.cliente_domicilio if vta else '') or ''
        calibre_str = prod.unidad_venta if prod and prod.unidad_venta else ''
        serie_str = item.serie or (item.subproducto.serie if item.subproducto else '') or ''
        cuim_str = item.cuim or (item.subproducto.cuim if item.subproducto else '') or ''
        credencial_str = item.credencial or ''

        row_data = [
            vta.ventas_id if vta else '',                             # 1. ID.Vta (A)
            vta.asiento_id if vta and vta.asiento_id else '',        # 2. ID.Asto (B)
            fecha_date,                                              # 3. Fecha (C) - Tipo fecha pura
            comprobante_str,                                         # 4. Comprobante (D)
            cli_obj.codigo_id if cli_obj else '',                    # 5. id.cod (E)
            vta.cliente_razon_social or (cli_obj.razon_social if cli_obj else ''), # 6. Cliente (F)
            domicilio_str,                                           # 7. Domicilio Completo (G)
            credencial_str,                                          # 8. Credencial (H)
            prod.pk if prod else '',                                 # 9. Codigo (I)
            prod.cod_prov if prod and prod.cod_prov else '',         # 10. Cod Prov (J)
            item.concepto or (prod.detalle if prod else ''),         # 11. Producto (K)
            calibre_str,                                             # 12. Calibre (L)
            serie_str,                                               # 13. Serie (M)
            cuim_str,                                                # 14. CUIM (N)
            float(cant),                                             # 15. Cantidad (O)
            float(prod.cto_rep) if prod and prod.cto_rep else 0.0,   # 16. cto.Rep. (P)
            fec_act_date,                                            # 17. fec.Act. / fec.Rep. (Q) - Tipo fecha pura
            item.moneda_origen or (vta.moneda if vta else 'PES'),    # 18. moneda (R)
            float(vta.cotizacion) if vta else 1.0,                   # 19. cotiz (S)
            float(alic_iva),                                         # 20. alic.iva (T)
            float(pcio_t),                                           # 21. pcio.T. (U)
            float(neto),                                             # 22. Neto (V)
            float(total),                                            # 23. Total (W)
            # Stock REAL de la sucursal de la venta
            float(stock_por_clave.get(
                (item.producto_id, vta.sucursal_id if vta else None), 0) or 0),   # 24. Stock (X)
            prov_obj.codigo_id if prov_obj else '',                  # 25. id.pro (Y)
            prov_obj.razon_social if prov_obj else '',               # 26. Proveedor (Z)
            rubro_obj.pk if rubro_obj else '',                       # 27. id.rubro (AA)
            rubro_obj.detalle if rubro_obj else '',                  # 28. rubro (AB)
            flia_obj.pk if flia_obj else '',                         # 29. id.flia (AC)
            flia_obj.detalle if flia_obj else '',                    # 30. familia (AD)
            marca_obj.pk if marca_obj else '',                       # 31. id.marca (AE)
            marca_obj.detalle if marca_obj else '',                  # 32. marca (AF)
            vendedor_obj.pk if vendedor_obj else '',                 # 33. id.vdor (AG)
            vendedor_obj.username if vendedor_obj else '',           # 34. Vendedor (AH)
            float(prod.cto_adq) if prod and prod.cto_adq else 0.0,   # 35. Cto.Adq. (AI)
            float(prod.margen) if prod and prod.margen else 0.0,     # 36. Margen (AJ)
            float(rubro_obj.margen) if rubro_obj and rubro_obj.margen else 0.0, # 37. Mg.Rubro (AK)
            vta.sucursal.nombre if vta and vta.sucursal else ''      # 38. Sucursal (AL)
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_light
            # Formatos de fecha (columnas 3 y 17)
            if col_idx in [3, 17]:
                cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(horizontal='center', vertical='center')
            # Formatos numéricos
            elif col_idx == 15:  # Cantidad
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif col_idx in [16, 19, 20, 21, 22, 23, 24, 35, 36, 37]:  # Moneda / importes / porcentajes
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right', vertical='center')

        row_idx += 1

    # Fila final de Totales
    ws.cell(row=row_idx, column=11, value="TOTALES GENERALES").font = Font(bold=True)
    ws.cell(row=row_idx, column=11).alignment = Alignment(horizontal='right', vertical='center')

    cell_cant = ws.cell(row=row_idx, column=15, value=float(tot_cant))
    cell_cant.number_format = '#,##0.00'
    cell_cant.font = Font(bold=True)

    cell_neto = ws.cell(row=row_idx, column=22, value=float(tot_neto))
    cell_neto.number_format = '#,##0.00'
    cell_neto.font = Font(bold=True)

    cell_total = ws.cell(row=row_idx, column=23, value=float(tot_total))
    cell_total.number_format = '#,##0.00'
    cell_total.font = Font(bold=True)

    border_total = Border(top=Side(style='double'), bottom=Side(style='double'))
    for c in range(1, ncols + 1):
        ws.cell(row=row_idx, column=c).border = border_total

    # Ajuste automático del ancho de columnas
    for col_idx in range(1, ncols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1]))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    timestamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    response['Content-Disposition'] = f'attachment; filename="productos_vendidos_{timestamp}.xlsx"'
    wb.save(response)
    return response
