import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from decimal import Decimal

# Mapeo de respaldo para códigos estándar de comprobantes AFIP/ARCA
MAPA_TIPO_COMPROBANTE = {
    '001': 'Factura A',
    '002': 'Nota de Débito A',
    '003': 'Nota de Crédito A',
    '004': 'Recibo A',
    '006': 'Factura B',
    '007': 'Nota de Débito B',
    '008': 'Nota de Crédito B',
    '009': 'Recibo B',
    '011': 'Factura C',
    '012': 'Nota de Débito C',
    '013': 'Nota de Crédito C',
    '015': 'Recibo C',
    '051': 'Factura M',
    '052': 'Nota de Débito M',
    '053': 'Nota de Crédito M',
    '099': 'Remito X',
    '000': 'Comprobante X',
}

def obtener_nombre_comprobante(tipo_obj):
    """
    Devuelve la denominación legible del tipo de comprobante.
    Prioriza el campo detalle del modelo TipoComprobante, con fallback a tabla AFIP.
    """
    if not tipo_obj:
        return 'Sin Comprobante'
    
    codigo = str(tipo_obj.codigo or '').strip().zfill(3)
    
    if tipo_obj.detalle and tipo_obj.detalle.strip():
        detalle = tipo_obj.detalle.strip()
        if detalle.isdigit() or len(detalle) <= 3:
            return MAPA_TIPO_COMPROBANTE.get(codigo, f"Comprobante {codigo}")
        return detalle.title()
    
    return MAPA_TIPO_COMPROBANTE.get(codigo, f"Comprobante {codigo}")


def exportar_ventas_listado_excel_service(ventas_qs, empresa, filtros):
    """
    Genera y devuelve una respuesta HttpResponse con la planilla .xlsx estilizada
    del Listado de Ventas con el detalle más completo posible.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Listado de Ventas"

    # Definición de encabezados de columnas
    headers = [
        'ID Venta',
        'Asiento',
        'Fecha',
        'Hora Carga',
        'Tipo Comprobante',
        'Pto. Vta.',
        'Número',
        'Comprobante Completo',
        'CAE',
        'Vto. CAE',
        'Cód. Cliente',
        'Cliente / Razón Social',
        'CUIT / DNI',
        'Condición IVA',
        'Domicilio',
        'Localidad',
        'Provincia',
        'Cond. Circuito',
        'Cond. Venta',
        'Sucursal',
        'Vendedor',
        'Usuario Carga',
        'Moneda',
        'Cotización',
        'Neto Gravado',
        'No Gravado',
        'Exento',
        'IVA',
        'Perc. IIBB',
        'Perc. IVA',
        'Perc. Ganancias',
        'Otras Perc.',
        'Total',
        'Cobrado',
        'Saldo',
        'Estado',
    ]

    ncols = len(headers)
    ultima_col = get_column_letter(ncols)

    # 1. Fila 1: Título institucional
    empresa_nombre = empresa.nombre.upper() if empresa else 'ERP IKIGAI'
    ws.merge_cells(f'A1:{ultima_col}1')
    ws['A1'] = f"{empresa_nombre} - LISTADO GENERAL DE VENTAS"
    ws['A1'].font = Font(size=14, bold=True, color="1E293B")
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')

    # 2. Fila 2: Subtítulo con parámetros de filtros y fecha de emisión
    desde = filtros.get('desde') or '—'
    hasta = filtros.get('hasta') or '—'
    suc = filtros.get('sucursal_nombre') or 'Todas'
    cli = filtros.get('cliente_nombre') or 'Todos'
    vdor = filtros.get('vendedor_nombre') or 'Todos'
    cond_txt = filtros.get('condic_nombre') or 'Todas'
    tipo_txt = filtros.get('tipo_nombre') or 'Todos'
    
    ws.merge_cells(f'A2:{ultima_col}2')
    ws['A2'] = (
        f"Período: {desde} al {hasta} | Sucursal: {suc} | Cliente: {cli} | "
        f"Vendedor: {vdor} | Tipo: {tipo_txt} | Condición: {cond_txt} | "
        f"Emisión: {timezone.localtime().strftime('%d/%m/%Y %H:%M')}"
    )
    ws['A2'].font = Font(size=9, italic=True, color="64748B")
    ws['A2'].alignment = Alignment(horizontal='left', vertical='center')

    # 3. Fila 4: Encabezados de la tabla con diseño Slate-900
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF", size=9)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    border_light = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Columnas numéricas a acumular (índices 1-based en Excel):
    # 25: Neto, 26: No Gravado, 27: Exento, 28: IVA, 29: P.IIBB, 30: P.IVA, 31: P.Gcia, 32: Otras, 33: Total, 34: Cobrado, 35: Saldo
    col_totales = {
        25: Decimal('0.00'), # Neto
        26: Decimal('0.00'), # No Gravado
        27: Decimal('0.00'), # Exento
        28: Decimal('0.00'), # IVA
        29: Decimal('0.00'), # P.IIBB
        30: Decimal('0.00'), # P.IVA
        31: Decimal('0.00'), # P.Gcia
        32: Decimal('0.00'), # Otras Perc.
        33: Decimal('0.00'), # Total
        34: Decimal('0.00'), # Cobrado
        35: Decimal('0.00'), # Saldo
    }

    row_idx = 5
    for vta in ventas_qs:
        cli_obj = vta.cliente
        suc_obj = vta.sucursal
        vdor_obj = vta.vendedor
        usu_obj = vta.usuario
        tipo_obj = vta.tipo

        nombre_comprobante = obtener_nombre_comprobante(tipo_obj)
        pto_str = f"{vta.punto:05d}" if vta.punto is not None else "00000"
        num_str = f"{vta.numero:08d}" if vta.numero is not None else "00000000"
        comprobante_completo = f"{nombre_comprobante} {pto_str}-{num_str}"

        # Clasificación contable
        condic_map = {1: 'Real', 2: 'Presupuestado', 3: 'Ajuste', 4: 'Auditoría'}
        condic_str = condic_map.get(vta.condic, str(vta.condic or ''))

        # Condición comercial
        cond_venta_str = dict(vta.CONDICION_VENTA_CHOICES).get(vta.condicion_venta, vta.condicion_venta or '')

        # Estado
        estado_map = {0: 'Activa', 1: 'Anulada', 2: 'Pendiente', 3: 'Rechazada'}
        estado_str = estado_map.get(vta.estado, str(vta.estado or ''))

        # Razón social y CUIT (con fallback a campos desnormalizados de Venta)
        razon_social = vta.cliente_razon_social or (cli_obj.razon_social if cli_obj else '')
        cuit_str = vta.cliente_cuit or (cli_obj.cuit if cli_obj else '')
        domicilio_str = vta.cliente_domicilio or (cli_obj.domicilio if cli_obj else '')
        localidad_str = cli_obj.localidad if cli_obj else ''
        provincia_str = cli_obj.jurisdiccion.nombre if cli_obj and cli_obj.jurisdiccion else ''
        cond_iva_str = cli_obj.get_condicion_iva_display() if cli_obj and hasattr(cli_obj, 'get_condicion_iva_display') else (cli_obj.condicion_iva if cli_obj else '')

        # Vendedor / Usuario
        vendedor_nombre = vdor_obj.get_full_name() or vdor_obj.username if vdor_obj else ''
        usuario_nombre = usu_obj.username if usu_obj else ''

        # Importes
        neto_val = Decimal(str(vta.neto or 0))
        no_grav_val = Decimal(str(vta.no_gravado or 0))
        exento_val = Decimal(str(vta.exento or 0))
        iva_val = Decimal(str(vta.iva or 0))
        p_iibb_val = Decimal(str(vta.p_iibb or 0))
        p_iva_val = Decimal(str(vta.p_iva or 0))
        p_gcia_val = Decimal(str(vta.p_gcia or 0))
        otras_perc_val = Decimal(str(vta.p_sircreb or 0)) + Decimal(str(vta.p_mun or 0)) + Decimal(str(vta.p_recbc or 0)) + Decimal(str(vta.otros or 0))
        total_val = Decimal(str(vta.total or 0))
        cobrado_val = Decimal(str(vta.cobrado or 0))
        saldo_val = Decimal(str(vta.saldo or 0))

        # Sumar a totales si no está anulada
        if vta.estado != 1:
            col_totales[25] += neto_val
            col_totales[26] += no_grav_val
            col_totales[27] += exento_val
            col_totales[28] += iva_val
            col_totales[29] += p_iibb_val
            col_totales[30] += p_iva_val
            col_totales[31] += p_gcia_val
            col_totales[32] += otras_perc_val
            col_totales[33] += total_val
            col_totales[34] += cobrado_val
            col_totales[35] += saldo_val

        # Armado de la fila
        row_data = [
            vta.ventas_id,                                                   # 1: ID Venta
            vta.asiento_id or '',                                            # 2: Asiento
            vta.fecha.strftime('%d/%m/%Y') if vta.fecha else '',             # 3: Fecha
            vta.fec_vta.strftime('%d/%m/%Y %H:%M') if vta.fec_vta else '',   # 4: Hora Carga
            nombre_comprobante,                                              # 5: Tipo Comprobante
            pto_str,                                                         # 6: Pto. Vta.
            num_str,                                                         # 7: Número
            comprobante_completo,                                            # 8: Comprobante Completo
            vta.cae or '',                                                   # 9: CAE
            vta.vto_cae.strftime('%d/%m/%Y') if vta.vto_cae else '',         # 10: Vto. CAE
            cli_obj.codigo_id if cli_obj else '',                            # 11: Cód. Cliente
            razon_social,                                                    # 12: Cliente / Razón Social
            cuit_str,                                                        # 13: CUIT / DNI
            cond_iva_str,                                                    # 14: Condición IVA
            domicilio_str,                                                   # 15: Domicilio
            localidad_str,                                                   # 16: Localidad
            provincia_str,                                                   # 17: Provincia
            condic_str,                                                      # 18: Cond. Circuito
            cond_venta_str,                                                  # 19: Cond. Venta
            suc_obj.nombre if suc_obj else '',                               # 20: Sucursal
            vendedor_nombre,                                                 # 21: Vendedor
            usuario_nombre,                                                  # 22: Usuario Carga
            vta.moneda or 'PES',                                             # 23: Moneda
            float(vta.cotizacion or 1.0),                                    # 24: Cotización
            float(neto_val),                                                 # 25: Neto Gravado
            float(no_grav_val),                                              # 26: No Gravado
            float(exento_val),                                               # 27: Exento
            float(iva_val),                                                  # 28: IVA
            float(p_iibb_val),                                               # 29: Perc. IIBB
            float(p_iva_val),                                                # 30: Perc. IVA
            float(p_gcia_val),                                               # 31: Perc. Ganancias
            float(otras_perc_val),                                           # 32: Otras Perc.
            float(total_val),                                                # 33: Total
            float(cobrado_val),                                              # 34: Cobrado
            float(saldo_val),                                                # 35: Saldo
            estado_str,                                                      # 36: Estado
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_light
            cell.font = Font(size=9)

            # Alineaciones y formatos
            if col_idx in (1, 2, 6, 7, 9, 11, 13, 23):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in (3, 4, 10):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx in (24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 34, 35):
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0.00'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

            # Si el comprobante está anulado, pintar en gris suave cursiva
            if vta.estado == 1:
                cell.font = Font(size=9, italic=True, color="94A3B8")

        row_idx += 1

    # 4. Fila final de Totales
    ws.cell(row=row_idx, column=24, value="TOTALES:").font = Font(bold=True, size=10, color="0F172A")
    ws.cell(row=row_idx, column=24).alignment = Alignment(horizontal='right', vertical='center')

    border_total = Border(
        top=Side(style='thin', color='0F172A'),
        bottom=Side(style='double', color='0F172A')
    )

    for col_idx in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        if col_idx in col_totales:
            cell.value = float(col_totales[col_idx])
            cell.font = Font(bold=True, size=10, color="0F172A")
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.fill = PatternFill("solid", fgColor="F1F5F9")
            cell.border = border_total
        else:
            if col_idx >= 24:
                cell.border = border_total
                cell.fill = PatternFill("solid", fgColor="F1F5F9")

    # 5. Ajuste automático de anchos de columna
    for col_idx in range(1, ncols + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(4, min(row_idx + 1, 50)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

    # Anchos fijos específicos para columnas que requieren holgura
    ws.column_dimensions['A'].width = 11  # ID Venta
    ws.column_dimensions['B'].width = 10  # Asiento
    ws.column_dimensions['C'].width = 12  # Fecha
    ws.column_dimensions['D'].width = 16  # Hora Carga
    ws.column_dimensions['E'].width = 22  # Tipo Comprobante
    ws.column_dimensions['H'].width = 30  # Comprobante Completo
    ws.column_dimensions['L'].width = 36  # Cliente
    ws.column_dimensions['M'].width = 15  # CUIT
    ws.column_dimensions['N'].width = 22  # Condición IVA
    ws.column_dimensions['O'].width = 30  # Domicilio
    ws.column_dimensions['T'].width = 18  # Sucursal
    ws.column_dimensions['U'].width = 18  # Vendedor

    # Generación de la respuesta HTTP
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    response['Content-Disposition'] = f'attachment; filename="ventas_listado_{timestamp}.xlsx"'
    wb.save(response)
    return response
