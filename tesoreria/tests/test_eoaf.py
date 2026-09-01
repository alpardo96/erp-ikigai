"""Plan 050 — Estado de Origen y Aplicación de Fondos.

Los asientos se construyen a mano en lugar de pasar por los circuitos de recibo / orden de pago /
mostrador: lo que se prueba acá es la DESCOMPOSICIÓN de un movimiento de fondos, no la carga.
"""

from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone

from contable.models import Asiento, AsientoLinea, Cuenta
from empresas.models import Empresa, Ejercicio, Sucursal
from facturacion.models import ClienteProveedor
from tesoreria.services.eoaf import detalle_de_cuenta, estado_origen_aplicacion_fondos

User = get_user_model()

CERO = Decimal('0.00')
DESDE = timezone.datetime(2026, 8, 1).date()
HASTA = timezone.datetime(2026, 8, 31).date()


class BaseEOAF(TestCase):
    def setUp(self):
        self.usuario = User.objects.create_user(username="tesorero", password="clave")
        self.empresa = Empresa.objects.create(
            nombre="FONDOS SA", cuit="30555444332", direccion="C 1", correo="f@t.com")
        self.sucursal = Sucursal.objects.create(empresa=self.empresa, nombre="CENTRAL")
        self.ejercicio = Ejercicio.objects.create(
            empresa=self.empresa, ejercicio="2026",
            inicio=timezone.datetime(2026, 1, 1).date(),
            cierre=timezone.datetime(2026, 12, 31).date())

        def cuenta(jerarquia, nombre, tipo, imputable=1, disponibilidad=''):
            return Cuenta.objects.create(
                jerarquia=jerarquia, cuenta=nombre, imputable=imputable, tipo=tipo,
                empresa=self.empresa, tipo_disponibilidad=disponibilidad)

        # Disponibilidades (el bolsillo)
        self.caja = cuenta("111001", "CAJA", "A", disponibilidad="EFE")
        self.banco = cuenta("111011", "BANCO NACION", "A", disponibilidad="BCO")
        self.caja_tesoreria = cuenta("111020", "CAJA TESORERIA", "A", disponibilidad="EFE")

        # Sumarizadoras + contrapartidas
        self.rubro_creditos = cuenta("112", "CREDITOS POR VENTAS", "A", imputable=0)
        self.clientes = cuenta("112001", "DEUDORES POR VENTAS", "A")
        self.rubro_deudas = cuenta("211", "DEUDAS COMERCIALES", "P", imputable=0)
        self.proveedores = cuenta("211001", "PROVEEDORES VARIOS", "P")
        self.rubro_gastos = cuenta("531", "GASTOS ADMINISTRATIVOS", "R", imputable=0)
        self.gastos_luz = cuenta("531001", "ADM-ENERGIA", "R")
        self.gastos_tel = cuenta("531002", "ADM-TELEFONIA", "R")

        self.cliente = ClienteProveedor.objects.create(
            empresa=self.empresa, razon_social="CLIENTE UNO SRL", cuit="20111111112")

    # ------------------------------------------------------------------ helpers
    def asentar(self, *lineas, fecha=None, condic=1, anulado=False, empresa=None):
        """`lineas` son tuplas (cuenta, debe, haber)."""
        asiento = Asiento.objects.create(
            empresa=empresa or self.empresa, ejercicio=self.ejercicio,
            fecha=fecha or timezone.datetime(2026, 8, 15).date(),
            concepto="MOVIMIENTO", monto=Decimal("0"), condic=condic, anulado=anulado)
        for orden, (cta, debe, haber) in enumerate(lineas, start=1):
            AsientoLinea.objects.create(asiento=asiento, orden=orden, cuenta=cta,
                                        debe=Decimal(str(debe)), haber=Decimal(str(haber)))
        return asiento

    def reporte(self, **kwargs):
        kwargs.setdefault('empresa_id', self.empresa.id)
        kwargs.setdefault('desde', DESDE)
        kwargs.setdefault('hasta', HASTA)
        return estado_origen_aplicacion_fondos(**kwargs)

    def fila(self, resultado, cuenta):
        for f in resultado['filas']:
            if f['jerarquia'] == cuenta.jerarquia:
                return f
        return None


class SignoTest(BaseEOAF):
    """Contrapartida al HABER = origen (Ingresos). Al DEBE = aplicación (Egresos)."""

    def test_cobranza_deja_al_cliente_en_ingresos(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        fila = self.fila(self.reporte(), self.clientes)

        self.assertEqual(fila['ingresos'], Decimal("1000.00"))
        self.assertEqual(fila['egresos'], CERO)
        self.assertEqual(fila['neto'], Decimal("1000.00"))
        self.assertEqual(fila['medios']['EFE'], Decimal("1000.00"))

    def test_pago_a_proveedor_deja_al_proveedor_en_egresos(self):
        self.asentar((self.proveedores, 800, 0), (self.caja, 0, 800))
        fila = self.fila(self.reporte(), self.proveedores)

        self.assertEqual(fila['egresos'], Decimal("800.00"))
        self.assertEqual(fila['ingresos'], CERO)
        self.assertEqual(fila['neto'], Decimal("-800.00"))

    def test_pago_de_gasto_por_banco_cae_en_la_columna_banco(self):
        self.asentar((self.gastos_luz, 500, 0), (self.banco, 0, 500))
        fila = self.fila(self.reporte(), self.gastos_luz)

        self.assertEqual(fila['egresos'], Decimal("500.00"))
        self.assertEqual(fila['medios']['BCO'], Decimal("-500.00"))
        self.assertEqual(fila['medios']['EFE'], CERO)


class TrasladosTest(BaseEOAF):

    def test_traslado_entre_disponibilidades_no_genera_filas(self):
        """Retiro de caja mostrador a tesorería: plata cambiando de bolsillo, no un origen."""
        self.asentar((self.caja_tesoreria, 5000, 0), (self.caja, 0, 5000))
        resultado = self.reporte()

        self.assertEqual(resultado['filas'], [])
        self.assertEqual(resultado['totales']['ingresos'], CERO)
        self.assertEqual(resultado['totales']['egresos'], CERO)

    def test_el_traslado_no_ensucia_un_periodo_con_movimientos_reales(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.asentar((self.caja_tesoreria, 400, 0), (self.caja, 0, 400))
        resultado = self.reporte()

        self.assertEqual(len(resultado['filas']), 2)   # 112 sumarizadora + 112001
        self.assertEqual(resultado['totales']['ingresos'], Decimal("1000.00"))


class ProrrateoTest(BaseEOAF):

    def test_varias_contrapartidas_reparten_el_desglose_sin_perder_centavos(self):
        """El prorrateo debe cerrar exacto: la última contrapartida absorbe el residuo."""
        self.asentar(
            (self.gastos_luz, Decimal("333.33"), 0),
            (self.gastos_tel, Decimal("333.33"), 0),
            (self.proveedores, Decimal("333.34"), 0),
            (self.caja, 0, Decimal("1000.00")),
        )
        resultado = self.reporte()

        efectivo = sum(f['medios']['EFE'] for f in resultado['filas'] if f['imputable'])
        self.assertEqual(efectivo, Decimal("-1000.00"), "El desglose no reconstruye el asiento")
        self.assertEqual(resultado['totales']['egresos'], Decimal("1000.00"))


class JerarquiaTest(BaseEOAF):

    def test_la_sumarizadora_acumula_a_sus_imputables(self):
        self.asentar((self.gastos_luz, 300, 0), (self.caja, 0, 300))
        self.asentar((self.gastos_tel, 200, 0), (self.caja, 0, 200))
        resultado = self.reporte()

        rubro = self.fila(resultado, self.rubro_gastos)
        self.assertEqual(rubro['egresos'], Decimal("500.00"))
        self.assertFalse(rubro['imputable'])

    def test_no_hay_doble_conteo_entre_niveles(self):
        """Los totales suman solo imputables: la sumarizadora ya las contiene."""
        self.asentar((self.gastos_luz, 300, 0), (self.caja, 0, 300))
        self.asentar((self.gastos_tel, 200, 0), (self.caja, 0, 200))
        resultado = self.reporte()

        self.assertEqual(resultado['totales']['egresos'], Decimal("500.00"))

    def test_cuenta_sin_movimiento_no_aparece(self):
        self.asentar((self.gastos_luz, 300, 0), (self.caja, 0, 300))
        resultado = self.reporte()

        self.assertIsNone(self.fila(resultado, self.gastos_tel))
        self.assertIsNone(self.fila(resultado, self.clientes))


class CoherenciaTest(BaseEOAF):

    def test_el_flujo_neto_iguala_la_variacion_de_disponibilidades(self):
        """El control que en el sistema legado NO cerraba.

        La suma de los flujos netos de las imputables tiene que dar exactamente lo que variaron
        las disponibilidades en el período. En el legado difería porque las imputables sumaban la
        'Dispon. Inicial' al neto y las sumarizadoras no, y porque el reporte solo mostraba tres
        de los seis tipos de disponibilidad.
        """
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.asentar((self.proveedores, 800, 0), (self.banco, 0, 800))
        self.asentar((self.gastos_luz, 250, 0), (self.caja, 0, 250))
        self.asentar((self.caja_tesoreria, 400, 0), (self.caja, 0, 400))   # traslado

        resultado = self.reporte()
        totales = resultado['totales']

        neto_por_medio = sum(totales['medios'].values(), CERO)
        self.assertEqual(totales['neto'], neto_por_medio)
        self.assertEqual(totales['neto'], totales['ingresos'] - totales['egresos'])
        # caja: +1000 −250 −400 +400(tesorería) = +750 ; banco: −800
        self.assertEqual(neto_por_medio, Decimal("-50.00"))


class FiltrosTest(BaseEOAF):

    def test_filtro_de_condicion(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000), condic=1)
        self.asentar((self.caja, 300, 0), (self.clientes, 0, 300), condic=2)

        self.assertEqual(self.reporte(condics=(1,))['totales']['ingresos'], Decimal("1000.00"))
        self.assertEqual(self.reporte(condics=(2,))['totales']['ingresos'], Decimal("300.00"))
        self.assertEqual(self.reporte(condics=(1, 2))['totales']['ingresos'], Decimal("1300.00"))

    def test_sin_condiciones_no_devuelve_nada(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.assertEqual(self.reporte(condics=())['filas'], [])

    def test_rango_de_fechas_por_la_fecha_del_comprobante(self):
        """Un comprobante retroactivo cae en el mes de SU fecha, no en el de carga."""
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000),
                     fecha=timezone.datetime(2026, 7, 15).date())
        self.asentar((self.caja, 500, 0), (self.clientes, 0, 500),
                     fecha=timezone.datetime(2026, 8, 15).date())

        self.assertEqual(self.reporte()['totales']['ingresos'], Decimal("500.00"))
        self.assertEqual(
            self.reporte(desde=timezone.datetime(2026, 7, 1).date())['totales']['ingresos'],
            Decimal("1500.00"))

    def test_asiento_anulado_no_entra(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000), anulado=True)
        self.assertEqual(self.reporte()['filas'], [])

    def test_aislamiento_entre_empresas(self):
        otra = Empresa.objects.create(nombre="OTRA SA", cuit="30999888771",
                                      direccion="C 2", correo="o@t.com")
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000), empresa=otra)
        self.assertEqual(self.reporte()['filas'], [])


class DrillDownTest(BaseEOAF):

    def test_el_detalle_suma_lo_que_muestra_la_grilla(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.asentar((self.banco, 500, 0), (self.clientes, 0, 500))
        self.asentar((self.clientes, 200, 0), (self.caja, 0, 200))   # devolución

        grilla = self.fila(self.reporte(), self.clientes)
        detalle = detalle_de_cuenta(self.empresa.id, DESDE, HASTA, self.clientes.id)

        self.assertEqual(len(detalle['filas']), 3)
        self.assertEqual(detalle['totales']['ingresos'], grilla['ingresos'])
        self.assertEqual(detalle['totales']['egresos'], grilla['egresos'])
        self.assertEqual(detalle['totales']['neto'], grilla['neto'])

    def test_saldo_corrido(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000),
                     fecha=timezone.datetime(2026, 8, 10).date())
        self.asentar((self.clientes, 200, 0), (self.caja, 0, 200),
                     fecha=timezone.datetime(2026, 8, 20).date())

        filas = detalle_de_cuenta(self.empresa.id, DESDE, HASTA, self.clientes.id)['filas']
        self.assertEqual([f['saldo'] for f in filas],
                         [Decimal("1000.00"), Decimal("800.00")])

    def test_filtro_por_medio(self):
        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.asentar((self.banco, 500, 0), (self.clientes, 0, 500))

        solo_banco = detalle_de_cuenta(self.empresa.id, DESDE, HASTA, self.clientes.id,
                                       medio='BCO')
        self.assertEqual(len(solo_banco['filas']), 1)
        self.assertEqual(solo_banco['totales']['ingresos'], Decimal("500.00"))


class VistasTest(BaseEOAF):
    """Fases 3 a 5: pantalla, drill-down y exportaciones."""

    def setUp(self):
        super().setUp()
        from django.test import Client
        self.http = Client()
        self.http.force_login(self.usuario)
        sesion = self.http.session
        sesion['empresa_id'] = self.empresa.id
        sesion['sucursal_id'] = self.sucursal.id
        sesion['ejercicio_id'] = self.ejercicio.id
        sesion.save()

        self.asentar((self.caja, 1000, 0), (self.clientes, 0, 1000))
        self.asentar((self.gastos_luz, 250, 0), (self.banco, 0, 250))

    @property
    def filtros(self):
        return {'desde': DESDE.isoformat(), 'hasta': HASTA.isoformat(), 'condic': ['1', '2']}

    def test_la_pantalla_no_autoejecuta_la_consulta(self):
        """Como el `cmdGenerar` del legado: se muestra el formulario, no el resultado."""
        from django.urls import reverse
        respuesta = self.http.get(reverse('eoaf_index'))

        self.assertEqual(respuesta.status_code, 200)
        contenido = respuesta.content.decode()
        self.assertIn('Generar', contenido)
        self.assertNotIn('DEUDORES POR VENTAS', contenido)

    def test_la_grilla_lista_las_cuentas_con_movimiento(self):
        from django.urls import reverse
        respuesta = self.http.get(reverse('eoaf_grilla'), self.filtros)

        self.assertEqual(respuesta.status_code, 200)
        contenido = respuesta.content.decode()
        self.assertIn('DEUDORES POR VENTAS', contenido)
        self.assertIn('ADM-ENERGIA', contenido)
        self.assertIn('1.000,00', contenido)   # formato es-AR

    def test_el_modal_de_una_cuenta_muestra_su_detalle(self):
        from django.urls import reverse
        respuesta = self.http.get(
            reverse('eoaf_cuenta_modal', args=[self.clientes.id]), self.filtros)

        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('DEUDORES POR VENTAS', respuesta.content.decode())

    def test_el_modal_rechaza_una_cuenta_de_otra_empresa(self):
        from django.urls import reverse
        otra = Empresa.objects.create(nombre="OTRA SA", cuit="30999888771",
                                      direccion="C 2", correo="o@t.com")
        ajena = Cuenta.objects.create(jerarquia="999", cuenta="AJENA", imputable=1,
                                      tipo="A", empresa=otra)
        respuesta = self.http.get(
            reverse('eoaf_cuenta_modal', args=[ajena.id]), self.filtros)

        self.assertIn('no pertenece a la empresa activa', respuesta.content.decode())

    def test_exportacion_a_excel(self):
        from io import BytesIO
        from django.urls import reverse
        import openpyxl

        respuesta = self.http.get(reverse('eoaf_excel'), self.filtros)
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])

        hoja = openpyxl.load_workbook(BytesIO(respuesta.content)).active
        encabezados = [c.value for c in hoja[3]]
        self.assertIn('Ingresos Fondos', encabezados)
        self.assertIn('Egresos Fondos', encabezados)
        self.assertIn('Flujo Neto', encabezados)
        self.assertNotIn('Disp.Inicial', encabezados)   # se quitó por decisión del usuario

        detalles = [hoja.cell(row=f, column=3).value for f in range(4, hoja.max_row + 1)]
        self.assertIn('DEUDORES POR VENTAS', detalles)
        self.assertIn('TOTALES', detalles)

    def test_exportacion_a_pdf(self):
        from django.urls import reverse
        respuesta = self.http.get(reverse('eoaf_pdf'), self.filtros)

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertTrue(respuesta.content.startswith(b'%PDF'))

    def test_las_exportaciones_respetan_el_filtro_de_condicion(self):
        from io import BytesIO
        from django.urls import reverse
        import openpyxl

        self.asentar((self.caja, 700, 0), (self.clientes, 0, 700), condic=2)

        respuesta = self.http.get(reverse('eoaf_excel'), {
            'desde': DESDE.isoformat(), 'hasta': HASTA.isoformat(), 'condic': ['1']})
        hoja = openpyxl.load_workbook(BytesIO(respuesta.content)).active

        fila_total = next(f for f in range(4, hoja.max_row + 1)
                          if hoja.cell(row=f, column=3).value == 'TOTALES')
        self.assertEqual(hoja.cell(row=fila_total, column=5).value, Decimal("1000.00"))

    def test_destildar_las_dos_condiciones_no_muestra_todo(self):
        """Un checkbox sin tildar no se envía, así que "ninguna condición" llega igual que
        "primera carga". El campo oculto `generado` distingue los dos casos: sin él, la grilla
        mostraría todo contradiciendo a los checkboxes."""
        from django.urls import reverse

        # Primera carga (nadie tocó el formulario): van las dos.
        primera = self.http.get(reverse('eoaf_grilla'), {
            'desde': DESDE.isoformat(), 'hasta': HASTA.isoformat()})
        self.assertIn('DEUDORES POR VENTAS', primera.content.decode())

        # Formulario enviado con las dos destildadas: no se muestra nada.
        destildado = self.http.get(reverse('eoaf_grilla'), {
            'desde': DESDE.isoformat(), 'hasta': HASTA.isoformat(), 'generado': '1'})
        contenido = destildado.content.decode()
        self.assertNotIn('DEUDORES POR VENTAS', contenido)
        self.assertIn('SIN CONDICIÓN', contenido)
