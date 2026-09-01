"""Exportaciones de la Caja Diaria: Excel y PDF.

El Excel replica el layout con el que el usuario ya trabaja (encabezado con N° de caja, fecha y
sucursal; grilla de movimientos; bloque de saldos al pie).

El PDF agrupa en INGRESOS / EGRESOS y, dentro de cada sección, por cuenta contable ordenada por
jerarquía, con UN subtotal por cuenta.
"""

from django.utils import timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from django.http import HttpResponse

from tesoreria.services.caja_diaria import agrupar_para_pdf

FORMATO_MONEDA = '#,##0.00'
AZUL = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
GRIS = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
BLANCO_NEGRITA = Font(bold=True, color='FFFFFF')
BORDE_FINO = Border(*(Side(style='thin', color='CBD5E1'),) * 4)


def _titulo_caja(sesion):
    fecha = sesion.fecha_operativa.strftime('%d/%m/%Y') if sesion.fecha_operativa else 'ABIERTA'
    return f"Caja N° {sesion.numero or sesion.id} de fecha {fecha} - Sucursal {sesion.caja.sucursal.nombre}"


def exportar_caja_diaria_excel(sesion, movimientos, saldos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Caja"

    ws['A1'] = _titulo_caja(sesion)
    ws['A1'].font = Font(size=12, bold=True)

    encabezados = [
        'id_asto', 'fecha', 'cod.', 'Razon', 'id_cta', 'Imputacion', 'Concepto',
        'efectivo', 'dolares', 'banco', 'valores', 'tarjetas', 'otros', 'total', 'saldo', 'cond',
    ]
    fila_encabezado = 3
    for columna, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=columna, value=titulo)
        celda.font = BLANCO_NEGRITA
        celda.fill = AZUL
        celda.alignment = Alignment(horizontal='center')
        celda.border = BORDE_FINO

    fila = fila_encabezado + 1
    for movimiento in movimientos:
        valores = [
            movimiento['asiento_id'], movimiento['fecha'], movimiento['codigo'], movimiento['razon'],
            movimiento['jerarquia'], movimiento['cuenta'], movimiento['descripcion'],
            movimiento['efectivo'], movimiento['dolares'], movimiento['banco'],
            movimiento['valores'], movimiento['tarjetas'], movimiento['otros'],
            movimiento['total'], movimiento['saldo'], movimiento['condic'],
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=columna, value=valor)
            if columna >= 8 and columna <= 15:
                celda.number_format = FORMATO_MONEDA
            if columna == 2 and movimiento['fecha']:
                celda.number_format = 'DD/MM/YYYY'
        fila += 1

    # Bloque de saldos
    fila += 2
    ws.cell(row=fila, column=8, value='Efectivo').font = Font(bold=True)
    ws.cell(row=fila, column=9, value='Dólares').font = Font(bold=True)
    ws.cell(row=fila, column=10, value='Banco').font = Font(bold=True)
    ws.cell(row=fila, column=11, value='Valores').font = Font(bold=True)
    ws.cell(row=fila, column=12, value='Tarjetas').font = Font(bold=True)
    ws.cell(row=fila, column=13, value='Otros').font = Font(bold=True)
    ws.cell(row=fila, column=14, value='Neto').font = Font(bold=True)

    etiquetas = (
        ('Saldos Iniciales', saldos['inicial']),
        ('Movimientos', saldos['movimiento']),
        ('Saldos Finales', saldos['final']),
    )
    for etiqueta, bloque in etiquetas:
        fila += 1
        celda = ws.cell(row=fila, column=7, value=etiqueta)
        celda.font = Font(bold=True)
        celda.alignment = Alignment(horizontal='right')
        for columna, clave in ((8, 'efectivo'), (9, 'dolares'), (10, 'banco'),
                               (11, 'valores'), (12, 'tarjetas'), (13, 'otros'), (14, 'neto')):
            celda = ws.cell(row=fila, column=columna, value=bloque.get(clave))
            celda.number_format = FORMATO_MONEDA
            celda.fill = GRIS

    fila += 1
    ws.cell(row=fila, column=7, value='Ingresos / Egresos').font = Font(bold=True)
    ws.cell(row=fila, column=8, value=saldos['movimiento']['ingresos']).number_format = FORMATO_MONEDA
    ws.cell(row=fila, column=9, value=saldos['movimiento']['egresos']).number_format = FORMATO_MONEDA

    anchos = [10, 12, 8, 32, 10, 30, 34, 15, 15, 15, 15, 15, 15, 16, 16, 7]
    for indice, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(indice)].width = ancho

    nombre = f"caja_diaria_{sesion.numero or sesion.id}_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    wb.save(respuesta)
    return respuesta


def exportar_caja_diaria_pdf(sesion, movimientos, saldos):
    from io import BytesIO
    from django.template.loader import render_to_string
    from xhtml2pdf import pisa

    contexto = {
        'sesion': sesion,
        'titulo': _titulo_caja(sesion),
        'secciones': agrupar_para_pdf(movimientos),
        'saldos': saldos,
        'fecha_emision': timezone.localtime(),
    }
    html = render_to_string('tesoreria/pdf/caja_diaria_pdf.html', contexto)
    resultado = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode('utf-8')), resultado)

    if pdf.err:
        return HttpResponse('Error al generar el PDF', status=400)

    nombre = f"caja_diaria_{sesion.numero or sesion.id}_{timezone.localdate().strftime('%Y%m%d')}.pdf"
    respuesta = HttpResponse(resultado.getvalue(), content_type='application/pdf')
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return respuesta
