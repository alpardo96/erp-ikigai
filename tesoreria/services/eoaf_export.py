"""Exportaciones del Estado de Origen y Aplicación de Fondos: Excel y PDF (Plan 050, fase 5).

El layout del Excel sigue al del sistema VFP para que las planillas históricas del usuario sigan
siendo comparables, **menos la columna `Disp.Inicial`**, que se quitó por decisión del usuario:
el reporte mide el flujo entre dos fechas y en el legado esa columna además estaba mal calculada
(las imputables la sumaban al flujo neto y las sumarizadoras no).
"""

from io import BytesIO

import openpyxl
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATO_MONEDA = '#,##0.00'
AZUL = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
BLANCO_NEGRITA = Font(bold=True, color='FFFFFF')
# Las sumarizadoras van en azul y negrita, como en el legado (`font.colorindex = 5`).
AZUL_NEGRITA = Font(bold=True, color='1D4ED8')


def _titulo(desde, hasta, condics):
    rotulos = {1: 'Real', 2: 'Presupuestado'}
    condicion = ' + '.join(rotulos[c] for c in sorted(condics)) or 'sin condición'
    return (f"Estado de Origen y Aplicación de Fondos — "
            f"del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')} ({condicion})")


def _nombre_archivo(desde, hasta, extension):
    return (f"origen_aplicacion_fondos_{desde.strftime('%Y%m%d')}_"
            f"{hasta.strftime('%Y%m%d')}.{extension}")


def exportar_eoaf_excel(datos, medios):
    """`datos` es lo que devuelve `estado_origen_aplicacion_fondos()`."""
    desde, hasta = datos['periodo']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Origen y Aplicacion"

    ws['A1'] = _titulo(desde, hasta, datos['condics'])
    ws['A1'].font = Font(size=12, bold=True)

    encabezados = (['Codigo', 'Jerarquia', 'Detalle', 'Imp',
                    'Ingresos Fondos', 'Egresos Fondos', 'Flujo Neto']
                   + [rotulo for _, rotulo in medios])
    fila_encabezado = 3
    for columna, titulo in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=columna, value=titulo)
        celda.font = BLANCO_NEGRITA
        celda.fill = AZUL
        celda.alignment = Alignment(horizontal='center')

    fila = fila_encabezado
    for item in datos['filas']:
        fila += 1
        valores = ([item['codigo'], item['jerarquia'], item['detalle'],
                    1 if item['imputable'] else 0,
                    item['ingresos'], item['egresos'], item['neto']]
                   + [item['medios'][clave] for clave, _ in medios])
        for columna, valor in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=columna, value=valor)
            if columna >= 5:
                celda.number_format = FORMATO_MONEDA
            if not item['imputable']:
                celda.font = AZUL_NEGRITA

    # Totales: sólo imputables (las sumarizadoras ya las contienen).
    totales = datos['totales']
    fila += 2
    celda = ws.cell(row=fila, column=3, value='TOTALES')
    celda.font = BLANCO_NEGRITA
    celda.fill = AZUL
    valores = ([totales['ingresos'], totales['egresos'], totales['neto']]
               + [totales['medios'][clave] for clave, _ in medios])
    for indice, valor in enumerate(valores):
        celda = ws.cell(row=fila, column=5 + indice, value=valor)
        celda.number_format = FORMATO_MONEDA
        celda.font = BLANCO_NEGRITA
        celda.fill = AZUL

    anchos = [10, 12, 42, 6] + [16] * (3 + len(medios))
    for indice, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = ancho
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=5)

    respuesta = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = (
        f'attachment; filename="{_nombre_archivo(desde, hasta, "xlsx")}"')
    wb.save(respuesta)
    return respuesta


def exportar_eoaf_pdf(datos, medios):
    from xhtml2pdf import pisa

    desde, hasta = datos['periodo']
    filas = [
        dict(item, medios_lista=[item['medios'][clave] for clave, _ in medios])
        for item in datos['filas']
    ]
    totales = dict(datos['totales'],
                   medios_lista=[datos['totales']['medios'][clave] for clave, _ in medios])

    html = render_to_string('tesoreria/pdf/eoaf_pdf.html', {
        'titulo': _titulo(desde, hasta, datos['condics']),
        'filas': filas,
        'totales': totales,
        'medios': medios,
        'fecha_emision': timezone.localtime(),
    })

    resultado = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode('utf-8')), resultado)
    if pdf.err:
        return HttpResponse('Error al generar el PDF', status=400)

    respuesta = HttpResponse(resultado.getvalue(), content_type='application/pdf')
    respuesta['Content-Disposition'] = (
        f'attachment; filename="{_nombre_archivo(desde, hasta, "pdf")}"')
    return respuesta
