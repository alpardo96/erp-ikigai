import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from decimal import Decimal
from datetime import datetime

def format_arg(val, decimals=2):
    if val is None: return "0,00"
    if isinstance(val, (int, float, str)): val = Decimal(str(val))
    return f"{val:,.{decimals}f}".replace(',', 'TEMP').replace('.', ',').replace('TEMP', '.')

def generar_excel_valores(valores):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valores de Terceros"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    headers = [
        "Vencimiento", "Banco", "Número", "Importe", "Estado", 
        "Origen", "Destino", "Asiento Rec", "Asiento Ent"
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for v in valores:
        origen = getattr(v, 'origen_razon_social', '') or v.nombre_firmante or ''
        destino = getattr(v, 'destino_razon_social', '') or '-'
        
        ws.append([
            v.fecha_vencimiento.strftime('%d/%m/%Y') if v.fecha_vencimiento else '',
            v.banco.nombre if v.banco else '',
            v.numero_cheque,
            float(v.importe),
            v.get_estado_display(),
            origen,
            destino,
            v.asiento_recepcion_id or '',
            v.asiento_entrega_id or ''
        ])

    # Format importe column
    for row in range(2, len(valores) + 2):
        ws.cell(row=row, column=4).number_format = '#,##0.00'

    # Adjust column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 30
    ws.column_dimensions['H'].width = 12
    ws.column_dimensions['I'].width = 12

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

class PDFBuilder:
    def __init__(self, buffer):
        self.c = canvas.Canvas(buffer, pagesize=landscape(A4))
        self.width, self.height = landscape(A4)
        self.c.setLineWidth(0.5)

    def draw_text(self, text, x, y, font_name="Helvetica", font_size=9, align="left"):
        self.c.setFont(font_name, font_size)
        y_rl = self.height - y
        if align == "left":
            self.c.drawString(x, y_rl, str(text))
        elif align == "right":
            self.c.drawRightString(x, y_rl, str(text))
        elif align == "center":
            self.c.drawCentredString(x, y_rl, str(text))

    def draw_line(self, x1, y1, x2, y2):
        self.c.line(x1, self.height - y1, x2, self.height - y2)

def generar_pdf_valores(valores, empresa):
    buffer = io.BytesIO()
    pdf = PDFBuilder(buffer)
    
    y = 40
    pdf.draw_text(empresa.razon_social, 30, y, "Helvetica-Bold", 14)
    pdf.draw_text(f"Reporte de Valores de Terceros - {datetime.now().strftime('%d/%m/%Y %H:%M')}", 30, y + 20, "Helvetica", 10)
    
    y += 40
    pdf.draw_line(30, y, pdf.width - 30, y)
    y += 15
    
    # Headers
    pdf.draw_text("Venc.", 30, y, "Helvetica-Bold", 9)
    pdf.draw_text("Banco", 90, y, "Helvetica-Bold", 9)
    pdf.draw_text("Número", 200, y, "Helvetica-Bold", 9)
    pdf.draw_text("Origen", 280, y, "Helvetica-Bold", 9)
    pdf.draw_text("Destino", 460, y, "Helvetica-Bold", 9)
    pdf.draw_text("Estado", 640, y, "Helvetica-Bold", 9)
    pdf.draw_text("Importe", pdf.width - 30, y, "Helvetica-Bold", 9, align="right")
    
    y += 5
    pdf.draw_line(30, y, pdf.width - 30, y)
    y += 15
    
    total = Decimal('0.00')
    
    for v in valores:
        if y > pdf.height - 40:
            pdf.c.showPage()
            y = 40
            
        origen = getattr(v, 'origen_razon_social', '') or v.nombre_firmante or ''
        destino = getattr(v, 'destino_razon_social', '') or '-'
        
        pdf.draw_text(v.fecha_vencimiento.strftime('%d/%m/%y') if v.fecha_vencimiento else '', 30, y, "Helvetica", 8)
        pdf.draw_text((v.banco.nombre if v.banco else '')[:18], 90, y, "Helvetica", 8)
        pdf.draw_text(v.numero_cheque, 200, y, "Helvetica", 8)
        pdf.draw_text(origen[:30], 280, y, "Helvetica", 8)
        pdf.draw_text(destino[:30], 460, y, "Helvetica", 8)
        pdf.draw_text(v.get_estado_display(), 640, y, "Helvetica", 8)
        pdf.draw_text(format_arg(v.importe), pdf.width - 30, y, "Helvetica", 8, align="right")
        
        total += v.importe
        y += 15
        
    y += 5
    pdf.draw_line(30, y, pdf.width - 30, y)
    y += 15
    pdf.draw_text("TOTAL:", pdf.width - 100, y, "Helvetica-Bold", 10)
    pdf.draw_text(format_arg(total), pdf.width - 30, y, "Helvetica-Bold", 10, align="right")
    
    pdf.c.save()
    buffer.seek(0)
    return buffer
