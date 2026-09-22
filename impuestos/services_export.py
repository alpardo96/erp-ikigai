import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any, List, Dict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from django.template.loader import render_to_string
from xhtml2pdf import pisa

from contable.models import LibroIvaVentas, LibroIvaCompras, LibroIvaAlic
from empresas.models import Empresa


def _dar_estilo_celda_encabezado(cell, color_hex="1E293B"):
    """Aplica formato visual corporativo a los encabezados de tabla."""
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=color_hex)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _dar_estilo_celda_total(cell, color_bg="F1F5F9"):
    """Aplica formato visual destacado a la fila de totales."""
    cell.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    cell.fill = PatternFill("solid", fgColor=color_bg)
    border_top = Border(
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="0F172A"),
    )
    cell.border = border_top


def exportar_libro_iva_excel(tipo: str, empresa: Empresa, anio: int, mes: int) -> HttpResponse:
    """
    Genera y retorna la planilla Excel (.xlsx) del Libro IVA Ventas o Compras
    con encabezado formal (Razón Social y CUIT), Asiento ID, discriminación de alícuotas
    (2.5%, 5%, 10.5%, 21%, 27%), IVA Computable / Débito Fiscal y totales contables.
    """
    periodo_yyyymm = f"{anio}{mes:02d}"
    es_ventas = tipo.upper() == 'VENTAS'
    c_v = 'V' if es_ventas else 'C'
    titulo_libro = "LIBRO IVA VENTAS" if es_ventas else "LIBRO IVA COMPRAS"
    col_sujeto = "Cliente / Razón Social" if es_ventas else "Proveedor / Razón Social"
    col_iva_comp = "IVA Débito Fiscal" if es_ventas else "IVA Computable"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Libro IVA"
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A6"
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    # 1. Encabezado de la Empresa y Período
    empresa_razon = empresa.nombre.upper() if empresa else "EMPRESA"
    empresa_cuit = f"CUIT: {empresa.cuit}" if empresa and empresa.cuit else ""

    # Fila 1: Razón Social y CUIT
    ws.merge_cells("A1:Q1")
    ws["A1"] = f"{empresa_razon}   {empresa_cuit}".strip()
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # Fila 2: Título del Libro
    ws.merge_cells("A2:Q2")
    ws["A2"] = titulo_libro
    color_titulo = "2563EB" if es_ventas else "4338CA"
    ws["A2"].font = Font(name="Calibri", size=13, bold=True, color=color_titulo)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Fila 3: Período y Fecha de Emisión
    ws.merge_cells("A3:Q3")
    fecha_emision = timezone.localtime().strftime("%d/%m/%Y %H:%M")
    ws["A3"] = f"Período Fiscal: {mes:02d}/{anio} ({periodo_yyyymm})   |   Emisión: {fecha_emision}"
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    # 2. Encabezados de Columnas
    headers = [
        "Fecha",
        "Cód.",
        "Comprobante",
        "Asiento",
        col_sujeto,
        "CUIT",
        "Neto Gravado",
        "Exento",
        "No Gravado",
        "IVA 2.5%",
        "IVA 5%",
        "IVA 10.5%",
        "IVA 21%",
        "IVA 27%",
        col_iva_comp,
        "Otros (Ret/Perc)",
        "Total",
    ]

    header_color = "1E3A8A" if es_ventas else "312E81"
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx, value=header_text)
        _dar_estilo_celda_encabezado(cell, header_color)

    # 3. Consulta de Comprobantes y Alícuotas
    if es_ventas:
        qs = LibroIvaVentas.objects.filter(
            empresa=empresa,
            periodo=periodo_yyyymm
        ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
        if not qs.exists():
            qs = LibroIvaVentas.objects.filter(
                empresa=empresa,
                fecha__year=anio,
                fecha__month=mes
            ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
    else:
        qs = LibroIvaCompras.objects.filter(
            empresa=empresa,
            periodo=periodo_yyyymm
        ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
        if not qs.exists():
            qs = LibroIvaCompras.objects.filter(
                empresa=empresa,
                fecha__year=anio,
                fecha__month=mes
            ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')

    # Pre-cargar alícuotas indexadas por asiento_id
    asientos_ids = [c.asiento_id for c in qs if c.asiento_id]
    alicuotas_qs = LibroIvaAlic.objects.filter(asiento_id__in=asientos_ids, c_v=c_v)
    alicuotas_map: Dict[int, List[LibroIvaAlic]] = {}
    for al in alicuotas_qs:
        alicuotas_map.setdefault(al.asiento_id, []).append(al)

    # 4. Volcado de Filas de Comprobantes
    row_num = 6
    border_fila = Border(
        bottom=Side(style="thin", color="E2E8F0"),
        left=Side(style="thin", color="F1F5F9"),
        right=Side(style="thin", color="F1F5F9"),
    )

    for c in qs:
        # Calcular desglose de IVA por alícuota
        iva25 = Decimal('0.00')
        iva5 = Decimal('0.00')
        iva105 = Decimal('0.00')
        iva21 = Decimal('0.00')
        iva27 = Decimal('0.00')
        iva_computable = Decimal('0.00')

        alics_del_asiento = alicuotas_map.get(c.asiento_id, [])
        if alics_del_asiento:
            for al in alics_del_asiento:
                alic_dec = Decimal(str(al.alicuota))
                iva_monto = Decimal(str(al.iva or 0))
                comp_monto = Decimal(str(al.computable or al.iva or 0))
                iva_computable += comp_monto

                if alic_dec == Decimal('2.5') or alic_dec == Decimal('2.50'):
                    iva25 += iva_monto
                elif alic_dec == Decimal('5') or alic_dec == Decimal('5.00'):
                    iva5 += iva_monto
                elif alic_dec == Decimal('10.5') or alic_dec == Decimal('10.50'):
                    iva105 += iva_monto
                elif alic_dec == Decimal('21') or alic_dec == Decimal('21.00'):
                    iva21 += iva_monto
                elif alic_dec == Decimal('27') or alic_dec == Decimal('27.00'):
                    iva27 += iva_monto
                else:
                    # Si no coincide exactamente, sumamos a 21
                    iva21 += iva_monto
        else:
            # Fallback si no hubiese alícuota en tabla satélite
            iva21 = Decimal(str(c.iva_total or 0))
            iva_computable = Decimal(str(c.iva_total or 0))

        # Valores de fila
        fecha_str = c.fecha.strftime("%d/%m/%Y") if c.fecha else ""
        cbte_nro = f"{c.punto:05d}-{c.numero}"
        razon = c.clienteproveedor.razon_social if c.clienteproveedor else ""
        cuit = c.cuit or (c.clienteproveedor.cuit if c.clienteproveedor else "") or ""

        row_vals = [
            fecha_str,
            str(c.codiva),
            cbte_nro,
            c.asiento_id,
            razon,
            cuit,
            float(c.neto_gravado or 0),
            float(c.exento or 0),
            float(c.no_gravado or 0),
            float(iva25),
            float(iva5),
            float(iva105),
            float(iva21),
            float(iva27),
            float(iva_computable),
            float(c.otros or 0),
            float(c.total or 0),
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = border_fila

            if col_idx in (1, 2, 3, 4, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                # Columnas numéricas / importes
                cell.number_format = "$ #,##0.00;($ #,##0.00);\"$ -\""
                cell.alignment = Alignment(horizontal="right", vertical="center")

        row_num += 1

    # 5. Fila de Totales
    last_data_row = row_num - 1
    if last_data_row >= 6:
        ws.cell(row=row_num, column=1, value="")
        ws.cell(row=row_num, column=2, value="")
        ws.cell(row=row_num, column=3, value="")
        ws.cell(row=row_num, column=4, value="")
        ws.cell(row=row_num, column=5, value="TOTALES DEL PERÍODO")
        ws.cell(row=row_num, column=6, value="")

        for col_idx in range(1, 7):
            _dar_estilo_celda_total(ws.cell(row=row_num, column=col_idx))

        # Fórmulas de suma para columnas numéricas 7 a 17
        for col_idx in range(7, 18):
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=row_num, column=col_idx)
            cell.value = f"=SUM({col_letter}6:{col_letter}{last_data_row})"
            cell.number_format = "$ #,##0.00;($ #,##0.00);\"$ -\""
            cell.alignment = Alignment(horizontal="right", vertical="center")
            _dar_estilo_celda_total(cell)

    # 6. Ajuste automático del ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Ignorar filas de títulos fusionadas 1, 2, 3
            if cell.row in (1, 2, 3):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Respuesta HTTP
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{titulo_libro.replace(' ', '_')}_{periodo_yyyymm}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def exportar_libro_iva_pdf(tipo: str, empresa: Empresa, anio: int, mes: int) -> HttpResponse:
    """
    Genera y retorna el archivo PDF del Libro IVA Ventas o Compras
    en formato A4 horizontal (Landscape) con membrete oficial, Razón Social, CUIT,
    totales y paginado.
    """
    periodo_yyyymm = f"{anio}{mes:02d}"
    es_ventas = tipo.upper() == 'VENTAS'
    titulo_libro = "LIBRO IVA VENTAS" if es_ventas else "LIBRO IVA COMPRAS"

    if es_ventas:
        comprobantes = LibroIvaVentas.objects.filter(
            empresa=empresa,
            periodo=periodo_yyyymm
        ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
        if not comprobantes.exists():
            comprobantes = LibroIvaVentas.objects.filter(
                empresa=empresa,
                fecha__year=anio,
                fecha__month=mes
            ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
    else:
        comprobantes = LibroIvaCompras.objects.filter(
            empresa=empresa,
            periodo=periodo_yyyymm
        ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')
        if not comprobantes.exists():
            comprobantes = LibroIvaCompras.objects.filter(
                empresa=empresa,
                fecha__year=anio,
                fecha__month=mes
            ).select_related('clienteproveedor').order_by('fecha', 'punto', 'numero')

    tot_neto = sum((Decimal(str(c.neto_gravado or 0)) for c in comprobantes), Decimal('0.00'))
    tot_exento = sum((Decimal(str(c.exento or 0)) for c in comprobantes), Decimal('0.00'))
    tot_no_gravado = sum((Decimal(str(c.no_gravado or 0)) for c in comprobantes), Decimal('0.00'))
    tot_iva = sum((Decimal(str(c.iva_total or 0)) for c in comprobantes), Decimal('0.00'))
    tot_otros = sum((Decimal(str(c.otros or 0)) for c in comprobantes), Decimal('0.00'))
    tot_total = sum((Decimal(str(c.total or 0)) for c in comprobantes), Decimal('0.00'))

    context = {
        'empresa': empresa,
        'tipo': tipo.upper(),
        'es_ventas': es_ventas,
        'titulo_libro': titulo_libro,
        'periodo': periodo_yyyymm,
        'anio': anio,
        'mes': mes,
        'comprobantes': comprobantes,
        'tot_neto': tot_neto,
        'tot_exento': tot_exento,
        'tot_no_gravado': tot_no_gravado,
        'tot_iva': tot_iva,
        'tot_otros': tot_otros,
        'tot_total': tot_total,
        'fecha_emision': timezone.localtime(),
    }

    html_string = render_to_string('impuestos/pdf/libro_iva_pdf.html', context)
    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html_string.encode("utf-8")), result)

    if not pdf.err:
        filename = f"{titulo_libro.replace(' ', '_')}_{periodo_yyyymm}.pdf"
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    return HttpResponse("Error al generar el archivo PDF de Libro IVA", status=400)
