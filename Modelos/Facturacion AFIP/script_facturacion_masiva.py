import os
import csv
import sys
import ssl
import glob
import json
import base64
import atexit
import time
import requests
import itertools
from datetime import datetime

# ============================================================================
# PARCHE SSL PARA AFIP
# ============================================================================
orig_init_poolmanager = requests.adapters.HTTPAdapter.init_poolmanager
def patched_init_poolmanager(self, *args, **kwargs):
    context = ssl.create_default_context()
    context.set_ciphers('DEFAULT@SECLEVEL=1')
    kwargs['ssl_context'] = context
    return orig_init_poolmanager(self, *args, **kwargs)
requests.adapters.HTTPAdapter.init_poolmanager = patched_init_poolmanager

# ============================================================================
# RESOLUCIÓN DE RUTAS
# ============================================================================
def obtener_ruta_base():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

RUTA_BASE = obtener_ruta_base()

# ============================================================================
# DUPLICADOR DE CONSOLA (LOG)
# ============================================================================
class DuplicadorConsola:
    def __init__(self, stream_original, archivo_log):
        self.stream_original = stream_original
        self.archivo_log = archivo_log

    def write(self, mensaje):
        self.stream_original.write(mensaje)
        self.archivo_log.write(mensaje)
        self.archivo_log.flush() 

    def flush(self):
        self.stream_original.flush()
        self.archivo_log.flush()

sys.path.insert(0, RUTA_BASE)

import arca_arg.settings as confg
from arca_arg.webservice import ArcaWebService
from arca_arg.settings import WSDL_FEV1_HOM, WSDL_FEV1_PROD

def buscar_certificados():
    carpeta_certs = os.path.join(RUTA_BASE, 'certificados')
    if not os.path.isdir(carpeta_certs):
        raise FileNotFoundError(f"No se encontró la carpeta 'certificados' en: {carpeta_certs}")

    archivos_key = glob.glob(os.path.join(carpeta_certs, '*.key'))
    if len(archivos_key) == 0:
        raise FileNotFoundError(f"No se encontró ningún archivo .key en: {carpeta_certs}")
    if len(archivos_key) > 1:
        raise FileNotFoundError(f"Se encontraron múltiples archivos .key en: {carpeta_certs}")

    archivos_crt = glob.glob(os.path.join(carpeta_certs, '*.crt'))
    if len(archivos_crt) == 0:
        raise FileNotFoundError(f"No se encontró ningún archivo .crt en: {carpeta_certs}")
    if len(archivos_crt) > 1:
        raise FileNotFoundError(f"Se encontraron múltiples archivos .crt en: {carpeta_certs}")

    ruta_key = archivos_key[0]
    ruta_crt = archivos_crt[0]
    print(f"  Clave privada encontrada: {os.path.basename(ruta_key)}")
    print(f"  Certificado encontrado:   {os.path.basename(ruta_crt)}")
    return ruta_key, ruta_crt

def inicializar_wsfe():
    print("\n--- Inicializando conexión con AFIP ---")
    ruta_key, ruta_crt = buscar_certificados()

    ruta_cuit = os.path.join(RUTA_BASE, 'cuitemisor.txt')
    cuit_detectado = None
    if os.path.exists(ruta_cuit):
        try:
            with open(ruta_cuit, 'r', encoding='utf-8') as f:
                lineas = f.readlines()
                lineas_validas = [l.strip() for l in lineas if l.strip() and not l.strip().startswith('#')]
                if lineas_validas:
                    cuit_detectado = lineas_validas[0]
        except Exception as e:
            print(f"  Advertencia al leer cuitemisor.txt: {e}")

    if cuit_detectado:
        confg.CUIT = cuit_detectado
        print(f"  CUIT cargado dinámicamente desde cuitemisor.txt: {confg.CUIT}")
    else:
        confg.CUIT = '30718098226'
        print(f"  Usando CUIT por defecto (respaldo/hardcodeado): {confg.CUIT}")

    confg.PRIVATE_KEY_PATH = ruta_key
    confg.CERT_PATH = ruta_crt
    confg.TA_FILES_PATH = os.path.join(RUTA_BASE, '')
    confg.PROD = True  # True = Producción | False = Homologación (pruebas)

    wsdl = WSDL_FEV1_PROD if confg.PROD else WSDL_FEV1_HOM
    wsfe = ArcaWebService(wsdl, 'wsfe')
    print("  Conexión establecida correctamente.\n")
    return wsfe

def facturar_lote(wsfe, lote_facturas):
    """
    Recibe el objeto wsfe y una LISTA de diccionarios con los datos a facturar.
    Se asume que todas las facturas en la lista comparten 'pto_vta' y 'cbte_tipo'.
    Devuelve la misma lista pero con los campos de resultado completados.
    """
    def to_float(val): return float(val) if val and str(val).strip() != '' else 0.0
    def to_int(val): return int(val) if val and str(val).strip() != '' else 0

    if not lote_facturas:
        return []

    # Extraemos el pto_vta y cbte_tipo del primer elemento
    pto_vta = to_int(lote_facturas[0]['pto_vta'])
    cbte_tipo = to_int(lote_facturas[0]['cbte_tipo'])

    try:
        # 1. Consultar el número del último comprobante emitido SÓLO UNA VEZ
        ultimo_nro_req = {
            'Auth': {'Token': wsfe.token, 'Sign': wsfe.sign, 'Cuit': wsfe.cuit},
            'PtoVta': pto_vta,
            'CbteTipo': cbte_tipo
        }
        
        res_ultimo = None
        retries = 3
        while retries > 0:
            try:
                res_ultimo = wsfe.send_request('FECompUltimoAutorizado', ultimo_nro_req)
                break
            except Exception as e:
                retries -= 1
                if retries == 0:
                    raise Exception(f"Fallo crítico consultando FECompUltimoAutorizado: {e}")
                print(f"    [!] Error de red al consultar último comprobante. Reintentando en 5s... ({e})")
                time.sleep(5)

        if hasattr(res_ultimo, 'FECompUltimoAutorizadoResult') and res_ultimo.FECompUltimoAutorizadoResult:
            res_ult = res_ultimo.FECompUltimoAutorizadoResult
        else:
            res_ult = res_ultimo

        if hasattr(res_ult, 'CbteNro'):
            ultimo_nro_afip = res_ult.CbteNro
        else:
            ultimo_nro_afip = 0  

        # 2. Armar el array de detalles para el lote
        detalles = []
        for index, datos_factura in enumerate(lote_facturas):
            concepto = to_int(datos_factura['concepto'])
            nuevo_nro = ultimo_nro_afip + 1 + index

            detalle = {
                'Concepto': concepto,
                'DocTipo': to_int(datos_factura['doc_tipo']),
                'DocNro': to_int(str(datos_factura['doc_nro']).replace('-', '').replace(' ', '')),
                'CbteDesde': nuevo_nro,
                'CbteHasta': nuevo_nro,
                'CbteFch': datos_factura['cbte_fch'],
                'ImpTotal': to_float(datos_factura['imp_total']),
                'ImpTotConc': to_float(datos_factura['imp_tot_conc']),
                'ImpNeto': to_float(datos_factura['imp_neto']),
                'ImpOpEx': to_float(datos_factura['imp_op_ex']),
                'ImpTrib': to_float(datos_factura['imp_trib']),
                'ImpIVA': to_float(datos_factura['imp_iva']),
                'MonId': 'PES',
                'MonCotiz': 1
            }

            # Condición IVA Receptor (RG 5616)
            if 'condicion_iva_receptor_id' in datos_factura and str(datos_factura['condicion_iva_receptor_id']).strip():
                detalle['CondicionIVAReceptorId'] = to_int(datos_factura['condicion_iva_receptor_id'])

            if concepto in (2, 3):
                detalle['FchServDesde'] = datos_factura['fch_serv_desde']
                detalle['FchServHasta'] = datos_factura['fch_serv_hasta']
                detalle['FchVtoPago'] = datos_factura['fch_vto_pago']

            if cbte_tipo in (2, 3, 7, 8, 12, 13):
                cbte_asoc_tipo = to_int(datos_factura.get('cbte_asoc_tipo', 0))
                cbte_asoc_pto_vta = to_int(datos_factura.get('cbte_asoc_pto_vta', 0))
                cbte_asoc_nro = to_int(datos_factura.get('cbte_asoc_nro', 0))

                if cbte_asoc_tipo > 0 and cbte_asoc_pto_vta > 0 and cbte_asoc_nro > 0:
                    detalle['CbtesAsoc'] = {
                        'CbteAsoc': [{
                            'Tipo': cbte_asoc_tipo,
                            'PtoVta': cbte_asoc_pto_vta,
                            'Nro': cbte_asoc_nro,
                            'Cuit': str(wsfe.cuit)
                        }]
                    }
                else:
                    raise ValueError(f"ID {datos_factura.get('id_interno', index)}: Faltan datos del comprobante asociado requeridos para N/C o N/D.")

            imp_iva = to_float(datos_factura['imp_iva'])
            if imp_iva > 0 and datos_factura['id_iva']:
                detalle['Iva'] = {
                    'AlicIva': [{
                        'Id': to_int(datos_factura['id_iva']),
                        'BaseImp': to_float(datos_factura['base_imp']),
                        'Importe': imp_iva
                    }]
                }
            
            detalles.append(detalle)

        cabecera = {'CantReg': len(detalles), 'PtoVta': pto_vta, 'CbteTipo': cbte_tipo}
        data = {
            'Auth': {'Token': wsfe.token, 'Sign': wsfe.sign, 'Cuit': wsfe.cuit},
            'FeCAEReq': {
                'FeCabReq': cabecera,
                'FeDetReq': {'FECAEDetRequest': detalles}
            }
        }

        # 3. Enviar solicitud de lote a AFIP con reintentos
        resultado = None
        retries = 3
        while retries > 0:
            try:
                resultado = wsfe.send_request('FECAESolicitar', data)
                break
            except Exception as e:
                retries -= 1
                if retries == 0:
                    raise Exception(f"Fallo crítico al solicitar CAE tras reintentos: {e}")
                print(f"    [!] Error de red al solicitar CAE del lote. Reintentando en 10s... ({e})")
                time.sleep(10)

        # 4. Analizar respuesta
        if hasattr(resultado, 'FECAESolicitarResult'):
            res = resultado.FECAESolicitarResult
        elif hasattr(resultado, 'FeDetResp') or hasattr(resultado, 'Errors'):
            res = resultado
        else:
            raise ValueError("Respuesta inesperada de AFIP.")

        if hasattr(res, 'Errors') and res.Errors:
            error_msg = "; ".join([f"Err {err.Code}: {err.Msg}" for err in res.Errors.Err])
            raise ValueError(f"Error general AFIP en el lote: {error_msg}")

        if hasattr(res, 'FeDetResp') and res.FeDetResp:
            det_responses = res.FeDetResp.FECAEDetResponse
            if not isinstance(det_responses, list):
                det_responses = [det_responses]

            for i, det_resp in enumerate(det_responses):
                fila = lote_facturas[i]
                if det_resp.Resultado == 'A':
                    fch = str(fila['cbte_fch'])
                    fch_formateada = f"{fch[:4]}-{fch[4:6]}-{fch[6:]}"
                    datos_qr = {
                        "ver": 1,
                        "fecha": fch_formateada,
                        "cuit": int(wsfe.cuit),
                        "ptoVta": int(fila['pto_vta']),
                        "tipoCmp": int(fila['cbte_tipo']),
                        "nroCmp": int(det_resp.CbteDesde),
                        "importe": float(fila['imp_total']),
                        "moneda": "PES",
                        "ctz": 1.0,
                        "tipoDocRec": int(fila['doc_tipo']),
                        "nroDocRec": int(fila['doc_nro']),
                        "tipoCodAut": "E",
                        "codAut": int(det_resp.CAE)
                    }
                    json_str = json.dumps(datos_qr)
                    b64_str = base64.b64encode(json_str.encode('utf-8')).decode('utf-8')
                    raw_cae_qr = f"https://www.afip.gob.ar/fe/qr/?p={b64_str}"

                    fila['cae'] = det_resp.CAE
                    fila['cae_vto'] = getattr(det_resp, 'CAEFchVto', '')
                    fila['cbte_nro'] = det_resp.CbteDesde
                    fila['raw_cae_qr'] = raw_cae_qr
                    fila['estado'] = 'FACTURADO'
                    fila['observaciones'] = ''
                else:
                    obs_msg = ""
                    if hasattr(det_resp, 'Observaciones') and det_resp.Observaciones:
                        obs_msg = "; ".join([f"Obs {obs.Code}: {obs.Msg}" for obs in det_resp.Observaciones.Obs])
                    
                    fila['cae'] = ''
                    fila['cae_vto'] = ''
                    fila['cbte_nro'] = det_resp.CbteDesde
                    fila['raw_cae_qr'] = ''
                    fila['estado'] = 'ERROR'
                    fila['observaciones'] = f"Comprobante rechazado. {obs_msg}"
        return lote_facturas

    except Exception as e:
        error_general = str(e)
        for fila in lote_facturas:
            fila['cae'] = ''
            fila['cae_vto'] = ''
            fila['cbte_nro'] = ''
            fila['raw_cae_qr'] = ''
            fila['estado'] = 'ERROR_CRITICO'
            fila['observaciones'] = error_general
        return lote_facturas


def guardar_resultados_xlsx(filas, campos, archivo_salida):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = Workbook()
        ws = wb.active
        ws.title = "Resultados Facturación"
        fuente_encabezado = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        relleno_encabezado = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
        alineacion_centro = Alignment(horizontal='center', vertical='center')
        borde_fino = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        relleno_exito = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        relleno_error = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

        for col_idx, campo in enumerate(campos, 1):
            celda = ws.cell(row=1, column=col_idx, value=campo)
            celda.font = fuente_encabezado
            celda.fill = relleno_encabezado
            celda.alignment = alineacion_centro
            celda.border = borde_fino

        for row_idx, fila in enumerate(filas, 2):
            for col_idx, campo in enumerate(campos, 1):
                valor = fila.get(campo, '')
                celda = ws.cell(row=row_idx, column=col_idx, value=valor)
                celda.border = borde_fino
                celda.alignment = Alignment(horizontal='center')

            estado = fila.get('estado', '')
            if estado == 'FACTURADO':
                for col_idx in range(1, len(campos) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = relleno_exito
            elif 'ERROR' in estado:
                for col_idx in range(1, len(campos) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = relleno_error

        for col_idx, campo in enumerate(campos, 1):
            ancho = max(len(str(campo)) + 4, 12)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = ancho

        ws.freeze_panes = 'A2'
        wb.save(archivo_salida)
        print(f"  XLSX generado: {os.path.basename(archivo_salida)}")

    except ImportError:
        print("  ADVERTENCIA: No se pudo generar el archivo .xlsx (falta openpyxl).")
    except Exception as e:
        print(f"  ERROR al generar XLSX: {e}")

# ============================================================================
# PROGRAMA PRINCIPAL
# ============================================================================
if __name__ == "__main__":
    carpeta_logs = os.path.join(RUTA_BASE, "logs")
    os.makedirs(carpeta_logs, exist_ok=True) 

    fecha_hora_log = datetime.now().strftime('%Y%m%d_%H%M%S')
    archivo_log_path = os.path.join(carpeta_logs, f"facturacion_masiva_{fecha_hora_log}.log")
    _log_file = open(archivo_log_path, 'w', encoding='utf-8')
    sys.stdout = DuplicadorConsola(sys.__stdout__, _log_file)
    sys.stderr = DuplicadorConsola(sys.__stderr__, _log_file)
    atexit.register(_log_file.close) 

    print("=" * 60)
    print("  FACTURACIÓN ELECTRÓNICA MASIVA AFIP - ARCA")
    print("  Optimizado para procesamiento por lotes.")
    print(f"  Log: {os.path.join('logs', os.path.basename(archivo_log_path))}")
    print("=" * 60)

    fecha_hoy = datetime.now().strftime('%Y%m%d')
    archivo_salida_csv = os.path.join(RUTA_BASE, f"facturacion_masiva_{fecha_hoy}.csv")
    archivo_salida_xlsx = os.path.join(RUTA_BASE, f"facturacion_masiva_{fecha_hoy}.xlsx")

    archivo_entrada_xlsx = os.path.join(RUTA_BASE, "plantilla_facturacion.xlsx")
    archivo_entrada_csv = os.path.join(RUTA_BASE, "plantilla_facturacion.csv")
    archivo_usado = None
    
    filas_a_procesar = []
    campos_entrada = []

    if os.path.exists(archivo_entrada_xlsx):
        archivo_usado = archivo_entrada_xlsx
        print(f"  Leyendo archivo Excel: {os.path.basename(archivo_usado)}")
        try:
            from openpyxl import load_workbook
            import datetime as dt_mod
            wb = load_workbook(archivo_usado, data_only=True)
            ws = wb.active
            campos_entrada = [str(cell.value).strip() for cell in ws[1] if cell.value is not None]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(row):  
                    fila = {}
                    for i, campo in enumerate(campos_entrada):
                        if i < len(row):
                            val = row[i]
                            if isinstance(val, dt_mod.datetime):
                                fila[campo] = val.strftime('%Y%m%d')
                            elif val is not None:
                                s_val = str(val).strip()
                                if s_val.endswith('.0'):
                                    s_val = s_val[:-2]
                                fila[campo] = s_val
                            else:
                                fila[campo] = ''
                        else:
                            fila[campo] = ''
                    filas_a_procesar.append(fila)
        except ImportError:
            print("  ERROR: Se encontró plantilla_facturacion.xlsx pero falta openpyxl.")
            input("\n  Presioná ENTER para salir...")
            sys.exit(1)
    elif os.path.exists(archivo_entrada_csv):
        archivo_usado = archivo_entrada_csv
        print(f"  Leyendo archivo CSV: {os.path.basename(archivo_usado)}")
        with open(archivo_usado, mode='r', encoding='utf-8') as f:
            primera_linea = f.readline()
            delimitador = ';' if ';' in primera_linea else ','
            f.seek(0)
            lector_csv = csv.DictReader(f, delimiter=delimitador)
            campos_entrada = lector_csv.fieldnames
            for fila in lector_csv:
                filas_a_procesar.append(fila)
    else:
        print(f"\n  ERROR: No se encontró la plantilla de facturación")
        input("\n  Presioná ENTER para salir...")
        sys.exit(1)

    # Filtrar vacíos
    filas_a_procesar = [f for f in filas_a_procesar if f.get('id_interno')]

    if not filas_a_procesar:
        print("  No hay filas válidas para procesar.")
        sys.exit(0)

    try:
        wsfe_cliente = inicializar_wsfe()
    except Exception as e:
        print(f"\n  ERROR AL INICIALIZAR:\n  {e}")
        input("\n  Presioná ENTER para salir...")
        sys.exit(1)

    campos_salida = campos_entrada + ['cae', 'cae_vto', 'cbte_nro', 'estado', 'observaciones', 'raw_cae_qr']
    
    # 1. Preparar archivo CSV y escribir cabeceras (Sobrescribe si existía de otra ejecución previa hoy)
    print(f"--- Preparando archivo de guardado en vivo: {os.path.basename(archivo_salida_csv)} ---")
    with open(archivo_salida_csv, mode='w', encoding='utf-8', newline='') as f_out:
        escritor = csv.DictWriter(f_out, fieldnames=campos_salida)
        escritor.writeheader()

    # 2. Agrupar filas por Punto de Venta y Tipo de Comprobante
    def to_int(val): return int(val) if val and str(val).strip() != '' else 0
    # Ordenar primero es obligatorio para itertools.groupby
    filas_a_procesar.sort(key=lambda x: (to_int(x.get('pto_vta')), to_int(x.get('cbte_tipo'))))

    lotes_agrupados = itertools.groupby(
        filas_a_procesar, 
        key=lambda x: (to_int(x.get('pto_vta')), to_int(x.get('cbte_tipo')))
    )

    CHUNK_SIZE = 100
    filas_procesadas = []

    # 3. Procesar lotes
    for (pto_vta, cbte_tipo), grupo_iter in lotes_agrupados:
        grupo_filas = list(grupo_iter)
        print(f"\n>> Procesando Punto Venta: {pto_vta} | Tipo Comprobante: {cbte_tipo} (Total: {len(grupo_filas)} facturas)")
        
        for i in range(0, len(grupo_filas), CHUNK_SIZE):
            chunk = grupo_filas[i:i + CHUNK_SIZE]
            print(f"\n  -> Enviando sub-lote a AFIP (facturas {i+1} a {min(i+CHUNK_SIZE, len(grupo_filas))}) ...")
            
            # Llama a AFIP con el chunk
            resultados_chunk = facturar_lote(wsfe_cliente, chunk)
            
            # Guadar EN VIVO
            with open(archivo_salida_csv, mode='a', encoding='utf-8', newline='') as f_out:
                escritor = csv.DictWriter(f_out, fieldnames=campos_salida)
                for fila_res in resultados_chunk:
                    # Imprimir solo info esencial para no saturar consola con 25k logs, pero sí ver progreso
                    if fila_res['estado'] == 'FACTURADO':
                        print(f"      ✓ ID {fila_res.get('id_interno')} | Cbte {fila_res.get('cbte_nro')} | CAE: {fila_res.get('cae')}")
                    else:
                        print(f"      ✗ ID {fila_res.get('id_interno')} | ERROR: {fila_res.get('observaciones')}")
                    
                    escritor.writerow(fila_res)
                    filas_procesadas.append(fila_res)

    print(f"\n--- Generando archivo XLSX final ---")
    guardar_resultados_xlsx(filas_procesadas, campos_salida, archivo_salida_xlsx)

    exitosas = sum(1 for f in filas_procesadas if f.get('estado') == 'FACTURADO')
    fallidas = sum(1 for f in filas_procesadas if f.get('estado') == 'ERROR')

    print(f"\n{'=' * 60}")
    print(f"  Proceso MASIVO finalizado.")
    print(f"  Comprobantes procesados en total: {len(filas_procesadas)}")
    print(f"    - Exitosas: {exitosas}")
    print(f"    - Fallidas: {fallidas}")
    print(f"  Archivos generados en: {RUTA_BASE}")
    print(f"{'=' * 60}")

    input("\n  Presioná ENTER para salir...")
