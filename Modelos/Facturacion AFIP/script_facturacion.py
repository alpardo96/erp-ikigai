import os
import csv
import sys
import ssl
import glob
import json
import base64
import atexit
import requests
from datetime import datetime

# ============================================================================
# PARCHE SSL PARA AFIP
# AFIP usa claves Diffie-Hellman cortas que Python rechaza por defecto.
# Interceptamos la creación de conexiones de requests para forzar SECLEVEL=1.
# ============================================================================
orig_init_poolmanager = requests.adapters.HTTPAdapter.init_poolmanager
def patched_init_poolmanager(self, *args, **kwargs):
    context = ssl.create_default_context()
    context.set_ciphers('DEFAULT@SECLEVEL=1')
    kwargs['ssl_context'] = context
    return orig_init_poolmanager(self, *args, **kwargs)
requests.adapters.HTTPAdapter.init_poolmanager = patched_init_poolmanager

# ============================================================================
# RESOLUCIÓN DE RUTAS
# Cuando se ejecuta como .exe (PyInstaller), la ruta base cambia.
# Usamos esta función para obtener siempre la ruta correcta.
# ============================================================================
def obtener_ruta_base():
    """
    Devuelve la ruta base del ejecutable o del script.
    Si se ejecuta como .exe empaquetado, devuelve la carpeta donde está el .exe.
    Si se ejecuta como .py, devuelve la carpeta donde está el script.
    """
    if getattr(sys, 'frozen', False):
        # Ejecutándose como .exe (PyInstaller)
        return os.path.dirname(sys.executable)
    else:
        # Ejecutándose como script .py
        return os.path.dirname(os.path.abspath(__file__))

RUTA_BASE = obtener_ruta_base()


# ============================================================================
# DUPLICADOR DE CONSOLA (LOG)
# Captura TODO lo que pasa por la terminal (prints, errores, tracebacks,
# respuestas de AFIP, etc.) y lo escribe simultáneamente en un archivo .log.
# ============================================================================
class DuplicadorConsola:
    """
    Clona la salida de la terminal escribiendo tanto en la consola
    como en un archivo de log al mismo tiempo.
    Captura stdout (salida normal) y stderr (errores/tracebacks).
    """
    def __init__(self, stream_original, archivo_log):
        self.stream_original = stream_original
        self.archivo_log = archivo_log

    def write(self, mensaje):
        self.stream_original.write(mensaje)
        self.archivo_log.write(mensaje)
        self.archivo_log.flush()  # Guardar inmediatamente en disco por si el programa se cierra inesperadamente

    def flush(self):
        self.stream_original.flush()
        self.archivo_log.flush()


# Agregar la ruta base al sys.path para que Python encuentre el paquete arca_arg
sys.path.insert(0, RUTA_BASE)

import arca_arg.settings as confg
from arca_arg.webservice import ArcaWebService
from arca_arg.settings import WSDL_FEV1_HOM, WSDL_FEV1_PROD


def buscar_certificados():
    """
    Busca automáticamente los archivos .key y .crt dentro de la carpeta 'certificados/'.
    No importa el nombre del archivo, solo la extensión.
    Devuelve una tupla (ruta_key, ruta_crt).
    Si no encuentra alguno, lanza un error descriptivo.
    """
    carpeta_certs = os.path.join(RUTA_BASE, 'certificados')

    if not os.path.isdir(carpeta_certs):
        raise FileNotFoundError(
            f"No se encontró la carpeta 'certificados' en: {carpeta_certs}\n"
            "Creá una carpeta llamada 'certificados' al lado del ejecutable y colocá allí tu .key y .crt"
        )

    # Buscar archivos .key
    archivos_key = glob.glob(os.path.join(carpeta_certs, '*.key'))
    if len(archivos_key) == 0:
        raise FileNotFoundError(
            f"No se encontró ningún archivo .key en: {carpeta_certs}\n"
            "Colocá tu clave privada (.key) en la carpeta 'certificados'."
        )
    if len(archivos_key) > 1:
        raise FileNotFoundError(
            f"Se encontraron múltiples archivos .key en: {carpeta_certs}\n"
            "Dejá solo UN archivo .key en la carpeta 'certificados'."
        )

    # Buscar archivos .crt
    archivos_crt = glob.glob(os.path.join(carpeta_certs, '*.crt'))
    if len(archivos_crt) == 0:
        raise FileNotFoundError(
            f"No se encontró ningún archivo .crt en: {carpeta_certs}\n"
            "Colocá tu certificado (.crt) en la carpeta 'certificados'."
        )
    if len(archivos_crt) > 1:
        raise FileNotFoundError(
            f"Se encontraron múltiples archivos .crt en: {carpeta_certs}\n"
            "Dejá solo UN archivo .crt en la carpeta 'certificados'."
        )

    ruta_key = archivos_key[0]
    ruta_crt = archivos_crt[0]

    print(f"  Clave privada encontrada: {os.path.basename(ruta_key)}")
    print(f"  Certificado encontrado:   {os.path.basename(ruta_crt)}")

    return ruta_key, ruta_crt


def inicializar_wsfe():
    """
    Inicializa el cliente del Web Service de Factura Electrónica (WSFE).
    Busca automáticamente los certificados por extensión en la carpeta 'certificados/'.
    """
    print("\n--- Inicializando conexión con AFIP ---")

    # Buscar certificados automáticamente
    ruta_key, ruta_crt = buscar_certificados()

    # #############################################
    #          Configuración de CUIT
    # #############################################
    # Intentamos cargar el CUIT dinámicamente desde 'cuitemisor.txt' en la misma carpeta.
    # Si el archivo no existe, está vacío o la línea empieza con '#' (comentada),
    # el programa no fallará y usará el CUIT predeterminado de respaldo.
    ruta_cuit = os.path.join(RUTA_BASE, 'cuitemisor.txt')
    cuit_detectado = None

    if os.path.exists(ruta_cuit):
        try:
            with open(ruta_cuit, 'r', encoding='utf-8') as f:
                # Leemos todas las líneas del archivo
                lineas = f.readlines()
                # Filtramos líneas que no estén vacías y que no comiencen con '#'
                lineas_validas = [l.strip() for l in lineas if l.strip() and not l.strip().startswith('#')]
                if lineas_validas:
                    cuit_detectado = lineas_validas[0]
        except Exception as e:
            print(f"  Advertencia al intentar leer cuitemisor.txt: {e}")

    # Asignamos el CUIT final
    if cuit_detectado:
        confg.CUIT = cuit_detectado
        print(f"  CUIT cargado dinámicamente desde cuitemisor.txt: {confg.CUIT}")
    else:
        confg.CUIT = '30718098226'  # CUIT de respaldo (hardcodeado)
        print(f"  Usando CUIT por defecto (respaldo/hardcodeado): {confg.CUIT}")

    confg.PRIVATE_KEY_PATH = ruta_key
    confg.CERT_PATH = ruta_crt
    confg.TA_FILES_PATH = os.path.join(RUTA_BASE, '')  # Guarda el ticket de acceso junto al .exe
    confg.PROD = False  # True = Producción | False = Homologación (pruebas)

    wsdl = WSDL_FEV1_PROD if confg.PROD else WSDL_FEV1_HOM
    wsfe = ArcaWebService(wsdl, 'wsfe')

    print("  Conexión establecida correctamente.\n")
    return wsfe


def facturar_comprobante(wsfe, datos_factura):
    """
    Recibe el objeto wsfe y un diccionario con los datos a facturar parseados del CSV.
    Devuelve un diccionario con el resultado (éxito/error, CAE, etc.)
    """
    def to_float(val): return float(val) if val and str(val).strip() != '' else 0.0
    def to_int(val): return int(val) if val and str(val).strip() != '' else 0

    try:
        pto_vta = to_int(datos_factura['pto_vta'])
        cbte_tipo = to_int(datos_factura['cbte_tipo'])
        concepto = to_int(datos_factura['concepto'])

        # 1. Consultar el número del último comprobante emitido
        ultimo_nro_req = {
            'Auth': {'Token': wsfe.token, 'Sign': wsfe.sign, 'Cuit': wsfe.cuit},
            'PtoVta': pto_vta,
            'CbteTipo': cbte_tipo
        }
        res_ultimo = wsfe.send_request('FECompUltimoAutorizado', ultimo_nro_req)

        # La respuesta puede venir envuelta o desenrollada por Zeep
        if hasattr(res_ultimo, 'FECompUltimoAutorizadoResult') and res_ultimo.FECompUltimoAutorizadoResult:
            res_ult = res_ultimo.FECompUltimoAutorizadoResult
        else:
            res_ult = res_ultimo

        if hasattr(res_ult, 'CbteNro'):
            nuevo_nro = res_ult.CbteNro + 1
        else:
            nuevo_nro = 1  # Si falla o es el primero

        # 2. Armar el diccionario para la solicitud
        cabecera = {'CantReg': 1, 'PtoVta': pto_vta, 'CbteTipo': cbte_tipo}
        detalle = {
            'Concepto': concepto,
            'DocTipo': to_int(datos_factura['doc_tipo']),
            'DocNro': int(datos_factura['doc_nro']),
            'CbteDesde': nuevo_nro,
            'CbteHasta': nuevo_nro,
            'CbteFch': datos_factura['cbte_fch'],
            'ImpTotal': to_float(datos_factura['imp_total']),
            'ImpTotConc': to_float(datos_factura['imp_tot_conc']),
            'ImpNeto': to_float(datos_factura['imp_neto']),
            'ImpOpEx': to_float(datos_factura['imp_op_ex']),
            'ImpTrib': to_float(datos_factura['imp_trib']),
            'ImpIVA': to_float(datos_factura['imp_iva']),
            'MonId': 'PES',
            'MonCotiz': 1
        }

        # Condición IVA Receptor (RG 5616)
        if 'condicion_iva_receptor_id' in datos_factura and str(datos_factura['condicion_iva_receptor_id']).strip():
            detalle['CondicionIVAReceptorId'] = to_int(datos_factura['condicion_iva_receptor_id'])

        # Fechas de servicio (obligatorias si el concepto es 2 o 3)
        if concepto in (2, 3):
            detalle['FchServDesde'] = datos_factura['fch_serv_desde']
            detalle['FchServHasta'] = datos_factura['fch_serv_hasta']
            detalle['FchVtoPago'] = datos_factura['fch_vto_pago']

        # Si es Nota de Crédito o Débito, requiere comprobante asociado
        # Tipos comunes: Débito (2, 7, 12) y Crédito (3, 8, 13)
        if cbte_tipo in (2, 3, 7, 8, 12, 13):
            cbte_asoc_tipo = to_int(datos_factura.get('cbte_asoc_tipo', 0))
            cbte_asoc_pto_vta = to_int(datos_factura.get('cbte_asoc_pto_vta', 0))
            cbte_asoc_nro = to_int(datos_factura.get('cbte_asoc_nro', 0))

            if cbte_asoc_tipo > 0 and cbte_asoc_pto_vta > 0 and cbte_asoc_nro > 0:
                detalle['CbtesAsoc'] = {
                    'CbteAsoc': [{
                        'Tipo': cbte_asoc_tipo,
                        'PtoVta': cbte_asoc_pto_vta,
                        'Nro': cbte_asoc_nro,
                        'Cuit': str(wsfe.cuit)
                    }]
                }
            else:
                return {
                    'exito': False, 'cae': '', 'cae_vto': '', 'cbte_nro': '',
                    'error': "Faltan datos del comprobante asociado (cbte_asoc_tipo, cbte_asoc_pto_vta, cbte_asoc_nro) requeridos para emitir Notas de Crédito/Débito."
                }

        # Alicuota IVA si aplica
        imp_iva = to_float(datos_factura['imp_iva'])
        if imp_iva > 0 and datos_factura['id_iva']:
            detalle['Iva'] = {
                'AlicIva': [{
                    'Id': to_int(datos_factura['id_iva']),
                    'BaseImp': to_float(datos_factura['base_imp']),
                    'Importe': imp_iva
                }]
            }

        data = {
            'Auth': {'Token': wsfe.token, 'Sign': wsfe.sign, 'Cuit': wsfe.cuit},
            'FeCAEReq': {
                'FeCabReq': cabecera,
                'FeDetReq': {'FECAEDetRequest': [detalle]}
            }
        }

        # 3. Enviar solicitud de autorización a AFIP
        resultado = wsfe.send_request('FECAESolicitar', data)

        # 4. Analizar respuesta de AFIP
        if hasattr(resultado, 'FECAESolicitarResult'):
            res = resultado.FECAESolicitarResult
        elif hasattr(resultado, 'FeDetResp') or hasattr(resultado, 'Errors'):
            res = resultado
        else:
            print(f"  ESTRUCTURA INESPERADA DE AFIP: {resultado}")
            return {
                'exito': False, 'cae': '', 'cae_vto': '', 'cbte_nro': '',
                'error': "Respuesta inesperada de AFIP."
            }

        # Si AFIP devuelve errores generales de validación
        if hasattr(res, 'Errors') and res.Errors:
            error_msg = "; ".join([f"Err {err.Code}: {err.Msg}" for err in res.Errors.Err])
            return {'exito': False, 'cae': '', 'cae_vto': '', 'cbte_nro': '', 'error': error_msg}

        # Analizamos la cabecera y el detalle devuelto
        if hasattr(res, 'FeDetResp') and res.FeDetResp:
            det_responses = res.FeDetResp.FECAEDetResponse
            det_resp = det_responses[0] if isinstance(det_responses, list) else det_responses

            # Si fue aprobado
            if det_resp.Resultado == 'A':
                # Armar URL de código QR de AFIP
                fch = str(datos_factura['cbte_fch'])
                fch_formateada = f"{fch[:4]}-{fch[4:6]}-{fch[6:]}"
                datos_qr = {
                    "ver": 1,
                    "fecha": fch_formateada,
                    "cuit": int(wsfe.cuit),
                    "ptoVta": int(datos_factura['pto_vta']),
                    "tipoCmp": int(datos_factura['cbte_tipo']),
                    "nroCmp": int(det_resp.CbteDesde),
                    "importe": float(datos_factura['imp_total']),
                    "moneda": "PES",
                    "ctz": 1.0,
                    "tipoDocRec": int(datos_factura['doc_tipo']),
                    "nroDocRec": int(datos_factura['doc_nro']),
                    "tipoCodAut": "E",
                    "codAut": int(det_resp.CAE)
                }
                json_str = json.dumps(datos_qr)
                b64_str = base64.b64encode(json_str.encode('utf-8')).decode('utf-8')
                raw_cae_qr = f"https://www.afip.gob.ar/fe/qr/?p={b64_str}"

                return {
                    'exito': True,
                    'cae': det_resp.CAE,
                    'cae_vto': getattr(det_resp, 'CAEFchVto', ''),
                    'cbte_nro': det_resp.CbteDesde,
                    'raw_cae_qr': raw_cae_qr,
                    'error': ''
                }
            else:
                # Si fue rechazado, recolectamos los mensajes de observación de AFIP
                obs_msg = ""
                if hasattr(det_resp, 'Observaciones') and det_resp.Observaciones:
                    obs_msg = "; ".join([f"Obs {obs.Code}: {obs.Msg}" for obs in det_resp.Observaciones.Obs])
                return {
                    'exito': False,
                    'cae': '',
                    'cae_vto': '',
                    'cbte_nro': det_resp.CbteDesde,
                    'error': f"Comprobante rechazado por AFIP. {obs_msg}"
                }
        else:
            return {
                'exito': False, 'cae': '', 'cae_vto': '', 'cbte_nro': '',
                'error': "Respuesta de AFIP sin detalle."
            }
    except Exception as e:
        return {'exito': False, 'cae': '', 'cae_vto': '', 'cbte_nro': '', 'error': str(e)}


def guardar_resultados_csv(filas, campos, archivo_salida):
    """
    Guarda los resultados de la facturación en un archivo CSV.
    """
    with open(archivo_salida, mode='w', encoding='utf-8', newline='') as file_out:
        escritor_csv = csv.DictWriter(file_out, fieldnames=campos)
        escritor_csv.writeheader()
        escritor_csv.writerows(filas)
    print(f"  CSV generado: {os.path.basename(archivo_salida)}")


def guardar_resultados_xlsx(filas, campos, archivo_salida):
    """
    Guarda los resultados de la facturación en un archivo Excel (.xlsx).
    Requiere la librería openpyxl.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "Resultados Facturación"

        # --- Estilos ---
        fuente_encabezado = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
        relleno_encabezado = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
        alineacion_centro = Alignment(horizontal='center', vertical='center')
        borde_fino = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        relleno_exito = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        relleno_error = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

        # --- Escribir encabezados ---
        for col_idx, campo in enumerate(campos, 1):
            celda = ws.cell(row=1, column=col_idx, value=campo)
            celda.font = fuente_encabezado
            celda.fill = relleno_encabezado
            celda.alignment = alineacion_centro
            celda.border = borde_fino

        # --- Escribir datos ---
        for row_idx, fila in enumerate(filas, 2):
            for col_idx, campo in enumerate(campos, 1):
                valor = fila.get(campo, '')
                celda = ws.cell(row=row_idx, column=col_idx, value=valor)
                celda.border = borde_fino
                celda.alignment = Alignment(horizontal='center')

            # Colorear la fila según el estado
            estado = fila.get('estado', '')
            if estado == 'FACTURADO':
                for col_idx in range(1, len(campos) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = relleno_exito
            elif estado == 'ERROR':
                for col_idx in range(1, len(campos) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = relleno_error

        # --- Ajustar ancho de columnas ---
        for col_idx, campo in enumerate(campos, 1):
            # Ancho mínimo basado en el nombre de la columna + margen
            ancho = max(len(str(campo)) + 4, 12)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = ancho

        # Congelar la primera fila (encabezados)
        ws.freeze_panes = 'A2'

        wb.save(archivo_salida)
        print(f"  XLSX generado: {os.path.basename(archivo_salida)}")

    except ImportError:
        print("  ADVERTENCIA: No se pudo generar el archivo .xlsx (falta la librería 'openpyxl').")
        print("  Instalala con: pip install openpyxl")
    except Exception as e:
        print(f"  ERROR al generar XLSX: {e}")


# ============================================================================
# PROGRAMA PRINCIPAL
# ============================================================================
if __name__ == "__main__":
    # --- Activar captura de consola a archivo .log ---
    # Se crea ANTES de cualquier print para capturar absolutamente todo.
    carpeta_logs = os.path.join(RUTA_BASE, "logs")
    os.makedirs(carpeta_logs, exist_ok=True)  # Crea la carpeta 'logs' si no existe

    fecha_hora_log = datetime.now().strftime('%Y%m%d_%H%M%S')
    archivo_log_path = os.path.join(carpeta_logs, f"facturacion_{fecha_hora_log}.log")
    _log_file = open(archivo_log_path, 'w', encoding='utf-8')
    sys.stdout = DuplicadorConsola(sys.__stdout__, _log_file)
    sys.stderr = DuplicadorConsola(sys.__stderr__, _log_file)
    atexit.register(_log_file.close)  # Cerrar el archivo al finalizar el programa

    print("=" * 60)
    print("  FACTURACIÓN ELECTRÓNICA AFIP - ARCA")
    print(f"  Log: {os.path.join('logs', os.path.basename(archivo_log_path))}")
    print("=" * 60)

    # Rutas de salida
    fecha_hoy = datetime.now().strftime('%Y%m%d')
    archivo_salida_csv = os.path.join(RUTA_BASE, f"facturacion_{fecha_hoy}.csv")
    archivo_salida_xlsx = os.path.join(RUTA_BASE, f"facturacion_{fecha_hoy}.xlsx")

    # Buscar archivo de entrada (XLSX o CSV)
    archivo_entrada_xlsx = os.path.join(RUTA_BASE, "plantilla_facturacion.xlsx")
    archivo_entrada_csv = os.path.join(RUTA_BASE, "plantilla_facturacion.csv")
    archivo_usado = None
    
    filas_a_procesar = []
    campos_entrada = []

    if os.path.exists(archivo_entrada_xlsx):
        archivo_usado = archivo_entrada_xlsx
        print(f"  Leyendo archivo Excel: {os.path.basename(archivo_usado)}")
        try:
            from openpyxl import load_workbook
            import datetime
            wb = load_workbook(archivo_usado, data_only=True)
            ws = wb.active
            # Leer encabezados
            campos_entrada = [str(cell.value).strip() for cell in ws[1] if cell.value is not None]
            # Leer datos
            for row in ws.iter_rows(min_row=2, values_only=True):
                if any(row):  # ignorar filas totalmente vacías
                    fila = {}
                    for i, campo in enumerate(campos_entrada):
                        if i < len(row):
                            val = row[i]
                            if isinstance(val, datetime.datetime):
                                # Si Excel lo detectó como fecha, formatearlo a AAAAMMDD
                                fila[campo] = val.strftime('%Y%m%d')
                            elif val is not None:
                                s_val = str(val).strip()
                                # Limpiar los ".0" que pone Excel a los números enteros
                                if s_val.endswith('.0'):
                                    s_val = s_val[:-2]
                                fila[campo] = s_val
                            else:
                                fila[campo] = ''
                        else:
                            fila[campo] = ''
                    filas_a_procesar.append(fila)
        except ImportError:
            print("  ERROR: Se encontró plantilla_facturacion.xlsx pero falta la librería 'openpyxl'.")
            input("\n  Presioná ENTER para salir...")
            sys.exit(1)
    elif os.path.exists(archivo_entrada_csv):
        archivo_usado = archivo_entrada_csv
        print(f"  Leyendo archivo CSV: {os.path.basename(archivo_usado)}")
        with open(archivo_usado, mode='r', encoding='utf-8') as f:
            # Detectar si Excel guardó con punto y coma
            primera_linea = f.readline()
            delimitador = ';' if ';' in primera_linea else ','
            f.seek(0)
            
            lector_csv = csv.DictReader(f, delimiter=delimitador)
            campos_entrada = lector_csv.fieldnames
            for fila in lector_csv:
                filas_a_procesar.append(fila)
    else:
        print(f"\n  ERROR: No se encontró el archivo 'plantilla_facturacion.xlsx' ni 'plantilla_facturacion.csv'")
        print(f"  Colocá la plantilla al lado del ejecutable.")
        input("\n  Presioná ENTER para salir...")
        sys.exit(1)

    try:
        # Inicializar conexión con AFIP
        wsfe_cliente = inicializar_wsfe()
    except FileNotFoundError as e:
        print(f"\n  ERROR DE CERTIFICADOS:\n  {e}")
        input("\n  Presioná ENTER para salir...")
        sys.exit(1)
    except Exception as e:
        print(f"\n  ERROR AL CONECTAR CON AFIP:\n  {e}")
        input("\n  Presioná ENTER para salir...")
        sys.exit(1)

    # Procesar las facturas
    filas_procesadas = []
    campos_salida = campos_entrada + ['cae', 'cae_vto', 'cbte_nro', 'estado', 'observaciones', 'raw_cae_qr']

    for fila in filas_a_procesar:
        # Ignorar filas donde el id_interno esté vacío
        if not fila.get('id_interno'):
            continue
            
        print(f"  Procesando ID interno {fila.get('id_interno')} (Doc: {fila.get('doc_nro', '')})...")

        respuesta = facturar_comprobante(wsfe_cliente, fila)

        if respuesta['exito']:
            fila['cae'] = respuesta['cae']
            fila['cae_vto'] = respuesta['cae_vto']
            fila['cbte_nro'] = respuesta['cbte_nro']
            fila['raw_cae_qr'] = respuesta.get('raw_cae_qr', '')
            fila['estado'] = 'FACTURADO'
            fila['observaciones'] = ''
            print(f"    ✓ FACTURADO - CAE: {respuesta['cae']}")
        else:
            fila['cae'] = ''
            fila['cae_vto'] = ''
            fila['cbte_nro'] = ''
            fila['raw_cae_qr'] = ''
            fila['estado'] = 'ERROR'
            fila['observaciones'] = respuesta['error']
            print(f"    ✗ ERROR: {respuesta['error']}")

        filas_procesadas.append(fila)

    # Guardar resultados en ambos formatos
    print(f"\n--- Generando archivos de resultados ---")
    guardar_resultados_csv(filas_procesadas, campos_salida, archivo_salida_csv)
    guardar_resultados_xlsx(filas_procesadas, campos_salida, archivo_salida_xlsx)

    exitosas = sum(1 for f in filas_procesadas if f.get('estado') == 'FACTURADO')
    fallidas = sum(1 for f in filas_procesadas if f.get('estado') == 'ERROR')

    print(f"\n{'=' * 60}")
    print(f"  Proceso finalizado. Comprobantes procesados: {len(filas_procesadas)}")
    print(f"    - Exitosas: {exitosas}")
    print(f"    - Fallidas: {fallidas}")
    print(f"  Archivos generados en: {RUTA_BASE}")
    print(f"{'=' * 60}")

    # Pausa para que el usuario vea el resultado cuando ejecuta el .exe
    input("\n  Presioná ENTER para salir...")
